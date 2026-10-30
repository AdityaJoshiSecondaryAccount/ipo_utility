"""
IPO Setup page E2E (/IPOSETUP, templates/IPOSETUP.html; edit page /edit/<id>, templates/edit.html).

The broker's list of IPOs. Runs on its OWN IPOs ("IPS ..."; group "IPS GROUP" for the orders the delete tests need);
setUp deletes and re-creates them, nothing else is touched. The list shows every IPO of the broker, so assertions on
the whole table are computed from the database at run time.

What the page offers:
  Table (server-side, sorted by name): checkbox + Sr No. | Name (link to the IPO's Orders page; names longer than 20
  characters are cut to 20 + "..." with the full name as title) | IPO Type | IPO Price | Pre-Open Price | Lot Retail /
  SHNI / BHNI ("-" when not set) | Total IPO size | Retail % | BHNI % / SHNI % ("-" for SME) | Remark ("-") | Edit, Delete.
  ADD IPO modal -> POST /AddIPO: only MAINBOARD (the SME choice is commented out); Name, IPO Upper Price Band, Total
  size, Retail / SHNI / BHNI lot sizes, Retail / BHNI / SHNI % (defaults 35 / 10 / 5), Remark. Pre-open price starts at
  the IPO price, the name is stored in capitals, a name that exists is refused, and the broker's IPO limit applies.
  Edit page -> POST /update/<id>; Delete (per row, a confirmation modal) -> /delete/<id>; bulk delete (checkboxes ->
  POST /BulkDeleteIPO). An IPO whose orders still have money to collect (order amounts minus the net of its accounting
  entries) is not deleted. Deleting removes its orders, order details and rate list; accounting entries keep the IPO's
  name and lose the link.
  Search box (server-side, debounced; also on Enter), column sort (server-side), page size (50 ... All, session) and
  page links, column filter boxes, Columns, Download (Excel / PDF).
"""
import datetime
import os
import re
import tempfile
import unittest

import openpyxl
from playwright.sync_api import expect

from home.models import Accounting, CurrentIpoName, GroupDetail, Order, OrderDetail, RateList
from test_cases.base import PersistentLiveServerTestCase, pause_for_inspection, settle

ALPHA, BETA, LONG = "IPS ALPHA", "IPS BETA SME", "IPS VERY LONG NAME FOR TRUNCATION"
GROUP = "IPS GROUP"

FIXTURE = {
    ALPHA: dict(IPOType="MAINBOARD", IPOPrice=100.0, PreOpenPrice=120.0, LotSizeRetail=10.0, LotSizeSHNI=140.0,
                LotSizeBHNI=700.0, TotalIPOSzie="500", RetailPercentage="35", BHNIPercentage="10", SHNIPercentage="5",
                Remark="ips-remark-alpha"),
    BETA: dict(IPOType="SME", IPOPrice=50.0, PreOpenPrice=50.0, LotSizeRetail=1200.0, LotSizeSHNI=None,
               LotSizeBHNI=None, TotalIPOSzie="20", RetailPercentage="100", BHNIPercentage="", SHNIPercentage="",
               Remark=""),
    LONG: dict(IPOType="MAINBOARD", IPOPrice=321.5, PreOpenPrice=330.0, LotSizeRetail=45.0, LotSizeSHNI=630.0,
               LotSizeBHNI=3150.0, TotalIPOSzie="1234.5", RetailPercentage="35", BHNIPercentage="10",
               SHNIPercentage="5", Remark="long one"),
}

ROWS_JS = """() => [...document.querySelectorAll('#example tbody tr')].filter(tr => tr.querySelector('.row-checkbox')).map(tr => {
    const td = [...tr.querySelectorAll(':scope > td')], clean = t => t.replace(/\\s+/g, ' ').trim();
    const link = td[1].querySelector('a');
    return {id: Number(tr.querySelector('.row-checkbox').value), cells: td.map(c => clean(c.textContent)),
            href: link?.getAttribute('href'), title: link?.getAttribute('title')};
})"""

NEW_IPO = {"name": "ips gamma new", "IPOPrice": "245.5", "TotalIPOSzie": "999.9", "LotSizeRetail": "60",
           "LotSizeSHNI": "840", "LotSizeBHNI": "4200", "RetailPercentage": "30", "BHNIPercentage": "20",
           "SHNIPercentage": "10", "Remark": "added by playwright"}
# The add form sits inside the <table> markup, so the browser moves its fields out of the <form> element (they still
# belong to it through the form owner): select them inside the #ADDIPO modal, never inside #addipoform.
ADD_FIELDS = {"name": "#name", "IPOPrice": "#IPOPrice", "TotalIPOSzie": '#ADDIPO input[name="TotalIPOSzie"]',
              "LotSizeRetail": "#LotSize", "LotSizeSHNI": "#LotSizeSH", "LotSizeBHNI": "#LotSizeBH",
              "RetailPercentage": '[id="Retail%"]', "BHNIPercentage": '[id="BHNI%"]', "SHNIPercentage": '[id="SHNI%"]',
              "Remark": '#ADDIPO input[name="Remark"]'}


def shown(value):
    """How the table prints a stored value (str of the Python value; None -> '-')."""
    return "-" if value is None else str(value)


class IPOSetupPlaywrightTestCase(PersistentLiveServerTestCase):
    PAGE_NAME = "ipo_setup"
    _ran_any = False

    # ------------------------------------------------------------------ setup / teardown
    @classmethod
    def tearDownClass(cls):
        cls.log_summary()
        if cls._ran_any:
            try:
                cls.page.goto(f"{cls.live_server_url}/IPOSETUP")
                settle(cls.page)
            except Exception:
                pass
            pause_for_inspection(cls.PAGE_NAME)
        super().tearDownClass()

    def setUp(self):
        super().setUp()
        type(self)._ran_any = True
        self.user.IPO_limit = None
        self.user.save()
        self.reset()
        self.js_errors = []
        self.page.on("pageerror", lambda e: self.js_errors.append(f"{e.message}\n{e.stack}"))
        self.clear_storage()

    def tearDown(self):
        self.user.IPO_limit = None
        self.user.save()
        self.clear_storage()
        super().tearDown()

    def clear_storage(self):
        try:
            if self.page.url.startswith(self.live_server_url):
                self.page.evaluate("localStorage.clear(); sessionStorage.clear()")
        except Exception:
            pass

    def ips(self):
        return CurrentIpoName.objects.filter(user=self.user, IPOName__startswith="IPS ")

    def reset(self):
        mine = self.ips()
        OrderDetail.objects.filter(user=self.user, Order__OrderIPOName__in=mine).delete()
        Order.objects.filter(user=self.user, OrderIPOName__in=mine).delete()
        RateList.objects.filter(user=self.user, RateListIPOName__in=mine).delete()
        Accounting.objects.filter(user=self.user, ipo__in=mine).delete()
        Accounting.objects.filter(user=self.user, ipo=None, ipo_name__startswith="IPS ").delete()
        mine.delete()
        self.group = GroupDetail.objects.get_or_create(user=self.user, GroupName=GROUP)[0]
        self.ipo = {name: CurrentIpoName.objects.create(user=self.user, IPOName=name, **fields)
                    for name, fields in FIXTURE.items()}

    def fillers(self, count=55):
        CurrentIpoName.objects.bulk_create([
            CurrentIpoName(user=self.user, IPOName=f"IPS FILL {i:02d}", IPOType="MAINBOARD", IPOPrice=10 + i,
                           PreOpenPrice=10 + i, LotSizeRetail=1, LotSizeSHNI=14, LotSizeBHNI=70, TotalIPOSzie="1",
                           RetailPercentage="35", BHNIPercentage="10", SHNIPercentage="5", Remark="")
            for i in range(count)])

    def with_order(self, ipo, amount):
        """An order (with one application) on `ipo` whose stored amount is `amount`."""
        order = Order.objects.create(user=self.user, OrderIPOName=ipo, OrderGroup=self.group, OrderType="BUY",
                                     OrderCategory="Kostak", InvestorType="RETAIL", Quantity=1, Rate=100,
                                     Amount=amount, OrderDate=datetime.date(2026, 3, 1),
                                     OrderTime=datetime.time(10, 0))
        OrderDetail.objects.create(user=self.user, Order=order, PreOpenPrice=ipo.PreOpenPrice)
        return order

    # ---------------------------------------------------------------------- page helpers
    def open(self):
        self.page.goto(f"{self.live_server_url}/IPOSETUP")
        settle(self.page)
        self.page.wait_for_selector("#example tbody")
        return self

    def rows(self):
        return self.page.evaluate(ROWS_JS)

    def row(self, ipo):
        return next(r for r in self.rows() if r["id"] == ipo.id)

    def shown_ids(self):
        return [r["id"] for r in self.rows()]

    def banner(self):
        return " | ".join(t.strip() for t in self.page.locator(".alert").all_inner_texts())

    def submit_and_wait(self, locator):
        with self.page.expect_navigation():
            locator.click()
        settle(self.page)

    def set_page_size(self, value):
        with self.page.expect_navigation():
            self.page.select_option("#Ip_page_size", value)
        settle(self.page)

    def search(self, text):
        box = self.page.locator('#example_filter input[type="search"]')
        box.fill(text)
        with self.page.expect_navigation():
            box.press("Enter")
        settle(self.page)

    def expected_cells(self, ipo, sr):
        name = ipo.IPOName.strip()
        sme = ipo.IPOType != "MAINBOARD"
        return [str(sr), name[:20] + "..." if len(name) > 20 else name, ipo.IPOType, str(ipo.IPOPrice),
                str(ipo.PreOpenPrice), shown(ipo.LotSizeRetail), shown(ipo.LotSizeSHNI), shown(ipo.LotSizeBHNI),
                ipo.TotalIPOSzie, ipo.RetailPercentage, "-" if sme else ipo.BHNIPercentage,
                "-" if sme else ipo.SHNIPercentage, ipo.Remark or "-", "Edit Delete"]

    def open_add_modal(self):
        self.page.locator("#add-ipo-btn").click()
        modal = self.page.locator("#ADDIPO")
        expect(modal).to_be_visible()
        return modal

    # ====================================================================== 1. the list
    def test_lists_every_ipo_sorted_by_name_with_exact_columns(self):
        self.open()
        expected = sorted(CurrentIpoName.objects.filter(user=self.user), key=lambda i: i.IPOName.strip().lower())
        self.assertLessEqual(len(expected), 50, "this test expects one page")
        rows = self.rows()
        self.assertEqual([r["id"] for r in rows], [i.id for i in expected])
        for sr, (row, ipo) in enumerate(zip(rows, expected), start=1):
            self.assertEqual(row["cells"], self.expected_cells(ipo, sr), ipo.IPOName)
            self.assertEqual(row["href"], f"/{ipo.id}/Order")
        expect(self.page.get_by_text(f"Showing 1 to {len(expected)} of {len(expected)} Rows")).to_be_visible()
        self.assertEqual(self.js_errors, [])

    def test_long_names_are_cut_with_the_full_name_as_title_and_link_to_the_orders_page(self):
        self.open()
        long_row = self.row(self.ipo[LONG])
        self.assertEqual(long_row["cells"][1], LONG[:20] + "...")
        self.assertEqual(long_row["title"], LONG)
        self.assertEqual(self.row(self.ipo[ALPHA])["cells"][1], ALPHA)
        with self.page.expect_navigation():
            self.page.locator(f'#example a[href="/{self.ipo[ALPHA].id}/Order"]').click()
        settle(self.page)
        self.assertRegex(self.page.url, rf"/{self.ipo[ALPHA].id}/Order")
        expect(self.page.locator("body")).not_to_contain_text("Server Error")

    # ====================================================================== 2. add an IPO
    def fill_add_form(self, values):
        for field, value in values.items():
            self.page.fill(ADD_FIELDS[field], value)

    def test_the_add_form_starts_with_the_default_percentages_and_only_mainboard(self):
        self.open()
        self.open_add_modal()
        self.assertEqual(self.page.input_value('[id="Retail%"]'), "35")
        self.assertEqual(self.page.input_value('[id="BHNI%"]'), "10")
        self.assertEqual(self.page.input_value('[id="SHNI%"]'), "5")
        self.assertEqual(self.page.input_value("#notype"), "MAINBOARD")
        self.assertEqual(self.page.locator("#smeipo").count(), 0, "the SME choice is not offered")
        for field in ("#LotSizeSH", "#LotSizeBH"):
            expect(self.page.locator(field)).to_be_visible()

    def test_adding_an_ipo_saves_every_field(self):
        before = CurrentIpoName.objects.filter(user=self.user).count()
        self.open()
        modal = self.open_add_modal()
        self.fill_add_form(NEW_IPO)
        self.submit_and_wait(modal.get_by_role("button", name="Submit Form"))
        self.assertIn("IPO Added successfully.", self.banner())
        self.assertRegex(self.page.url, r"/IPOSETUP$")
        self.assertEqual(CurrentIpoName.objects.filter(user=self.user).count(), before + 1)
        new = CurrentIpoName.objects.get(user=self.user, IPOName="IPS GAMMA NEW")       # stored in capitals
        self.assertEqual(
            (new.IPOType, new.IPOPrice, new.PreOpenPrice, new.TotalIPOSzie, new.LotSizeRetail, new.LotSizeSHNI,
             new.LotSizeBHNI, new.RetailPercentage, new.BHNIPercentage, new.SHNIPercentage, new.Remark, new.Active),
            ("MAINBOARD", 245.5, 245.5, "999.9", 60.0, 840.0, 4200.0, "30", "20", "10", "added by playwright", True))
        self.assertIn(new.id, self.shown_ids())

    def test_the_add_form_refuses_missing_required_fields(self):
        before = CurrentIpoName.objects.filter(user=self.user).count()
        self.open()
        modal = self.open_add_modal()
        for field in ("#name", "#IPOPrice", "#LotSize"):
            self.page.fill(field, "")
        url = self.page.url
        modal.get_by_role("button", name="Submit Form").click()
        self.page.wait_for_timeout(500)
        self.assertEqual(self.page.url, url, "an incomplete form is not sent")
        invalid = self.page.evaluate("[...document.getElementById('addipoform').elements].filter(e => !e.checkValidity()).map(e => e.name)")
        for name in ("name", "IPOPrice", "TotalIPOSzie", "LotSizeRetail", "LotSizeSHNI", "LotSizeBHNI"):
            self.assertIn(name, invalid)
        self.assertEqual(CurrentIpoName.objects.filter(user=self.user).count(), before)

    def test_number_fields_only_take_digits(self):
        self.open()
        self.open_add_modal()
        self.page.locator("#IPOPrice").press_sequentially("12a.5b")
        self.assertEqual(self.page.input_value("#IPOPrice"), "12.5")
        self.page.locator("#LotSize").press_sequentially("1x0")
        self.assertEqual(self.page.input_value("#LotSize"), "10")
        self.page.locator('[id="Retail%"]').fill("")
        self.page.locator('[id="Retail%"]').press_sequentially("4z0")
        self.assertEqual(self.page.input_value('[id="Retail%"]'), "40")

    def test_close_resets_the_add_form(self):
        self.open()
        modal = self.open_add_modal()
        self.page.fill("#name", "IPS NOT SAVED")
        modal.locator("button", has_text=re.compile(r"^Close$")).click()
        expect(modal).to_be_hidden()
        self.assertFalse(CurrentIpoName.objects.filter(user=self.user, IPOName="IPS NOT SAVED").exists())
        self.open_add_modal()
        self.assertEqual(self.page.input_value("#name"), "")

    def test_a_name_that_exists_is_refused_whatever_its_case(self):
        before = CurrentIpoName.objects.filter(user=self.user).count()
        self.open()
        modal = self.open_add_modal()
        self.fill_add_form({**NEW_IPO, "name": ALPHA.lower()})
        self.submit_and_wait(modal.get_by_role("button", name="Submit Form"))
        self.assertIn("IPO Already Exist.", self.banner())
        self.assertEqual(CurrentIpoName.objects.filter(user=self.user).count(), before)
        self.assertEqual(CurrentIpoName.objects.get(id=self.ipo[ALPHA].id).IPOPrice, 100.0, "the old IPO is untouched")

    def test_the_brokers_ipo_limit_is_enforced(self):
        count = CurrentIpoName.objects.filter(user=self.user).count()
        self.user.IPO_limit = count
        self.user.save()
        self.open()
        modal = self.open_add_modal()
        self.fill_add_form(NEW_IPO)
        self.submit_and_wait(modal.get_by_role("button", name="Submit Form"))
        self.assertIn(f"You have reached the limit of {count} IPO Limits.", self.banner())
        self.assertEqual(CurrentIpoName.objects.filter(user=self.user).count(), count)

    @unittest.expectedFailure
    def test_the_server_refuses_a_non_numeric_price_without_crashing(self):
        """KNOWN BUG: /AddIPO saves whatever it is sent; a non-numeric price raises a server error (500)."""
        self.open()
        token = self.page.locator('input[name="csrfmiddlewaretoken"]').first.input_value()
        response = self.page.request.post(f"{self.live_server_url}/AddIPO", form={
            "csrfmiddlewaretoken": token, "IPOType": "MAINBOARD", "name": "IPS BAD PRICE", "IPOPrice": "abc",
            "TotalIPOSzie": "1", "LotSizeRetail": "1", "LotSizeSHNI": "1", "LotSizeBHNI": "1",
            "RetailPercentage": "35", "BHNIPercentage": "10", "SHNIPercentage": "5"}, max_redirects=0)
        self.assertLess(response.status, 500)
        self.assertFalse(CurrentIpoName.objects.filter(user=self.user, IPOName="IPS BAD PRICE").exists())

    # ====================================================================== 3. edit an IPO
    def open_edit(self, name):
        ipo = self.ipo[name]
        self.open()
        with self.page.expect_navigation():
            self.page.locator(f'#example tr:has(.row-checkbox[value="{ipo.id}"])').get_by_role("button", name="Edit").click()
        settle(self.page)
        self.assertRegex(self.page.url, rf"/edit/{ipo.id}")
        return ipo

    def save_edit(self):
        self.submit_and_wait(self.page.locator('form[action^="/update/"] button[type="submit"]'))

    def test_the_edit_page_is_filled_with_the_ipo(self):
        ipo = self.open_edit(ALPHA)
        value = self.page.input_value
        self.assertEqual(value("#name"), ALPHA)
        self.assertEqual(value("#category"), "MAINBOARD")
        self.assertEqual(float(value("#IPOPrice")), 100.0)
        self.assertEqual(value('input[name="TotalIPOSize"]'), "500")
        self.assertEqual((float(value("#LotSize")), float(value("#LotSizeS")), float(value("#LotSizeB"))), (10.0, 140.0, 700.0))
        self.assertEqual((value('input[name="RetailPercentage"]'), value('[id="b%"]'), value('[id="s%"]')), ("35", "10", "5"))
        self.assertEqual(value('input[name="Remark"]'), "ips-remark-alpha")
        expect(self.page.locator('form[action^="/update/"]')).to_have_attribute("action", re.compile(rf"^/update/{ipo.id}"))

    def test_editing_saves_every_field(self):
        ipo = self.open_edit(ALPHA)
        for selector, value in (("#name", "ips alpha renamed"), ("#IPOPrice", "111.5"),
                                ('input[name="TotalIPOSize"]', "600"), ("#LotSize", "12"), ("#LotSizeS", "168"),
                                ("#LotSizeB", "840"), ('input[name="RetailPercentage"]', "40"), ('[id="b%"]', "12"),
                                ('[id="s%"]', "6"), ('input[name="Remark"]', "edited remark")):
            self.page.fill(selector, value)
        self.save_edit()
        self.assertIn("IPO Edit successfully.", self.banner())
        self.assertRegex(self.page.url, r"/IPOSETUP\?page=1$")
        ipo.refresh_from_db()
        self.assertEqual(
            (ipo.IPOName, ipo.IPOPrice, ipo.TotalIPOSzie, ipo.LotSizeRetail, ipo.LotSizeSHNI, ipo.LotSizeBHNI,
             ipo.RetailPercentage, ipo.BHNIPercentage, ipo.SHNIPercentage, ipo.Remark, ipo.IPOType),
            ("IPS ALPHA RENAMED", 111.5, "600", 12.0, 168.0, 840.0, "40", "12", "6", "edited remark", "MAINBOARD"))
        self.assertEqual(ipo.PreOpenPrice, 120.0, "the pre-open price is not edited here")
        row = self.row(ipo)
        self.assertEqual(row["cells"][1:4], ["IPS ALPHA RENAMED", "MAINBOARD", "111.5"])

    def test_saving_under_its_own_name_is_allowed_but_another_ipos_name_is_refused(self):
        ipo = self.open_edit(ALPHA)
        self.page.fill('input[name="Remark"]', "same name ok")
        self.save_edit()
        self.assertIn("IPO Edit successfully.", self.banner())
        ipo.refresh_from_db()
        self.assertEqual((ipo.IPOName, ipo.Remark), (ALPHA, "same name ok"))

        self.open_edit(ALPHA)
        self.page.fill("#name", LONG.lower())
        self.page.fill('input[name="Remark"]', "should not save")
        self.save_edit()
        self.assertIn("IPO Already Exist.", self.banner())
        self.assertRegex(self.page.url, rf"/edit/{ipo.id}")
        ipo.refresh_from_db()
        self.assertEqual((ipo.IPOName, ipo.Remark), (ALPHA, "same name ok"))

    def test_the_edit_page_refuses_an_empty_name_and_back_leaves_without_saving(self):
        ipo = self.open_edit(ALPHA)
        self.page.fill("#name", "")
        url = self.page.url
        self.page.locator('form[action^="/update/"] button[type="submit"]').click()
        self.page.wait_for_timeout(500)
        self.assertEqual(self.page.url, url)
        self.page.fill("#name", "IPS SHOULD NOT SAVE")
        self.submit_and_wait(self.page.get_by_role("button", name=re.compile("Back|Cancel", re.I)).first)
        self.assertRegex(self.page.url, r"/IPOSETUP$")
        ipo.refresh_from_db()
        self.assertEqual(ipo.IPOName, ALPHA)

    @unittest.expectedFailure
    def test_editing_an_sme_ipo_keeps_it_sme(self):
        """KNOWN BUG: the SME option of the type select is disabled, so the form sends no type and an SME IPO
        is saved with an empty type."""
        ipo = self.open_edit(BETA)
        self.page.fill('input[name="Remark"]', "sme edited")
        self.save_edit()
        ipo.refresh_from_db()
        self.assertEqual(ipo.Remark, "sme edited")
        self.assertEqual(ipo.IPOType, "SME")

    @unittest.expectedFailure
    def test_after_editing_the_list_returns_to_the_page_the_edit_came_from(self):
        """KNOWN BUG: the edit form posts to /update/<id>?page<N> (no '='), so the list always reopens on page 1."""
        self.fillers(55)
        self.open()
        with self.page.expect_navigation():
            self.page.locator(".pagination .page-link", has_text=re.compile(r"^2$")).first.click()
        settle(self.page)
        target = self.rows()[0]["id"]
        with self.page.expect_navigation():
            self.page.locator(f'#example tr:has(.row-checkbox[value="{target}"])').get_by_role("button", name="Edit").click()
        settle(self.page)
        self.save_edit()
        self.assertRegex(self.page.url, r"page=2")

    @unittest.expectedFailure
    def test_changing_the_ipo_price_recalculates_the_order_amounts(self):
        """KNOWN BUG: /update/<id> saves the new IPO price but does not recalculate the stored order amounts."""
        ipo = self.ipo[ALPHA]
        order = self.with_order(ipo, 0)
        OrderDetail.objects.filter(Order=order).update(AllotedQty=10, PreOpenPrice=120)
        from home.views import calculate
        calculate(ipo.id, self.user)
        order.refresh_from_db()
        self.assertAlmostEqual(order.Amount, (120 - 100) * 10 - 100, places=2)          # (pre-open - price) x qty - rate
        self.open_edit(ALPHA)
        self.page.fill("#IPOPrice", "110")
        self.save_edit()
        order.refresh_from_db()
        self.assertAlmostEqual(order.Amount, (120 - 110) * 10 - 100, places=2)

    # ====================================================================== 4. delete
    def delete_button(self, ipo):
        return self.page.locator(f'#example tr:has(.row-checkbox[value="{ipo.id}"])').get_by_role("button", name="Delete")

    def open_delete_modal(self, ipo):
        self.delete_button(ipo).click()
        modal = self.page.locator(f'[id="{ipo.id}"].modal')
        expect(modal).to_be_visible()
        expect(modal).to_contain_text(f"Are you sure you want to delete {ipo.IPOName} IPO?")
        return modal

    def test_delete_asks_first_and_cancel_keeps_the_ipo(self):
        ipo = self.ipo[ALPHA]
        self.open()
        modal = self.open_delete_modal(ipo)
        modal.get_by_role("button", name="Cancel").click()
        expect(modal).to_be_hidden()
        self.assertTrue(CurrentIpoName.objects.filter(id=ipo.id).exists())
        expect(self.page.locator(".modal-backdrop")).to_have_count(0)
        self.page.locator(f'.row-checkbox[value="{ipo.id}"]').check(timeout=3000)      # the page is usable again

    def test_deleting_an_ipo_removes_its_orders_and_keeps_its_accounting_by_name(self):
        ipo = self.ipo[ALPHA]
        order = self.with_order(ipo, 0)
        rates = RateList.objects.create(user=self.user, RateListIPOName=ipo)
        now = datetime.datetime(2026, 3, 1, 10, 0, tzinfo=datetime.timezone.utc)
        entries = [Accounting.objects.create(user=self.user, ipo=ipo, group=self.group, amount=250, amount_type=kind,
                                             date_time=now, remark="ips ledger") for kind in ("credit", "debit")]
        self.open()
        modal = self.open_delete_modal(ipo)
        self.submit_and_wait(modal.get_by_role("button", name="Delete"))
        self.assertIn(f"IPO '{ALPHA}' successfully deleted.", self.banner())
        self.assertFalse(CurrentIpoName.objects.filter(id=ipo.id).exists())
        self.assertFalse(Order.objects.filter(id=order.id).exists())
        self.assertFalse(OrderDetail.objects.filter(Order_id=order.id).exists())
        self.assertFalse(RateList.objects.filter(id=rates.id).exists())
        for entry in entries:
            entry.refresh_from_db()
            self.assertEqual((entry.ipo_id, entry.ipo_name, entry.status), (None, ALPHA, True))
        self.assertNotIn(ipo.id, self.shown_ids())
        for other in (BETA, LONG):
            self.assertTrue(CurrentIpoName.objects.filter(id=self.ipo[other].id).exists())

    def test_an_ipo_with_money_still_to_collect_is_not_deleted(self):
        ipo = self.ipo[ALPHA]
        order = self.with_order(ipo, 500)
        Accounting.objects.create(user=self.user, ipo=ipo, group=self.group, amount=200, amount_type="credit",
                                  date_time=datetime.datetime(2026, 3, 1, tzinfo=datetime.timezone.utc))
        self.open()
        modal = self.open_delete_modal(ipo)
        self.submit_and_wait(modal.get_by_role("button", name="Delete"))
        self.assertIn(f"IPO '{ALPHA}' cannot be deleted as it has a non-zero collection amount of 300.0", self.banner())
        self.assertTrue(CurrentIpoName.objects.filter(id=ipo.id).exists())
        self.assertTrue(Order.objects.filter(id=order.id).exists())

    def test_ticking_rows_swaps_add_for_delete(self):
        self.open()
        add, delete = self.page.locator("#add-ipo-btn"), self.page.locator("#bulk-delete-btn")
        expect(add).to_be_visible()
        expect(delete).to_be_hidden()
        self.page.locator(f'.row-checkbox[value="{self.ipo[ALPHA].id}"]').check()
        expect(delete).to_be_visible()
        expect(add).to_be_hidden()
        self.page.locator("#selectAll").check()
        boxes = self.page.locator(".row-checkbox")
        self.assertEqual(boxes.evaluate_all("els => els.filter(e => e.checked).length"), boxes.count())
        self.page.locator(f'.row-checkbox[value="{self.ipo[BETA].id}"]').uncheck()
        expect(self.page.locator("#selectAll")).not_to_be_checked()
        self.page.locator("#selectAll").check()
        self.page.locator("#selectAll").uncheck()
        expect(delete).to_be_hidden()
        expect(add).to_be_visible()

    def test_bulk_delete_removes_the_settled_ipos_and_names_the_ones_it_kept(self):
        blocked = self.ipo[LONG]
        self.with_order(blocked, 750)
        self.with_order(self.ipo[ALPHA], 0)
        self.open()
        for name in (ALPHA, BETA, LONG):
            self.page.locator(f'.row-checkbox[value="{self.ipo[name].id}"]').check()
        self.page.locator("#bulk-delete-btn").click()
        modal = self.page.locator("#bulkDeleteModal")
        expect(modal).to_be_visible()
        expect(modal).to_contain_text("Are you sure you want to delete the selected IPO(s)?")
        with self.page.expect_navigation():
            modal.locator("#confirmBulkDelete").click()
        settle(self.page)
        banner = self.banner()
        self.assertIn("Successfully deleted 2 IPO(s).", banner)
        self.assertIn(f"1 IPO(s) could not be deleted (Non-zero collection amount): {LONG} (Collection: 750.0).", banner)
        self.assertEqual(list(self.ips().values_list("id", flat=True)), [blocked.id])
        self.assertFalse(Order.objects.filter(OrderIPOName_id=self.ipo[ALPHA].id).exists())

    @unittest.expectedFailure
    def test_cancelling_the_bulk_delete_leaves_the_page_usable(self):
        """KNOWN BUG: the Delete button opens the dialog through both Bootstrap 4 and 5 (two backdrops); Cancel
        removes one, the other stays over the page and blocks every click."""
        self.open()
        self.page.locator(f'.row-checkbox[value="{self.ipo[ALPHA].id}"]').check()
        self.page.locator("#bulk-delete-btn").click()
        modal = self.page.locator("#bulkDeleteModal")
        expect(modal).to_be_visible()
        modal.get_by_role("button", name="Cancel").click()
        expect(modal).to_be_hidden()
        self.assertTrue(self.ips().filter(id=self.ipo[ALPHA].id).exists(), "Cancel deletes nothing")
        self.assertEqual(self.page.locator(".modal-backdrop").count(), 0)

    # ====================================================================== 5. search, sort, pages
    def test_search_looks_at_name_type_numbers_and_remark(self):
        cases = [("ips alpha", {ALPHA}), ("ips-remark", {ALPHA}), ("321.5", {LONG}), ("1234.5", {LONG}),
                 ("SME", {BETA}), ("truncation", {LONG}), ("zzz-nothing", set())]
        for term, names in cases:
            with self.subTest(search=term):
                self.open()
                self.search(term)
                ids = set(self.shown_ids())
                mine = {i.id for i in self.ips()}
                self.assertEqual(ids & mine, {self.ipo[n].id for n in names})
                if term != "SME":                       # other pages have SME IPOs too
                    self.assertEqual(ids - mine, set())
                self.assertEqual(self.page.input_value('#example_filter input[type="search"]'), term)
        self.open()
        self.search("")
        self.assertEqual(len(self.rows()), CurrentIpoName.objects.filter(user=self.user).count())

    def test_typing_searches_by_itself(self):
        self.open()
        box = self.page.locator('#example_filter input[type="search"]')
        box.click()
        with self.page.expect_navigation(timeout=5000):
            box.press_sequentially("IPS BETA", delay=40)
        settle(self.page)
        self.assertEqual(self.shown_ids(), [self.ipo[BETA].id])

    def sort_by(self, title):
        header = self.page.locator(".dataTables_scrollHead thead tr:not(.filters) th", has_text=re.compile(rf"^\s*{re.escape(title)}\s*$"))
        with self.page.expect_navigation():
            header.first.click()
        settle(self.page)

    def test_clicking_a_header_sorts_on_the_server(self):
        # Compare the sort KEYS, not ids: in the full pipeline many IPOs share a price, and the server
        # breaks ties in its own order.
        by_id = {i.id: i for i in CurrentIpoName.objects.filter(user=self.user)}
        cases = [("IPO Price", lambda i: float(i.IPOPrice or 0)), ("Name", lambda i: i.IPOName.strip().lower())]
        for title, key in cases:
            with self.subTest(column=title):
                self.open()
                self.sort_by(title)
                first = [key(by_id[i]) for i in self.shown_ids()]
                self.assertEqual(sorted(self.shown_ids()), sorted(by_id), "every IPO listed once")
                self.assertIn(first, (sorted(first), sorted(first, reverse=True)))
                self.sort_by(title)
                second = [key(by_id[i]) for i in self.shown_ids()]
                self.assertEqual({tuple(first), tuple(second)}, {tuple(sorted(first)), tuple(sorted(first, reverse=True))},
                                 "the two clicks give both directions")

    def test_fifty_rows_a_page_the_page_links_and_all(self):
        self.fillers(55)
        total = CurrentIpoName.objects.filter(user=self.user).count()
        self.open()
        self.assertEqual(len(self.rows()), 50)
        expect(self.page.get_by_text(f"Showing 1 to 50 of {total} Rows")).to_be_visible()
        with self.page.expect_navigation():
            self.page.locator(".pagination .page-link", has_text=re.compile(r"^2$")).first.click()
        settle(self.page)
        rows = self.rows()
        self.assertEqual(len(rows), min(50, total - 50))
        self.assertEqual(rows[0]["cells"][0], "51")
        self.set_page_size("All")
        expect(self.page.get_by_text(f"Showing 1 to {total} of {total} Rows")).to_be_visible()
        self.open()                                                   # remembered for the session
        expect(self.page.get_by_text(f"Showing 1 to {total} of {total} Rows")).to_be_visible()
        self.set_page_size("50")
        self.assertEqual(len(self.rows()), 50)

    @unittest.expectedFailure
    def test_page_size_all_shows_every_ipo(self):
        """KNOWN BUG: the table remembers its 50-row length (stateSave) and its own pager is removed, so with "All"
        the server sends every IPO but only 50 are shown and the rest cannot be reached."""
        self.fillers(55)
        total = CurrentIpoName.objects.filter(user=self.user).count()
        self.open()
        self.set_page_size("All")
        expect(self.page.get_by_text(f"Showing 1 to {total} of {total} Rows")).to_be_visible()
        self.assertEqual(len(self.rows()), total)

    @unittest.expectedFailure
    def test_the_search_is_kept_when_moving_to_page_2(self):
        """KNOWN BUG: the search is sent by POST but the page links are plain GET links, so page 2 drops it."""
        self.fillers(55)
        self.open()
        self.search("IPS FILL")
        self.assertEqual(len(self.rows()), 50)
        with self.page.expect_navigation():
            self.page.locator(".pagination .page-link", has_text=re.compile(r"^2$")).first.click()
        settle(self.page)
        names = [r["cells"][1] for r in self.rows()]
        self.assertEqual(len(names), 5)
        self.assertTrue(all(n.startswith("IPS FILL") for n in names), names)

    def test_column_filter_box_narrows_the_page(self):
        self.open()
        self.page.locator('.dataTables_scrollHead .filters input[placeholder^="IPO Type"]').click()
        self.page.keyboard.type("SME")
        self.page.wait_for_timeout(400)
        types = {r["cells"][2] for r in self.rows()}
        self.assertEqual(types, {"SME"})
        self.assertIn(self.ipo[BETA].id, self.shown_ids())

    # ====================================================================== 6. download
    def test_the_excel_download_lists_every_ipo_on_the_page(self):
        expected = {i.id: i for i in CurrentIpoName.objects.filter(user=self.user)}
        self.open()
        self.page.get_by_role("button", name="Download").click()
        self.page.locator(".dt-button-collection .dt-button", has_text="Excel").click()
        popup = self.page.locator(".swal2-popup")
        expect(popup).to_be_visible()
        with self.page.expect_download() as info:
            popup.locator(".swal2-confirm").click()
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, info.value.suggested_filename)
            info.value.save_as(path)
            self.assertTrue(path.endswith(".xlsx"), path)
            rows = list(openpyxl.load_workbook(path).active.iter_rows(values_only=True))
        head = next(i for i, r in enumerate(rows) if r and any(str(c).strip().startswith("IPO Price") for c in r if c))
        cols = {str(c).strip(): i for i, c in enumerate(rows[head]) if c}
        self.assertNotIn("Action", cols)
        data = [r for r in rows[head + 1:] if any(c is not None for c in r)]
        self.assertEqual(len(data), len(expected))
        by_price_type = sorted((str(r[cols["IPO Type"]]), float(r[cols["IPO Price"]])) for r in data)
        self.assertEqual(by_price_type, sorted((i.IPOType, i.IPOPrice) for i in expected.values()))
        names = {str(r[cols["Name"]]) for r in data}
        self.assertIn(ALPHA, names)

    @unittest.expectedFailure
    def test_the_excel_download_has_full_ipo_names(self):
        """KNOWN BUG: the export copies the table text, so names longer than 20 characters are cut to '...'."""
        self.open()
        self.page.get_by_role("button", name="Download").click()
        self.page.locator(".dt-button-collection .dt-button", has_text="Excel").click()
        with self.page.expect_download() as info:
            self.page.locator(".swal2-popup .swal2-confirm").click()
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, info.value.suggested_filename)
            info.value.save_as(path)
            cells = [str(c) for r in openpyxl.load_workbook(path).active.iter_rows(values_only=True) for c in r if c]
        self.assertIn(LONG, cells)
