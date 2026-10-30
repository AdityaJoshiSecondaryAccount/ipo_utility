"""
Client Setup page E2E (/ClientSetup, templates/ClientSetup.html; /EditClient/<id>, templates/EditClient.html).

The broker's list of clients (PAN holders). Runs on its OWN groups (CLS ALPHA, CLS BETA, CLS BULK) and clients whose PANs
start with "CLS" (plus one IPO, CLS IPO, whose order uses a client), reset in setUp, so no other page's clients are touched.

What the page does:
  Table (server side, 50 rows a page by default): checkbox + Sr No. | PAN No. | Name | Group | Client-ID/DP-ID |
      Edit (-> /EditClient/<id>) and Delete (modal -> /DeleteClient/<id>; refused while an order application uses the PAN).
  Toolbar: Show (50 / 100 / 200 / 250 / All rows; POST, kept in the session), Group filter (POST, kept in the session),
      Search (server side over PAN / Name / Group / Client-ID, submitted 0.5 s after typing), Columns, Download
      (Excel / PDF), "Delete All Clients" (deletes every unused client of the chosen group and downloads a status
      sheet) which turns into "Delete Selected" once rows are ticked (/BulkDeleteClients).
  ADD CLIENT modal -> /AddClient: PAN (10 characters, required), Name (required), Group (required), Client-ID/DP-ID
      (16 characters, optional). PAN and Name are stored in capitals; a PAN that already exists is refused; the
      broker's client limit (CustomUser.Client_limit) is enforced.
  /EditClient/<id> -> /UpdateClient/<id>: the same fields; a PAN used by another client is refused.
  /autocomplete1 (POST PAN) returns the name of the client with that PAN (used by the order detail pages).
"""
import os
import re
import tempfile
import unittest

import openpyxl
from playwright.sync_api import expect

from home.models import ClientDetail, CurrentIpoName, GroupDetail, Order, OrderDetail
from test_cases.base import PersistentLiveServerTestCase, pause_for_inspection, settle

ALPHA, BETA, BULK = "CLS ALPHA", "CLS BETA", "CLS BULK"
CLS_IPO = "CLS IPO"

# key -> (PAN, name, group, client id)
CLIENTS = {
    "a1": ("CLSAA1001A", "CLS ANAND", ALPHA, "1203320000000001"),
    "a2": ("CLSAA1002B", "CLS BHAVNA", ALPHA, ""),
    "a3": ("CLSAA1003C", "CLS CHIRAG", ALPHA, "IN30000000000003"),
    "b1": ("CLSBB2001D", "CLS DEEPA", BETA, "1203320000000004"),
    "b2": ("CLSBB2002E", "CLS EKTA", BETA, ""),
}
USED = "b1"            # has an order application on CLS IPO -> cannot be deleted

ROWS_JS = """() => [...document.querySelectorAll('#example tbody tr')].filter(tr => !tr.querySelector('td.dataTables_empty')).map(tr => {
    const td = [...tr.querySelectorAll(':scope > td')];
    const clean = t => t.replace(/\\s+/g, ' ').trim();
    return {id: Number(tr.querySelector('.client-checkbox')?.value), cells: td.map(c => clean(c.textContent))};
})"""


class ClientSetupPlaywrightTestCase(PersistentLiveServerTestCase):
    PAGE_NAME = "client_setup"
    _ran_any = False

    # ------------------------------------------------------------------ setup / teardown
    @classmethod
    def tearDownClass(cls):
        cls.log_summary()
        if cls._ran_any:
            try:
                cls.page.goto(f"{cls.live_server_url}/ClientSetup")
                settle(cls.page)
            except Exception:
                pass
            pause_for_inspection(cls.PAGE_NAME)
        super().tearDownClass()

    def setUp(self):
        super().setUp()
        type(self)._ran_any = True
        self.user.Client_limit = None
        self.user.save()
        self.groups = {n: GroupDetail.objects.get_or_create(user=self.user, GroupName=n)[0] for n in (ALPHA, BETA, BULK)}
        self.cls_ipo = CurrentIpoName.objects.get_or_create(
            user=self.user, IPOName=CLS_IPO,
            defaults=dict(IPOType="MAINBOARD", IPOPrice=100, PreOpenPrice=150, LotSizeRetail=100, LotSizeSHNI=1400,
                          LotSizeBHNI=7000, Active=True))[0]
        self.reset()
        self.js_errors = []
        self.page.on("pageerror", lambda e: self.js_errors.append(f"{e.message}\n{e.stack}"))

    def tearDown(self):
        self.user.Client_limit = None
        self.user.save()
        try:
            if self.page.url.startswith(self.live_server_url):
                self.page.evaluate("sessionStorage.clear()")
        except Exception:
            pass
        super().tearDown()

    def mine(self):
        """Every client this page created: in a CLS group, or with a CLS PAN / name."""
        from django.db.models import Q
        return ClientDetail.objects.filter(user=self.user).filter(
            Q(Group__in=self.groups.values()) | Q(PANNo__startswith="CLS") | Q(Name__startswith="CLS "))

    def reset(self):
        OrderDetail.objects.filter(user=self.user, Order__OrderIPOName=self.cls_ipo).delete()
        Order.objects.filter(user=self.user, OrderIPOName=self.cls_ipo).delete()
        self.mine().delete()
        self.c = {}
        for key, (pan, name, group, dp) in CLIENTS.items():
            self.c[key] = ClientDetail.objects.create(user=self.user, PANNo=pan, Name=name, Group=self.groups[group],
                                                      ClientIdDpId=dp)
        order = Order.objects.create(user=self.user, OrderIPOName=self.cls_ipo, OrderGroup=self.groups[BETA],
                                     OrderType="BUY", OrderCategory="Kostak", InvestorType="RETAIL", Quantity=1,
                                     Rate=100, OrderDate="2026-02-10", OrderTime="09:00:00")
        self.used_detail = OrderDetail.objects.create(user=self.user, Order=order, OrderDetailPANNo=self.c[USED],
                                                      PreOpenPrice=150)

    # ---------------------------------------------------------------------- page helpers
    def open(self, group=None):
        self.page.goto(f"{self.live_server_url}/ClientSetup")
        settle(self.page)
        expect(self.page.locator("h3.cl")).to_have_text("CLIENT")
        if group:
            self.choose("#group_filter", group)
        return self

    def choose(self, selector, value):
        with self.page.expect_navigation():
            self.page.select_option(selector, value)
        settle(self.page)

    def search(self, text):
        with self.page.expect_navigation():
            self.page.locator("#table-search").fill(text)
        settle(self.page)

    def rows(self):
        return self.page.evaluate(ROWS_JS)

    def shown_ids(self):
        return [r["id"] for r in self.rows()]

    def ids(self, *keys):
        return {self.c[k].id for k in keys}

    def banner(self):
        return " | ".join(t.strip() for t in self.page.locator(".alert").all_text_contents())

    def total_text(self):
        return self.page.locator("p", has_text=re.compile(r"Showing .* Rows")).first.inner_text().strip()

    @staticmethod
    def expected_cells(client, sr):
        return [str(sr), client.PANNo, client.Name, client.Group.GroupName if client.Group_id else "",
                client.ClientIdDpId, "Edit Delete"]

    def add_client(self, pan, name, group, dp=None, submit=True):
        self.page.get_by_role("button", name="ADD CLIENT").click()
        modal = self.page.locator("#ADDCLIENT")
        expect(modal).to_be_visible()
        modal.locator("#PANNo").fill(pan)
        modal.locator("#Name").fill(name)
        if group:
            modal.locator("#group").select_option(group)
        if dp is not None:
            modal.locator('input[name="ClientIdDpId"]').fill(dp)
        if submit:
            with self.page.expect_navigation():
                modal.get_by_role("button", name="Submit Form").click()
            settle(self.page)
        return modal

    def open_edit(self, key):
        client = self.c[key]
        self.open(self.groups[ALPHA if client.Group.GroupName == ALPHA else BETA].GroupName)
        with self.page.expect_navigation():
            self.page.locator(f'tr:has(.client-checkbox[value="{client.id}"])').get_by_role("button", name="Edit").click()
        settle(self.page)
        expect(self.page.locator("h3")).to_have_text("Update Client Details")
        return client

    def submit_edit(self):
        with self.page.expect_navigation():
            self.page.get_by_role("button", name="Submit").click()
        settle(self.page)

    # ====================================================================== 1. the list
    def test_a_group_shows_its_clients_with_exact_cells(self):
        self.open(ALPHA)
        want = list(ClientDetail.objects.filter(user=self.user, Group=self.groups[ALPHA]).order_by("id"))
        rows = self.rows()
        self.assertEqual(sorted(r["id"] for r in rows), sorted(c.id for c in want))
        by_id = {c.id: c for c in want}
        for sr, row in enumerate(rows, start=1):
            self.assertEqual(row["cells"], self.expected_cells(by_id[row["id"]], sr))
        self.assertEqual(self.total_text(), f"Showing 1 to {len(want)} of {len(want)} Rows")
        self.assertEqual(self.page.input_value("#group_filter"), ALPHA)
        headers = self.page.locator(".dataTables_scrollHead thead tr").first.locator("th").all_inner_texts()
        self.assertEqual([h.strip() for h in headers], ["Sr No.", "PAN No.", "Name", "Group", "Client-ID/DP-ID", "Action"])
        self.assertEqual(self.js_errors, [])

    def test_all_groups_counts_every_client_of_the_broker(self):
        self.open()
        total = ClientDetail.objects.filter(user=self.user).count()
        self.assertRegex(self.total_text(), rf"^Showing 1 to {min(total, 50)} of {total} Rows$")
        self.assertEqual([r["cells"][0] for r in self.rows()], [str(i) for i in range(1, min(total, 50) + 1)])
        options = self.page.eval_on_selector_all("#group_filter option", "os => os.map(o => o.value)")
        self.assertEqual(options[0], "All")
        self.assertEqual(options[1:], sorted(GroupDetail.objects.filter(user=self.user).values_list("GroupName", flat=True)))

    def test_the_group_filter_is_remembered_and_all_group_brings_everything_back(self):
        self.open(BETA)
        self.assertEqual(set(self.shown_ids()), self.ids("b1", "b2"))
        self.page.goto(f"{self.live_server_url}/ClientSetup")             # a fresh visit keeps the group
        settle(self.page)
        self.assertEqual(self.page.input_value("#group_filter"), BETA)
        self.assertEqual(set(self.shown_ids()), self.ids("b1", "b2"))
        self.choose("#group_filter", "All")
        total = ClientDetail.objects.filter(user=self.user).count()
        self.assertIn(f"of {total} Rows", self.total_text())

    # ====================================================================== 2. search, page size, pages
    def test_search_looks_at_pan_name_group_and_client_id(self):
        cases = [("CLSAA1002", ["a2"]), ("clsbb", ["b1", "b2"]), ("cls chirag", ["a3"]), ("CLS BETA", ["b1", "b2"]),
                 ("IN30000000000003", ["a3"]), ("1203320000000004", ["b1"]), ("CLS ZZZ NOBODY", [])]
        for term, keys in cases:
            with self.subTest(search=term):
                self.open()
                self.search(term)
                mine = {i for i in self.shown_ids() if i in {c.id for c in self.c.values()}}
                self.assertEqual(mine, self.ids(*keys))
                self.assertEqual(self.page.input_value("#table-search"), term)

    def test_search_and_group_filter_work_together_and_a_fresh_visit_clears_the_search(self):
        self.open(ALPHA)
        self.search("BHAVNA")
        self.assertEqual(set(self.shown_ids()), self.ids("a2"))
        self.page.goto(f"{self.live_server_url}/ClientSetup")
        settle(self.page)
        self.assertEqual(self.page.input_value("#table-search"), "")
        self.assertEqual(set(self.shown_ids()), self.ids("a1", "a2", "a3"))

    def bulk_clients(self, count):
        ClientDetail.objects.bulk_create([
            ClientDetail(user=self.user, PANNo=f"CLSCC{3000 + i}F", Name=f"CLS BULK {i:03d}", Group=self.groups[BULK],
                         ClientIdDpId="") for i in range(count)])

    def test_pages_of_fifty_page_links_and_page_size(self):
        self.bulk_clients(55)
        self.open(BULK)
        self.assertEqual(len(self.rows()), 50)
        self.assertEqual(self.total_text(), "Showing 1 to 50 of 55 Rows")
        with self.page.expect_navigation():
            self.page.locator(".pagination .page-link", has_text=re.compile(r"^2$")).click()
        settle(self.page)
        rows = self.rows()
        self.assertEqual([r["cells"][0] for r in rows], [str(i) for i in range(51, 56)])
        self.assertEqual(self.total_text(), "Showing 51 to 55 of 55 Rows")
        self.assertEqual(self.page.input_value("#group_filter"), BULK, "the group stays while paging")
        # page 1 and page 2 together are exactly the 55 clients
        self.choose("#client_page_size", "All")
        self.assertEqual(len(self.rows()), 55)
        self.assertEqual(self.page.input_value("#client_page_size"), "All")
        self.choose("#client_page_size", "100")
        self.assertEqual(len(self.rows()), 55)
        self.assertEqual(self.page.input_value("#client_page_size"), "100")
        names = sorted(r["cells"][2] for r in self.rows())
        self.assertEqual(names, [f"CLS BULK {i:03d}" for i in range(55)])

    # ====================================================================== 3. add a client
    def test_adding_a_client_saves_every_field_in_capitals(self):
        self.open()
        self.add_client("clsdd4001g", "cls new person", BETA, "1203320000000099")
        self.assertIn("Client Added successfully.", self.banner())
        client = ClientDetail.objects.get(user=self.user, PANNo="CLSDD4001G")
        self.assertEqual((client.Name, client.Group.GroupName, client.ClientIdDpId),
                         ("CLS NEW PERSON", BETA, "1203320000000099"))
        self.open(BETA)
        row = next(r for r in self.rows() if r["id"] == client.id)
        self.assertEqual(row["cells"][1:5], ["CLSDD4001G", "CLS NEW PERSON", BETA, "1203320000000099"])

    def test_the_client_id_is_optional(self):
        self.open()
        self.add_client("CLSDD4002H", "CLS NO DP", ALPHA)
        self.assertEqual(ClientDetail.objects.get(user=self.user, PANNo="CLSDD4002H").ClientIdDpId, "")

    def test_a_pan_that_already_exists_is_refused_in_any_case(self):
        self.open()
        before = ClientDetail.objects.filter(user=self.user).count()
        self.add_client(CLIENTS["a1"][0].lower(), "CLS SOMEONE ELSE", BETA)
        self.assertIn("Client Already Exist.", self.banner())
        self.assertEqual(ClientDetail.objects.filter(user=self.user).count(), before)
        self.c["a1"].refresh_from_db()
        self.assertEqual((self.c["a1"].Name, self.c["a1"].Group.GroupName), ("CLS ANAND", ALPHA))

    def test_the_form_refuses_missing_or_short_fields(self):
        self.open()
        before = ClientDetail.objects.filter(user=self.user).count()
        cases = [
            ({"pan": "", "name": "CLS X", "group": ALPHA}, "#PANNo", "valueMissing"),
            ({"pan": "CLSEE5001J", "name": "", "group": ALPHA}, "#Name", "valueMissing"),
            ({"pan": "CLSEE5001J", "name": "CLS X", "group": None}, "#group", "valueMissing"),
            ({"pan": "CLSEE50", "name": "CLS X", "group": ALPHA}, "#PANNo", "tooShort"),
            ({"pan": "CLSEE5001J", "name": "CLS X", "group": ALPHA, "dp": "12345"}, 'input[name="ClientIdDpId"]', "tooShort"),
        ]
        for fields, field, problem in cases:
            with self.subTest(field=field, problem=problem):
                self.open()
                modal = self.add_client(fields["pan"], fields["name"], fields["group"], fields.get("dp"), submit=False)
                url = self.page.url
                modal.get_by_role("button", name="Submit Form").click()
                self.page.wait_for_timeout(400)
                self.assertEqual(self.page.url, url)
                self.assertTrue(modal.locator(field).evaluate(f"e => e.validity.{problem}"))
                expect(modal.locator("form")).to_have_class(re.compile("was-validated"))
        self.assertEqual(ClientDetail.objects.filter(user=self.user).count(), before)
        # the PAN box takes at most 10 characters
        self.assertEqual(self.page.locator("#PANNo").get_attribute("maxlength"), "10")

    def test_close_leaves_the_list_unchanged(self):
        self.open()
        before = ClientDetail.objects.filter(user=self.user).count()
        modal = self.add_client("CLSEE5002K", "CLS CLOSED", ALPHA, submit=False)
        modal.locator(".modal-footer").get_by_role("button", name="Close").click()
        expect(modal).to_be_hidden()
        self.assertEqual(ClientDetail.objects.filter(user=self.user).count(), before)

    def test_the_client_limit_is_enforced(self):
        count = ClientDetail.objects.filter(user=self.user).count()
        self.user.Client_limit = count
        self.user.save()
        self.open()
        self.add_client("CLSEE5003L", "CLS OVER LIMIT", ALPHA)
        self.assertIn(f"You have reached the limit of {count} Client limits.", self.banner())
        self.assertFalse(ClientDetail.objects.filter(user=self.user, PANNo="CLSEE5003L").exists())
        self.user.Client_limit = count + 1
        self.user.save()
        self.open()
        self.add_client("CLSEE5003L", "CLS UNDER LIMIT", ALPHA)
        self.assertTrue(ClientDetail.objects.filter(user=self.user, PANNo="CLSEE5003L").exists())

    @unittest.expectedFailure
    def test_a_pan_with_the_wrong_shape_is_refused(self):
        """KNOWN BUG: only the length is checked; '12345ABCDE' (not 5 letters, 4 digits, 1 letter) is saved."""
        self.open()
        self.add_client("12345ABCDE", "CLS BAD PAN", ALPHA)
        self.assertFalse(ClientDetail.objects.filter(user=self.user, PANNo="12345ABCDE").exists())

    # ====================================================================== 4. edit a client
    def test_the_edit_page_is_filled_with_the_client(self):
        client = self.open_edit("a3")
        self.assertEqual(self.page.input_value("#PANNo"), client.PANNo)
        self.assertEqual(self.page.input_value("#Name"), client.Name)
        self.assertEqual(self.page.input_value("#Group"), ALPHA)
        self.assertEqual(self.page.input_value('input[name="ClientIdDpId"]'), client.ClientIdDpId)

    def test_editing_saves_every_field(self):
        client = self.open_edit("a2")
        self.page.fill("#PANNo", "clsff6001m")
        self.page.fill("#Name", "cls bhavna shah")
        self.page.select_option("#Group", BETA)
        self.page.fill('input[name="ClientIdDpId"]', "1203320000000777")
        self.submit_edit()
        self.assertRegex(self.page.url, r"/ClientSetup")
        client.refresh_from_db()
        self.assertEqual((client.PANNo, client.Name, client.Group.GroupName, client.ClientIdDpId),
                         ("CLSFF6001M", "CLS BHAVNA SHAH", BETA, "1203320000000777"))

    def test_a_new_pan_shows_on_the_orders_that_use_the_client(self):
        client = self.open_edit(USED)
        self.page.fill("#PANNo", "CLSFF6002N")
        self.submit_edit()
        self.used_detail.refresh_from_db()
        self.assertEqual(self.used_detail.OrderDetailPANNo_id, client.id)
        self.assertEqual(self.used_detail.OrderDetailPANNo.PANNo, "CLSFF6002N")

    def test_editing_to_another_clients_pan_is_refused(self):
        client = self.open_edit("a1")
        self.page.fill("#PANNo", CLIENTS["a2"][0].lower())
        self.page.fill("#Name", "CLS SHOULD NOT SAVE")
        self.submit_edit()
        self.assertIn(f"/EditClient/{client.id}", self.page.url)
        self.assertIn("Client Already Exist.", " ".join(self.page.locator(".alert").all_text_contents()))
        client.refresh_from_db()
        self.assertEqual((client.PANNo, client.Name), (CLIENTS["a1"][0], "CLS ANAND"))

    def test_keeping_the_same_pan_is_allowed(self):
        client = self.open_edit("a1")
        self.page.fill("#Name", "CLS ANAND KUMAR")
        self.submit_edit()
        client.refresh_from_db()
        self.assertEqual((client.PANNo, client.Name), (CLIENTS["a1"][0], "CLS ANAND KUMAR"))

    def test_the_edit_form_refuses_an_empty_pan_and_back_returns(self):
        client = self.open_edit("a1")
        self.page.fill("#PANNo", "")
        url = self.page.url
        self.page.get_by_role("button", name="Submit").click()
        self.page.wait_for_timeout(400)
        self.assertEqual(self.page.url, url)
        with self.page.expect_navigation():
            self.page.get_by_role("button", name="Back").click()
        self.assertRegex(self.page.url, r"/ClientSetup$")
        client.refresh_from_db()
        self.assertEqual(client.PANNo, CLIENTS["a1"][0])

    # ====================================================================== 5. delete
    def delete_row(self, key):
        client = self.c[key]
        row = self.page.locator(f'tr:has(.client-checkbox[value="{client.id}"])')
        row.get_by_role("button", name="Delete").click()
        modal = self.page.locator(f"#deleteModal_{client.id}")
        expect(modal).to_be_visible()
        expect(modal).to_contain_text(f"Are you sure you want to delete {client.PANNo} Client?")
        return modal

    def test_delete_asks_first_and_removes_an_unused_client(self):
        self.open(ALPHA)
        modal = self.delete_row("a2")
        modal.get_by_role("button", name="Cancel").click()
        expect(modal).to_be_hidden()
        self.assertTrue(ClientDetail.objects.filter(id=self.c["a2"].id).exists())
        modal = self.delete_row("a2")
        with self.page.expect_navigation():
            modal.get_by_role("button", name="Delete").click()
        settle(self.page)
        self.assertIn("Client deleted successfully.", self.banner())
        self.assertFalse(ClientDetail.objects.filter(id=self.c["a2"].id).exists())
        self.assertEqual(set(self.shown_ids()), self.ids("a1", "a3"))

    def test_a_client_with_orders_cannot_be_deleted(self):
        self.open(BETA)
        modal = self.delete_row(USED)
        with self.page.expect_navigation():
            modal.get_by_role("button", name="Delete").click()
        settle(self.page)
        client = self.c[USED]
        self.assertIn(f"Client-{client.PANNo} ({client.Name}) cannot be deleted as it has order(s) in: {CLS_IPO}", self.banner())
        self.assertTrue(ClientDetail.objects.filter(id=client.id).exists())

    def test_ticking_rows_turns_delete_all_into_delete_selected(self):
        self.open(ALPHA)
        button = self.page.locator("#deleteAllClientsBtn")
        expect(button).to_have_text("Delete All Clients")
        self.page.locator(f'.client-checkbox[value="{self.c["a1"].id}"]').check()
        expect(button).to_have_text("Delete Selected")
        self.page.locator("#select_all_clients").check()
        self.assertEqual(self.page.locator(".client-checkbox:checked").count(), 3)
        self.page.locator(f'.client-checkbox[value="{self.c["a1"].id}"]').uncheck()
        expect(self.page.locator("#select_all_clients")).not_to_be_checked()
        self.page.locator(f'.client-checkbox[value="{self.c["a1"].id}"]').check()
        expect(self.page.locator("#select_all_clients")).to_be_checked()
        self.page.locator("#select_all_clients").uncheck()
        expect(button).to_have_text("Delete All Clients")

    def test_delete_selected_removes_the_ticked_clients_and_keeps_used_ones(self):
        self.open(BETA)
        for key in ("b1", "b2"):
            self.page.locator(f'.client-checkbox[value="{self.c[key].id}"]').check()
        self.page.once("dialog", lambda d: d.dismiss())
        self.page.locator("#deleteAllClientsBtn").click()
        self.page.wait_for_timeout(400)
        self.assertEqual(self.mine().filter(id__in=self.ids("b1", "b2")).count(), 2, "cancel keeps them")
        self.page.once("dialog", lambda d: d.accept())
        with self.page.expect_navigation():
            self.page.locator("#deleteAllClientsBtn").click()
        settle(self.page)
        self.assertFalse(ClientDetail.objects.filter(id=self.c["b2"].id).exists())
        self.assertTrue(ClientDetail.objects.filter(id=self.c[USED].id).exists())
        banner = self.banner()
        self.assertIn("Successfully deleted 1 client(s).", banner)
        self.assertIn(f"1 client(s) could not be deleted (Used in Orders): {self.c[USED].PANNo}({CLS_IPO})", banner)

    def test_delete_all_clients_of_a_group_deletes_the_unused_ones_and_downloads_a_report(self):
        self.open(BETA)                                     # the delete follows the chosen group
        self.page.locator("#deleteAllClientsBtn").click()
        modal = self.page.locator("#DeleteAllClient")
        expect(modal).to_be_visible()
        expect(modal).to_contain_text("Are you sure you want to delete all clients?")
        modal.get_by_role("button", name="Cancel").click()
        expect(modal).to_be_hidden()
        self.assertEqual(self.mine().filter(Group=self.groups[BETA]).count(), 2)
        alpha_before = set(ClientDetail.objects.filter(user=self.user).exclude(Group=self.groups[BETA]).values_list("id", flat=True))

        report = self.page.request.get(f"{self.live_server_url}/DeleteAllClient?page=1")
        self.assertEqual(report.status, 200)
        self.assertIn("PAN_Deletion_Status.xlsx", report.headers["content-disposition"])
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "status.xlsx")
            with open(path, "wb") as handle:
                handle.write(report.body())
            sheet = list(openpyxl.load_workbook(path).active.iter_rows(values_only=True))
        self.assertEqual(sheet[0], ("PAN", "Status"))
        self.assertEqual(dict(sheet[1:]), {CLIENTS["b1"][0]: f"Not Deleted (Used in: {CLS_IPO})",
                                           CLIENTS["b2"][0]: "Deleted"})
        self.assertEqual(set(self.mine().filter(Group=self.groups[BETA]).values_list("id", flat=True)), self.ids(USED))
        self.assertEqual(set(ClientDetail.objects.filter(user=self.user).exclude(Group=self.groups[BETA]).values_list("id", flat=True)),
                         alpha_before, "clients of other groups are untouched")

    def test_the_delete_all_button_in_the_modal_reloads_the_page(self):
        self.open(BETA)
        self.page.locator("#deleteAllClientsBtn").click()
        modal = self.page.locator("#DeleteAllClient")
        expect(modal).to_be_visible()
        with self.page.expect_navigation(timeout=15000):
            modal.get_by_role("button", name="Delete").click()
        settle(self.page)
        self.assertRegex(self.page.url, r"/ClientSetup\?page=")
        self.assertEqual(set(self.shown_ids()), self.ids(USED))

    # ====================================================================== 6. table tools
    def test_clicking_a_header_sorts_the_rows(self):
        self.open(ALPHA)
        header = self.page.locator(".dataTables_scrollHead thead tr").first.locator("th", has_text="Name")
        header.click()
        names = [r["cells"][2] for r in self.rows()]
        self.assertEqual(names, sorted(names))
        header.click()
        names = [r["cells"][2] for r in self.rows()]
        self.assertEqual(names, sorted(names, reverse=True))

    def test_column_boxes_filter_the_rows_on_the_page(self):
        self.open(ALPHA)
        boxes = self.page.locator(".dataTables_scrollHead thead tr").nth(1).locator("input")
        boxes.nth(1).fill("ANAND")
        boxes.nth(1).dispatch_event("change")
        self.assertEqual(set(self.shown_ids()), self.ids("a1"))
        boxes.nth(1).fill("")
        boxes.nth(1).dispatch_event("change")
        self.assertEqual(set(self.shown_ids()), self.ids("a1", "a2", "a3"))

    def test_the_excel_download_lists_the_clients_shown(self):
        self.open(ALPHA)
        self.page.get_by_role("button", name="Download").click()
        self.page.locator(".dt-button-collection .dt-button", has_text="Excel").click()
        popup = self.page.locator(".swal2-popup")
        expect(popup).to_be_visible()
        with self.page.expect_download() as info:
            popup.locator(".swal2-confirm").click()
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, info.value.suggested_filename)
            info.value.save_as(path)
            clean = lambda v: "" if v is None else str(v).replace("\xa0", " ").strip()
            rows = [tuple(clean(v) for v in r) for r in openpyxl.load_workbook(path).active.iter_rows(values_only=True) if any(r)]
        head = next(i for i, r in enumerate(rows) if r and "PAN No." in r)
        self.assertEqual(list(rows[head]), ["PAN No.", "Name", "Group", "Client-ID/DP-ID"])
        got = sorted(rows[head + 1:])
        want = sorted((c.PANNo, c.Name, ALPHA, c.ClientIdDpId) for c in (self.c["a1"], self.c["a2"], self.c["a3"]))
        self.assertEqual(got, want)

    def test_the_columns_button_hides_a_column(self):
        self.open(ALPHA)
        self.page.get_by_role("button", name="Columns").click()
        self.page.locator(".dt-button-collection").get_by_text("Client-ID/DP-ID", exact=True).click()
        self.page.keyboard.press("Escape")
        expect(self.page.locator(".dataTables_scrollHead thead tr").first.locator("th", has_text="Client-ID/DP-ID")).to_have_count(0)

    # ====================================================================== 7. name look-up
    def test_the_name_lookup_returns_the_client_of_a_pan(self):
        self.open()
        token = next(c["value"] for c in self.page.context.cookies() if c["name"] == "csrftoken")
        post = lambda pan: self.page.request.post(f"{self.live_server_url}/autocomplete1", form={"PAN": pan},
                                                  headers={"X-CSRFToken": token, "Referer": f"{self.live_server_url}/ClientSetup"})
        self.assertEqual(post(CLIENTS["a3"][0]).json(), ["CLS CHIRAG"])
        self.assertEqual(post("CLSZZ9999Z").json(), [])
