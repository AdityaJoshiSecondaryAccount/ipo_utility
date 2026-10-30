"""
Group Setup page E2E (/GroupSetup, templates/GroupSetup.html; edit page /EditGroup/<id>, templates/EditGroup.html;
"KOSTAK GROUP" in the top bar).

Runs on its OWN groups (all named "GRS ...", plus the IPO "GRS IPO" for the orders that block a delete), reset in
setUp, so the groups of every other page (ALPHA TRADERS, GWD ..., ACCT ..., TRD ...) are never edited or deleted.

What the page shows: every group of the broker, sorted by name, 50 per page (server side; "Show" 50 / 100 / 200 / 250 /
All, kept in the session), columns Sr No. (with a select checkbox) | Group Name | Mobile No | Email | Address | Remark
("-" when empty) | Edit, Delete. A plain visit (no ?page=) resets the search and page size; the search box (debounced, or
Enter) and the header sort are POSTed to /GroupSetup and kept in the session.
  ADD KOSTAK GROUP modal  -> /AddGroup     name required (no "," or "/"), 10 digit mobile, email, address, remark
  BULK UPLOAD modal       -> /BulkUploadGroup  .csv / .xls / .xlsx with a GroupName column; sample at /DownloadGroupSample
  BULK DELETE (n)         -> /BulkDeleteGroup  (JSON); appears instead of BULK UPLOAD while rows are ticked
  Delete (per row modal)  -> /DeleteGroup/<id>  refused while the group has clients or orders; its ledger entries keep
                             the group's name (group_name) and lose the link
  Edit                    -> /EditGroup/<id> -> /UpdateGroup/<id>
Flash messages fade out after 3 seconds, so they are read from the HTML the server sent, not from the screen.
"""
import csv
import datetime
import html
import os
import re
import tempfile
import unittest

import openpyxl
from django.db.models import Q
from django.utils import timezone
from playwright.sync_api import expect

from home.models import Accounting, ClientDetail, CurrentIpoName, GroupDetail, Order, OrderDetail
from test_cases.base import PersistentLiveServerTestCase, pause_for_inspection, settle

IPO = "GRS IPO"
# key -> (name, mobile, email, address, remark)
GROUPS = {
    "alpha": ("GRS ALPHA", "9876543210", "alpha@grs.test", "1 Alpha Road", "first"),
    "beta": ("GRS BETA", "", "", "", ""),
    "gamma": ("GRS GAMMA", "9123456780", "gamma@grs.test", "Gamma Street", "zeta remark"),
    "orders": ("GRS ORDERS", "9000000001", "", "Order Lane", ""),
    "clients": ("GRS CLIENTS", "", "clients@grs.test", "", "has client"),
}
SAMPLE_HEADER = ["GroupName", "MobileNo", "Email", "Address", "Remark"]

ROWS_JS = """() => [...document.querySelectorAll('#example tbody tr')].filter(tr => tr.querySelector('.group-checkbox')).map(tr => ({
    id: Number(tr.querySelector('.group-checkbox').value),
    cells: [...tr.querySelectorAll(':scope > td')].map(c => c.textContent.replace(/\\s+/g, ' ').trim())}))"""

HEADER = ".dataTables_scrollHead thead tr:not(.filters) th"


def sort_key(g):
    return (g.GroupName.strip().lower(), g.id)


class GroupSetupPlaywrightTestCase(PersistentLiveServerTestCase):
    PAGE_NAME = "group_setup"
    _ran_any = False

    # ------------------------------------------------------------------ setup / teardown
    @classmethod
    def tearDownClass(cls):
        cls.log_summary()
        if cls._ran_any:
            try:
                cls.page.goto(f"{cls.live_server_url}/GroupSetup")
                settle(cls.page)
            except Exception:
                pass
            pause_for_inspection(cls.PAGE_NAME)
        super().tearDownClass()

    def setUp(self):
        super().setUp()
        type(self)._ran_any = True
        self.user.Group_limit = None
        self.user.save()
        self.ipo = CurrentIpoName.objects.get_or_create(
            user=self.user, IPOName=IPO,
            defaults=dict(IPOType="MAINBOARD", IPOPrice=100.0, PreOpenPrice=150.0, LotSizeRetail=100,
                          LotSizeSHNI=1400, LotSizeBHNI=7000, Active=True))[0]
        self.reset()
        self.js_errors = []
        self.page.on("pageerror", lambda e: self.js_errors.append(e.message))
        self.clear_storage()

    def tearDown(self):
        self.user.Group_limit = None
        self.user.save()
        self.clear_storage()
        super().tearDown()

    def clear_storage(self):
        try:
            if self.page.url.startswith(self.live_server_url):
                self.page.evaluate("localStorage.clear(); sessionStorage.clear()")
        except Exception:
            pass

    def mine(self):
        return GroupDetail.objects.filter(user=self.user, GroupName__contains="GRS")

    def reset(self):
        mine = self.mine()
        OrderDetail.objects.filter(user=self.user).filter(Q(Order__OrderGroup__in=mine) | Q(Order__OrderIPOName=self.ipo)).delete()
        Order.objects.filter(user=self.user).filter(Q(OrderGroup__in=mine) | Q(OrderIPOName=self.ipo)).delete()
        ClientDetail.objects.filter(user=self.user).filter(Q(Group__in=mine) | Q(PANNo__startswith="GRSP")).delete()
        Accounting.objects.filter(user=self.user).filter(Q(group__in=mine) | Q(group_name__contains="GRS") | Q(ipo=self.ipo)).delete()
        mine.delete()
        self.g = {}
        for key, (name, mobile, email, address, remark) in GROUPS.items():
            self.g[key] = GroupDetail.objects.create(user=self.user, GroupName=name, MobileNo=mobile, Email=email,
                                                     Address=address, Remark=remark)
        for rate in (500, 510):
            order = Order.objects.create(user=self.user, OrderIPOName=self.ipo, OrderGroup=self.g["orders"], OrderType="BUY",
                                         OrderCategory="Kostak", InvestorType="RETAIL", Quantity=2, Rate=rate,
                                         OrderDate=datetime.date(2026, 2, 1), OrderTime=datetime.time(10, 0))
            OrderDetail.objects.bulk_create([OrderDetail(user=self.user, Order=order, PreOpenPrice=150.0) for _ in range(2)])
        self.client_row = ClientDetail.objects.create(user=self.user, PANNo="GRSPA1234A", Name="GRS CLIENT", Group=self.g["clients"])
        self.ledger = Accounting.objects.create(user=self.user, ipo=self.ipo, group=self.g["beta"], amount=500, amount_type="credit",
                                                remark="GRS ledger", date_time=timezone.now())

    # ---------------------------------------------------------------------- page helpers
    def open(self, search=None):
        """A plain visit (resets search and page size), then optionally a search."""
        self.page.goto(f"{self.live_server_url}/GroupSetup")
        settle(self.page)
        expect(self.page.locator("h3", has_text="KOSTAK GROUP")).to_be_visible()
        self.page.wait_for_timeout(300)                         # the search box is wired 100 ms after load
        if search is not None:
            self.search(search)

    def act(self, action):
        """Run an action that loads a page; return the flash messages of the page the server sent."""
        with self.page.expect_navigation() as nav:
            action()
        response = nav.value
        body = response.text() if response else ""
        settle(self.page)
        self.page.wait_for_timeout(200)
        return [html.unescape(re.sub(r"\s+", " ", m)).strip() for m in
                re.findall(r'id="successalerts"[^>]*>.*?</button>(.*?)</div>', body, flags=re.S)]

    def search(self, text):
        box = self.page.locator(".dataTables_filter input")
        box.fill(text)
        self.act(lambda: box.press("Enter"))

    def rows(self):
        return self.page.evaluate(ROWS_JS)

    def names(self):
        return [r["cells"][1] for r in self.rows()]

    def expected_cells(self, g, sr):
        return [str(sr), g.GroupName.strip(), g.MobileNo, g.Email, g.Address, g.Remark or "-", "Edit Delete"]

    def open_add(self):
        self.page.get_by_role("button", name="ADD KOSTAK GROUP").click()
        modal = self.page.locator("#ADDGROUP")
        expect(modal).to_be_visible()
        return modal

    def fill_add(self, name, mobile="", email="", address="", remark=""):
        modal = self.open_add()
        modal.locator("#GroupName").fill(name)
        if mobile:
            modal.locator("#MobileNo").fill(mobile)
        modal.locator('input[name="Email"]').fill(email)
        modal.locator('input[name="Address"]').fill(address)
        modal.locator('input[name="Remark"]').fill(remark)
        return modal

    def submit_add(self, modal):
        return self.act(lambda: modal.get_by_role("button", name="Submit Form").click())

    def open_edit(self, key):
        g = self.g[key]
        self.open(search="GRS")
        with self.page.expect_navigation():
            self.page.locator(f'tr:has(.group-checkbox[value="{g.id}"])').get_by_role("button", name="Edit").click()
        settle(self.page)
        self.assertRegex(self.page.url, rf"/EditGroup/{g.id}\?page=")
        return g

    def submit_edit(self):
        return self.act(lambda: self.page.get_by_role("button", name="Submit").click())

    def strip_client_checks(self, *selectors):
        """Turn off the browser's checks on these inputs so the SERVER's validation is what gets tested."""
        for sel in selectors:
            self.page.evaluate("""s => { const e = document.querySelector(s); ['pattern','minlength','maxlength','oninput','required']
                .forEach(a => e.removeAttribute(a)); e.type = 'text'; e.oninput = null; e.setCustomValidity(''); }""", sel)

    # ====================================================================== 1. the list
    def test_lists_the_groups_with_exact_columns_sorted_by_name(self):
        self.open(search="GRS")
        headers = self.page.locator(HEADER).all_inner_texts()
        self.assertEqual([h.strip() for h in headers], ["Sr No.", "Group Name", "Mobile No", "Email", "Address", "Remark", "Action"])
        want = sorted(self.mine(), key=sort_key)
        rows = self.rows()
        self.assertEqual([r["id"] for r in rows], [g.id for g in want])
        for sr, (row, g) in enumerate(zip(rows, want), start=1):
            self.assertEqual(row["cells"], self.expected_cells(g, sr), g.GroupName)
        expect(self.page.get_by_text(f"Showing 1 to {len(want)} of {len(want)} Rows")).to_be_visible()
        self.assertEqual(self.js_errors, [])

    def test_a_plain_visit_lists_every_group_of_the_broker(self):
        self.open()
        everything = sorted(GroupDetail.objects.filter(user=self.user), key=sort_key)
        rows = self.rows()
        self.assertEqual([r["id"] for r in rows], [g.id for g in everything[:50]])
        expect(self.page.get_by_text(f"of {len(everything)} Rows")).to_be_visible()
        for name in ("ALPHA TRADERS", "BETA INVESTORS"):
            self.assertTrue(GroupDetail.objects.filter(user=self.user, GroupName=name).exists(), f"{name} must survive")

    # ====================================================================== 2. search
    def test_search_looks_at_name_mobile_email_address_and_remark(self):
        cases = [("grs b", ["beta"]), ("alpha@grs", ["alpha"]), ("9123", ["gamma"]), ("Order Lane", ["orders"]),
                 ("zeta", ["gamma"]), ("zz-nothing", [])]
        for term, keys in cases:
            with self.subTest(search=term):
                self.open(search=term)
                # other pages' groups may match too (e.g. Status's "ZETA BUYERS"): judge only this page's groups
                mine = {g.id for g in self.g.values()}
                self.assertEqual({r["id"] for r in self.rows()} & mine, {self.g[k].id for k in keys})
                self.assertEqual(self.page.input_value(".dataTables_filter input"), term)

    def test_the_search_is_kept_by_the_page_links_and_cleared_by_a_plain_visit(self):
        self.open(search="GRS")
        self.page.goto(f"{self.live_server_url}/GroupSetup?page=1")
        settle(self.page)
        self.assertEqual(set(self.names()), {g[0] for g in GROUPS.values()})
        self.open()
        self.assertEqual(self.page.input_value(".dataTables_filter input"), "")
        self.assertGreater(len(self.rows()), len(GROUPS) - 1)

    def test_typing_searches_by_itself_after_a_short_pause(self):
        self.open()
        box = self.page.locator(".dataTables_filter input")
        box.click()
        with self.page.expect_navigation(timeout=5000):
            box.press_sequentially("grs gam", delay=40)
        settle(self.page)
        self.assertEqual(self.names(), ["GRS GAMMA"])

    # ====================================================================== 3. sort
    def sort_by(self, title):
        self.act(lambda: self.page.locator(HEADER, has_text=re.compile(rf"^\s*{title}\s*$")).first.click())

    def test_clicking_a_header_sorts_on_the_server_and_keeps_the_search(self):
        keys = {"Mobile No": lambda g: (g.MobileNo, g.id), "Email": lambda g: (g.Email.strip().lower(), g.id),
                "Address": lambda g: (g.Address, g.id), "Remark": lambda g: (g.Remark, g.id)}
        for title, key in keys.items():
            with self.subTest(column=title):
                self.open(search="GRS")
                self.sort_by(title)
                self.assertEqual([r["id"] for r in self.rows()], [g.id for g in sorted(self.mine(), key=key)])
                self.assertEqual(self.page.input_value(".dataTables_filter input"), "GRS")
        self.open(search="GRS")
        self.sort_by("Group Name")
        self.sort_by("Group Name")
        self.assertEqual(self.names(), sorted((g[0] for g in GROUPS.values()), reverse=True))

    # ====================================================================== 4. add
    def test_adding_a_group_saves_every_field(self):
        self.open()
        modal = self.fill_add("grs delta", "9012345678", "delta@grs.test", "Delta Road", "added in test")
        messages = self.submit_add(modal)
        self.assertIn("Group Added successfully.", messages)
        g = GroupDetail.objects.get(user=self.user, GroupName="GRS DELTA")
        self.assertEqual((g.MobileNo, g.Email, g.Address, g.Remark), ("9012345678", "delta@grs.test", "Delta Road", "added in test"))
        self.search("GRS DELTA")
        self.assertEqual(self.rows()[0]["cells"], self.expected_cells(g, 1))

    def test_a_new_name_is_trimmed_and_upper_cased(self):
        self.open()
        self.submit_add(self.fill_add("   grs epsilon  "))
        self.assertTrue(GroupDetail.objects.filter(user=self.user, GroupName="GRS EPSILON").exists())

    def test_an_existing_name_in_any_case_is_refused(self):
        before = self.mine().count()
        self.open()
        messages = self.submit_add(self.fill_add("grs alpha"))
        self.assertIn("Group Already Exist.", messages)
        self.assertEqual(self.mine().count(), before)

    @unittest.expectedFailure
    def test_an_existing_name_with_spaces_around_it_is_refused(self):
        """KNOWN BUG: the duplicate check upper-cases but does not trim, so ' grs alpha ' becomes a second GRS ALPHA."""
        self.open()
        self.submit_add(self.fill_add("  grs alpha "))
        self.assertEqual(GroupDetail.objects.filter(user=self.user, GroupName="GRS ALPHA").count(), 1)

    def test_the_add_form_checks_name_mobile_and_email_in_the_browser(self):
        self.open()
        modal = self.open_add()
        name, mobile, email = modal.locator("#GroupName"), modal.locator("#MobileNo"), modal.locator('input[name="Email"]')
        submit = modal.get_by_role("button", name="Submit Form")
        url = self.page.url
        submit.click()
        self.assertTrue(name.evaluate("e => e.validity.valueMissing"), "a name is required")
        for bad in ("GRS A,B", "GRS A/B"):
            name.fill(bad)
            self.assertFalse(name.evaluate("e => e.checkValidity()"), bad)
            expect(modal.locator("#nameError")).to_contain_text("Commas (,) and slashes (/) are not allowed")
        name.fill("GRS OK")
        self.assertTrue(name.evaluate("e => e.checkValidity()"))
        mobile.press_sequentially("98ab76")
        self.assertEqual(mobile.input_value(), "9876", "only digits can be typed")
        self.assertFalse(mobile.evaluate("e => e.checkValidity()"), "fewer than 10 digits")
        mobile.press_sequentially("543210")
        self.assertTrue(mobile.evaluate("e => e.checkValidity()"))
        email.fill("not-an-email")
        self.assertFalse(email.evaluate("e => e.checkValidity()"))
        email.fill("")
        mobile.fill("123")
        submit.click()
        self.page.wait_for_timeout(400)
        self.assertEqual(self.page.url, url, "an invalid form is not sent")
        self.assertFalse(GroupDetail.objects.filter(user=self.user, GroupName="GRS OK").exists())

    @unittest.expectedFailure
    def test_typing_a_mobile_number_raises_no_script_error(self):
        """KNOWN BUG: validateMobileNo() writes to #mobileNoError, which does not exist -> TypeError on every key."""
        self.open()
        self.open_add().locator("#MobileNo").press_sequentially("98765")
        self.page.wait_for_timeout(200)
        self.assertEqual(self.js_errors, [])

    def test_the_server_refuses_a_bad_email_or_mobile_too(self):
        for field, value, message in (("Email", "bad@", "Invalid email format."),
                                      ("MobileNo", "12345", "Invalid Mobile Number. Please enter exactly 10 digits.")):
            with self.subTest(field=field):
                self.open()
                modal = self.open_add()
                self.strip_client_checks(f'#ADDGROUP input[name="{field}"]')
                modal.locator("#GroupName").fill("GRS SERVER CHECK")
                modal.locator(f'input[name="{field}"]').fill(value)
                self.assertIn(message, self.submit_add(modal))
                self.assertFalse(GroupDetail.objects.filter(user=self.user, GroupName="GRS SERVER CHECK").exists())

    def test_the_group_limit_stops_a_new_group(self):
        limit = GroupDetail.objects.filter(user=self.user).count()
        self.user.Group_limit = limit
        self.user.save()
        self.open()
        messages = self.submit_add(self.fill_add("GRS OVER LIMIT"))
        self.assertIn(f"You have reached the limit of {limit} Group Limits.", messages)
        self.assertFalse(GroupDetail.objects.filter(user=self.user, GroupName="GRS OVER LIMIT").exists())

    def test_close_leaves_the_add_modal_without_saving(self):
        self.open()
        modal = self.fill_add("GRS CLOSED")
        modal.locator(".modal-footer").get_by_role("button", name="Close").click()
        expect(modal).to_be_hidden()
        self.assertFalse(GroupDetail.objects.filter(user=self.user, GroupName="GRS CLOSED").exists())

    # ====================================================================== 5. edit
    def test_the_edit_page_is_filled_in_and_saves_every_field(self):
        g = self.open_edit("gamma")
        values = [self.page.input_value(f'input[name="{n}"]') for n in ("GroupName", "MobileNo", "Email", "Address", "Remark")]
        self.assertEqual(values, list(GROUPS["gamma"]))
        for field, value in (("GroupName", "grs gamma two"), ("MobileNo", "9988776655"), ("Email", "g2@grs.test"),
                             ("Address", "New Street"), ("Remark", "")):
            self.page.fill(f'input[name="{field}"]', value)
        self.submit_edit()
        self.assertRegex(self.page.url, r"/GroupSetup\?page=")
        g.refresh_from_db()
        self.assertEqual((g.GroupName, g.MobileNo, g.Email, g.Address, g.Remark),
                         ("GRS GAMMA TWO", "9988776655", "g2@grs.test", "New Street", ""))

    def test_renaming_keeps_the_groups_orders_clients_and_ledger(self):
        for key, new in (("orders", "GRS ORDERS RENAMED"), ("clients", "GRS CLIENTS RENAMED"), ("beta", "GRS BETA RENAMED")):
            with self.subTest(group=key):
                g = self.open_edit(key)
                self.page.fill('input[name="GroupName"]', new.lower())
                self.submit_edit()
                g.refresh_from_db()
                self.assertEqual(g.GroupName, new)
        self.assertEqual(Order.objects.filter(user=self.user, OrderGroup__GroupName="GRS ORDERS RENAMED").count(), 2)
        self.client_row.refresh_from_db()
        self.assertEqual(self.client_row.Group.GroupName, "GRS CLIENTS RENAMED")
        self.ledger.refresh_from_db()
        self.assertEqual(self.ledger.group.GroupName, "GRS BETA RENAMED")

    def test_renaming_to_another_groups_name_is_refused(self):
        g = self.open_edit("beta")
        self.page.fill('input[name="GroupName"]', "grs alpha")
        messages = self.submit_edit()
        self.assertIn("Group Already Exist.", messages)
        self.assertRegex(self.page.url, rf"/EditGroup/{g.id}")
        g.refresh_from_db()
        self.assertEqual(g.GroupName, "GRS BETA")
        # keeping its own name is fine
        self.page.fill('input[name="Remark"]', "own name kept")
        self.submit_edit()
        g.refresh_from_db()
        self.assertEqual((g.GroupName, g.Remark), ("GRS BETA", "own name kept"))

    @unittest.expectedFailure
    def test_a_renamed_name_is_trimmed(self):
        """KNOWN BUG: UpdateGroup checks the trimmed name but saves the raw one, so '  grs beta two ' keeps its spaces."""
        g = self.open_edit("beta")
        self.page.fill('input[name="GroupName"]', "  grs beta two ")
        self.submit_edit()
        g.refresh_from_db()
        self.assertEqual(g.GroupName, "GRS BETA TWO")

    @unittest.expectedFailure
    def test_the_edit_page_refuses_a_slash_like_the_add_form(self):
        """KNOWN BUG: the add form refuses '/' in a name (it breaks /BUY/<ipo>/<group> links) but the edit page saves it."""
        g = self.open_edit("beta")
        self.page.fill('input[name="GroupName"]', "GRS A/B")
        try:
            self.submit_edit()
        except Exception:
            pass                                                # refused in the browser: no page load
        g.refresh_from_db()
        self.assertEqual(g.GroupName, "GRS BETA")

    def test_the_edit_page_refuses_a_comma_bad_email_and_bad_mobile(self):
        g = self.open_edit("beta")
        name = self.page.locator('input[name="GroupName"]')
        name.fill("GRS A,B")
        self.assertFalse(name.evaluate("e => e.checkValidity()"))
        url = self.page.url
        self.page.get_by_role("button", name="Submit").click()
        self.page.wait_for_timeout(400)
        self.assertEqual(self.page.url, url)
        name.fill("GRS BETA")
        self.strip_client_checks('input[name="Email"]')
        self.page.fill('input[name="Email"]', "bad@")
        self.assertIn("Invalid email format.", self.submit_edit())
        self.assertRegex(self.page.url, rf"/EditGroup/{g.id}")
        self.strip_client_checks('input[name="MobileNo"]')
        self.page.fill('input[name="MobileNo"]', "abcdefghij")
        self.assertIn("Invalid Mobile Number. Please enter exactly 10 digits.", self.submit_edit())
        g.refresh_from_db()
        self.assertEqual((g.GroupName, g.Email, g.MobileNo), ("GRS BETA", "", ""))

    def test_back_leaves_the_edit_page_without_saving(self):
        g = self.open_edit("alpha")
        self.page.fill('input[name="Address"]', "not saved")
        with self.page.expect_navigation():
            self.page.get_by_role("button", name="Back").click()
        self.assertRegex(self.page.url, r"/GroupSetup$")
        g.refresh_from_db()
        self.assertEqual(g.Address, "1 Alpha Road")

    # ====================================================================== 6. delete
    def open_delete(self, key):
        g = self.g[key]
        self.open(search="GRS")
        self.page.locator(f'tr:has(.group-checkbox[value="{g.id}"])').get_by_role("button", name="Delete").click()
        modal = self.page.locator(f'div.modal[id="{g.id}"]')
        expect(modal).to_be_visible()
        expect(modal).to_contain_text(f"Are you sure you want to delete {g.GroupName} group?")
        return g, modal

    def test_deleting_a_group_keeps_its_ledger_entries_by_name(self):
        g, modal = self.open_delete("beta")
        self.act(lambda: modal.locator("button.btn-outline-danger", has_text="Delete").click())
        self.assertFalse(GroupDetail.objects.filter(id=g.id).exists())
        self.ledger.refresh_from_db()
        self.assertEqual((self.ledger.group_id, self.ledger.group_name, self.ledger.status), (None, "GRS BETA", True))
        self.assertNotIn("GRS BETA", self.names())

    def test_cancel_in_the_delete_modal_keeps_the_group(self):
        g, modal = self.open_delete("alpha")
        modal.get_by_role("button", name="Cancel").click()
        expect(modal).to_be_hidden()
        self.assertTrue(GroupDetail.objects.filter(id=g.id).exists())

    def test_a_group_with_clients_or_orders_cannot_be_deleted(self):
        for key, reason in (("clients", "has client(s). First Remove Client(s)."),
                            ("orders", "has created order(s). First Remove Order(s).")):
            with self.subTest(group=key):
                g, modal = self.open_delete(key)
                messages = self.act(lambda: modal.locator("button.btn-outline-danger", has_text="Delete").click())
                self.assertIn(f"Group {g.GroupName} cannot be deleted as it {reason}", messages)
                self.assertTrue(GroupDetail.objects.filter(id=g.id).exists())

    def tick(self, *keys):
        for key in keys:
            self.page.locator(f'.group-checkbox[value="{self.g[key].id}"]').check()

    def test_ticking_rows_swaps_bulk_upload_for_bulk_delete(self):
        self.open(search="GRS")
        upload, delete = self.page.locator("#bulk_upload_btn"), self.page.locator("#bulk_delete_btn")
        expect(upload).to_be_visible()
        expect(delete).to_be_hidden()
        self.tick("alpha")
        expect(delete).to_have_text("BULK DELETE (1)")
        expect(upload).to_be_hidden()
        expect(self.page.locator("#bulk_delete_text")).to_have_text("Are you sure you want to delete 1 selected group?")
        self.page.locator("#select_all_groups").check()
        expect(delete).to_have_text(f"BULK DELETE ({len(GROUPS)})")
        self.page.locator(f'.group-checkbox[value="{self.g["beta"].id}"]').uncheck()
        expect(self.page.locator("#select_all_groups")).not_to_be_checked()
        self.page.locator("#select_all_groups").check()
        self.page.locator("#select_all_groups").uncheck()
        expect(delete).to_be_hidden()
        expect(upload).to_be_visible()

    def test_bulk_delete_removes_the_free_groups_and_names_the_blocked_ones(self):
        self.open(search="GRS")
        self.tick("alpha", "gamma", "orders", "clients")
        self.page.locator("#bulk_delete_btn").click()
        modal = self.page.locator("#BULKDELETEGROUP")
        expect(modal).to_be_visible()
        expect(modal.locator("#bulk_delete_text")).to_have_text("Are you sure you want to delete 4 selected groups?")
        modal.get_by_role("button", name="Cancel").click()
        expect(modal).to_be_hidden()
        self.assertEqual(self.mine().count(), len(GROUPS))
        self.page.locator("#bulk_delete_btn").click()
        messages = self.act(lambda: modal.locator("#confirm_bulk_delete").click())
        text = " ".join(messages)
        self.assertIn("Successfully deleted 2 groups. 2 groups could not be deleted:", text)
        self.assertIn("GRS ORDERS (has orders)", text)
        self.assertIn("GRS CLIENTS (has clients)", text)
        self.assertEqual(set(self.mine().values_list("GroupName", flat=True)), {"GRS BETA", "GRS ORDERS", "GRS CLIENTS"})

    def test_bulk_delete_of_only_blocked_groups_reports_an_error(self):
        self.open(search="GRS")
        self.tick("orders")
        self.page.locator("#bulk_delete_btn").click()
        messages = self.act(lambda: self.page.locator("#confirm_bulk_delete").click())
        self.assertIn("Successfully deleted 0 groups. 1 group could not be deleted: GRS ORDERS (has orders).", messages)
        self.assertEqual(self.mine().count(), len(GROUPS))

    # ====================================================================== 7. bulk upload
    def upload(self, path):
        self.page.locator("#bulk_upload_btn").click()
        modal = self.page.locator("#BULKUPLOADGROUP")
        expect(modal).to_be_visible()
        modal.locator("#bulkFile").set_input_files(path)
        return self.act(lambda: modal.get_by_role("button", name="Upload", exact=True).click())

    def test_the_upload_modal_offers_the_sample_and_needs_a_file(self):
        self.open()
        self.page.locator("#bulk_upload_btn").click()
        modal = self.page.locator("#BULKUPLOADGROUP")
        expect(modal.get_by_role("link", name="Download Sample CSV")).to_have_attribute("href", "DownloadGroupSample")
        sample = self.page.request.get(f"{self.live_server_url}/DownloadGroupSample")
        self.assertEqual(sample.status, 200)
        self.assertIn("Sample_Group_Upload.csv", sample.headers["content-disposition"])
        self.assertEqual(sample.text().strip().split(","), SAMPLE_HEADER)
        url = self.page.url
        modal.get_by_role("button", name="Upload", exact=True).click()
        self.page.wait_for_timeout(400)
        self.assertEqual(self.page.url, url)
        self.assertTrue(modal.locator("#bulkFile").evaluate("e => e.validity.valueMissing"))

    def test_a_csv_adds_new_groups_and_counts_duplicates_and_bad_rows(self):
        rows = [["grs up one", "9811111111", "one@grs.test", "Up Road", "csv"],
                ["  GRS UP TWO ", "12345", "not-an-email", "", ""],          # bad mobile / email are dropped, the group is kept
                ["grs alpha", "", "", "", ""],                               # duplicate
                ["", "9822222222", "", "", "no name"]]                       # no name
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "groups.csv")
            with open(path, "w", newline="") as handle:
                csv.writer(handle).writerows([SAMPLE_HEADER] + rows)
            self.open()
            messages = self.upload(path)
        self.assertIn("Successfully added 2 groups.", messages)
        self.assertIn("Skipped 1 duplicate groups.", messages)
        self.assertIn("Failed to add 1 invalid entries.", messages)
        one = GroupDetail.objects.get(user=self.user, GroupName="GRS UP ONE")
        two = GroupDetail.objects.get(user=self.user, GroupName="GRS UP TWO")
        self.assertEqual((one.MobileNo, one.Email, one.Address, one.Remark), ("9811111111", "one@grs.test", "Up Road", "csv"))
        self.assertEqual((two.MobileNo, two.Email), ("", ""))
        self.assertEqual(GroupDetail.objects.filter(user=self.user, GroupName="GRS ALPHA").count(), 1)

    def test_an_excel_file_is_accepted_too(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "groups.xlsx")
            book = openpyxl.Workbook()
            for row in ([SAMPLE_HEADER, ["grs xl one", 9833333333, "xl@grs.test", "XL Road", "excel"]]):
                book.active.append(row)
            book.save(path)
            self.open()
            messages = self.upload(path)
        self.assertIn("Successfully added 1 groups.", messages)
        g = GroupDetail.objects.get(user=self.user, GroupName="GRS XL ONE")
        self.assertEqual((g.MobileNo, g.Email), ("9833333333", "xl@grs.test"))

    def test_a_wrong_file_type_or_layout_is_reported(self):
        with tempfile.TemporaryDirectory() as tmp:
            text = os.path.join(tmp, "groups.txt")
            with open(text, "w") as handle:
                handle.write("GroupName\nGRS TXT\n")
            self.open()
            self.assertIn("Invalid file format. Please upload a .csv, .xls, or .xlsx file.", self.upload(text))
            wrong = os.path.join(tmp, "groups.csv")
            with open(wrong, "w") as handle:
                handle.write("Name,MobileNo\nGRS WRONG,9844444444\n")
            self.open()
            self.assertIn('Invalid file structure. Column "GroupName" is missing.', self.upload(wrong))
        self.assertFalse(GroupDetail.objects.filter(user=self.user, GroupName__in=["GRS TXT", "GRS WRONG"]).exists())

    def test_the_group_limit_stops_a_bulk_upload_part_way(self):
        limit = GroupDetail.objects.filter(user=self.user).count() + 1
        self.user.Group_limit = limit
        self.user.save()
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "groups.csv")
            with open(path, "w", newline="") as handle:
                csv.writer(handle).writerows([SAMPLE_HEADER, ["GRS LIM ONE"], ["GRS LIM TWO"]])
            self.open()
            messages = self.upload(path)
        self.assertIn("Successfully added 1 groups.", messages)
        self.assertIn(f"You have reached your limit of {limit} Group Limits. Process stopped.", messages)
        self.assertEqual(sorted(self.mine().filter(GroupName__startswith="GRS LIM").values_list("GroupName", flat=True)), ["GRS LIM ONE"])

    # ====================================================================== 8. pages and page size
    def fillers(self, count=55):
        GroupDetail.objects.bulk_create([GroupDetail(user=self.user, GroupName=f"GRS FILL {i:02d}") for i in range(1, count + 1)])

    def test_pages_of_fifty_and_the_page_size(self):
        self.fillers(55)
        self.open(search="GRS FILL")
        rows = self.rows()
        self.assertEqual([r["cells"][0] for r in rows], [str(i) for i in range(1, 51)])
        self.assertEqual(rows[0]["cells"][1], "GRS FILL 01")
        expect(self.page.get_by_text("Showing 1 to 50 of 55 Rows")).to_be_visible()
        with self.page.expect_navigation():
            self.page.locator(".pagination .page-link", has_text=re.compile(r"^2$")).click()
        settle(self.page)
        self.assertEqual([r["cells"][1] for r in self.rows()], [f"GRS FILL {i}" for i in range(51, 56)])
        self.assertEqual(self.rows()[0]["cells"][0], "51")
        self.page.wait_for_timeout(300)
        self.act(lambda: self.page.select_option("#Gp_page_size", "All"))
        self.assertEqual(len(self.rows()), 55)
        expect(self.page.get_by_text("Showing 1 to 55 of 55 Rows")).to_be_visible()
        self.open()                                                 # a plain visit goes back to 50
        self.assertEqual(self.page.input_value("#Gp_page_size"), "50")

    # ====================================================================== 9. column boxes, download
    def type_in_column_box(self, placeholder, text):
        box = self.page.locator(f'.dataTables_scrollHead .filters input[placeholder^="{placeholder}"]')
        box.press_sequentially(text)
        self.page.wait_for_timeout(400)

    def test_column_boxes_filter_the_rows_once_a_sort_is_chosen(self):
        self.open(search="GRS")
        self.sort_by("Group Name")
        url = self.page.url
        self.type_in_column_box("Mobile No", "9123")
        self.assertEqual(self.page.url, url)
        self.assertEqual(self.names(), ["GRS GAMMA"])

    @unittest.expectedFailure
    def test_column_boxes_work_on_a_freshly_opened_page(self):
        """KNOWN BUG: the filter's redraw fires DataTables' 'order' event; with no sort chosen yet the page submits the
        sort form (column 0 != ''), so the page reloads and what was typed is lost."""
        self.open(search="GRS")
        try:
            self.type_in_column_box("Mobile No", "9123")
        except Exception:
            pass
        settle(self.page)
        self.assertEqual(self.names(), ["GRS GAMMA"])

    def test_the_excel_download_has_the_listed_groups(self):
        self.open(search="GRS")
        self.page.get_by_role("button", name="Download").click()
        self.page.locator(".dt-button-collection .dt-button", has_text="Excel").click()
        popup = self.page.locator(".swal2-popup")
        expect(popup).to_be_visible()
        expect(popup.locator(".swal2-confirm")).to_contain_text(f"Download {len(GROUPS)} Rows")
        with self.page.expect_download() as info:
            popup.locator(".swal2-confirm").click()
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, info.value.suggested_filename)
            info.value.save_as(path)
            self.assertTrue(path.endswith(".xlsx"), path)
            rows = list(openpyxl.load_workbook(path).active.iter_rows(values_only=True))
        clean = lambda v: "" if v is None else str(v).replace("\xa0", " ").strip()
        rows = [tuple(clean(c) for c in r) for r in rows]
        head = next(i for i, r in enumerate(rows) if "Group Name" in r)
        self.assertEqual(list(rows[head]), ["Group Name", "Mobile No", "Email", "Address"])
        got = sorted(r for r in rows[head + 1:] if any(r))
        want = sorted((g[0], g[1], g[2], g[3]) for g in GROUPS.values())
        self.assertEqual(got, want)
