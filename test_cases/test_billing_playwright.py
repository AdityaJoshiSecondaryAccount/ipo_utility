"""
Client Wise Billing E2E (/<IPOid>/Billing and /<IPOid>/Billing/<group>/<category>/<investor>[/<rate>]).

Runs on its OWN IPOs ("BILLING TEST IPO" - MAINBOARD, "BILLING SME IPO" - SME) so the global
Pre-Open edit can never disturb the data the other pages use. Orders are seeded with known
allotments; the app's own calculate() fills in the amounts and every amount is then checked
against an independent re-implementation of the settlement formulas written in this file.

Formulas (IPO price P, pre-open price O, rate R, allotted quantity Q, order quantity N):
  Kostak    BUY  ((O - P) * Q) - R                     SELL  -(BUY)       Q None -> 0   (Q = 0 -> -R / +R)
  Subject To BUY ((O - P) * Q) - R                     SELL  -(BUY)       Q None or 0 -> 0
             Premium method: the cost is R * Q instead of R
  Premium   BUY  (O - (P + R)) * N                     SELL  -(BUY)
  CALL  BUY  -R*N + max((O - (P+K))*N, 0)    SELL  +R*N + min((P+K-O)*N, 0)        (K = strike)
  PUT   BUY  -R*N + max((P+K-O)*N, 0)        SELL  +R*N + min((O-(P+K))*N, 0)
Kostak / Subject To amounts live on each OrderDetail (application); Premium / CALL / PUT on the Order.
"""
import csv
import datetime
import io
import os
import re
import shutil
import subprocess
import tempfile
import unittest
from collections import Counter
from urllib.parse import quote

import openpyxl
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.test import Client
from playwright.sync_api import expect

from home.models import ClientDetail, CurrentIpoName, GroupDetail, Order, OrderDetail
from home.views import calculate
from test_cases import oracle
from test_cases.base import PASSWORD, PersistentLiveServerTestCase, pause_for_inspection, settle

BILL_IPO, SME_IPO = "BILLING TEST IPO", "BILLING SME IPO"
IPO_PRICE, PRE_OPEN = 100.0, 250.0
ALPHA, BETA, SIGMA, PAGER = "ALPHA TRADERS", "BETA INVESTORS", "SIGMA CAPITAL", "BILL PAGER"
SEED_DATE, SEED_TIME = datetime.date(2026, 1, 20), datetime.time(10, 0, 0)
CUSTOMER_USER = "playwright_customer"

# Seed orders. Kostak / Subject To: one OrderDetail per entry of `alloted` (None = not allotted yet).
# Premium / CALL / PUT: a single order of `qty` shares (method = strike for options).
SPECS = [
    dict(group=ALPHA, type="BUY", cat="Kostak", inv="RETAIL", rate=450, alloted=[10, 0, None],
         pans=["ABCDE1234F", None, None]),
    dict(group=ALPHA, type="SELL", cat="Kostak", inv="RETAIL", rate=455, alloted=[10, None]),
    dict(group=BETA, type="BUY", cat="Kostak", inv="BHNI", rate=700, alloted=[1, 2], pans=["FGHIJ5678K", None]),
    dict(group=SIGMA, type="BUY", cat="Kostak", inv="SHNI", rate=710, alloted=[5]),
    dict(group=ALPHA, type="BUY", cat="Subject To", inv="RETAIL", rate=1200, alloted=[10, 0]),
    dict(group=BETA, type="BUY", cat="Subject To", inv="SHNI", rate=5, method="Premium", alloted=[10, 20]),
    dict(group=BETA, type="SELL", cat="Subject To", inv="SHNI", rate=1210, alloted=[10]),
    dict(group=BETA, type="SELL", cat="Subject To", inv="BHNI", rate=6, method="Premium", alloted=[10]),
    dict(group=ALPHA, type="BUY", cat="Premium", inv="PREMIUM", rate=20, qty=100),
    dict(group=SIGMA, type="SELL", cat="Premium", inv="PREMIUM", rate=30, qty=50),
] + [
    dict(group=ALPHA, type=t, cat=c, inv="OPTIONS", rate=10, qty=100, method=str(k))
    for c in ("CALL", "PUT") for t in ("BUY", "SELL") for k in (120, 200)
] + [
    dict(group=PAGER, type="BUY", cat="Kostak", inv="RETAIL", rate=111, alloted=[1] * 120),   # > 2 pages of 50
]
SME_SPECS = [
    dict(group=ALPHA, type="BUY", cat="Kostak", inv="RETAIL", rate=300, alloted=[5, 6]),
    dict(group=BETA, type="SELL", cat="Kostak", inv="RETAIL", rate=310, alloted=[2]),
    dict(group=ALPHA, type="BUY", cat="Premium", inv="PREMIUM", rate=10, qty=20),
]


class BillingPlaywrightTestCase(PersistentLiveServerTestCase):
    PAGE_NAME = "billing"
    _ran_any = False

    # ------------------------------------------------------------------ setup / teardown
    @classmethod
    def tearDownClass(cls):
        cls.log_summary()
        if cls._ran_any:
            try:
                cls.page.goto(f"{cls.live_server_url}/{cls._ipo_id}/Billing")
                settle(cls.page)
            except Exception:
                pass
            pause_for_inspection(cls.PAGE_NAME)
        super().tearDownClass()

    def setUp(self):
        super().setUp()
        type(self)._ran_any = True
        self.js_errors = []
        self.page.on("pageerror", lambda e: self.js_errors.append(f"{e.message}\n{e.stack}"))
        self.page.on("console", self._on_console)
        defaults = dict(IPOPrice=IPO_PRICE, PreOpenPrice=PRE_OPEN, LotSizeRetail=100, LotSizeSHNI=1400,
                        LotSizeBHNI=7000, Active=True)
        self.bipo = CurrentIpoName.objects.get_or_create(
            user=self.user, IPOName=BILL_IPO, defaults=dict(defaults, IPOType="MAINBOARD"))[0]
        self.sme = CurrentIpoName.objects.get_or_create(
            user=self.user, IPOName=SME_IPO, defaults=dict(defaults, IPOType="SME"))[0]
        type(self)._ipo_id = self.bipo.id
        self.seed(self.bipo, SPECS)
        self.seed(self.sme, SME_SPECS)
        self.reset_prices()

    def tearDown(self):
        self.page.remove_listener("console", self._on_console)
        super().tearDown()

    def _on_console(self, msg):
        if msg.type == "error" and not msg.text.startswith("Failed to load resource"):
            self.js_errors.append(msg.text)

    # ------------------------------------------------------------------------- seeding
    def group(self, name):
        return GroupDetail.objects.get_or_create(user=self.user, GroupName=name)[0]

    def seed(self, ipo, specs):
        """Idempotent. Details start at the canonical pre-open price."""
        for s in specs:
            order = Order.objects.filter(
                user=self.user, OrderIPOName=ipo, OrderGroup__GroupName=s["group"], OrderType=s["type"],
                OrderCategory=s["cat"], InvestorType=s["inv"], Rate=s["rate"], Method=s.get("method")).first()
            if order is not None:
                continue
            alloted = s.get("alloted")
            order = Order.objects.create(
                user=self.user, OrderIPOName=ipo, OrderGroup=self.group(s["group"]), OrderType=s["type"],
                OrderCategory=s["cat"], InvestorType=s["inv"], Rate=s["rate"],
                Quantity=len(alloted) if alloted is not None else s["qty"], Method=s.get("method"),
                OrderDate=SEED_DATE, OrderTime=SEED_TIME)
            for i, qty in enumerate(alloted or []):
                pan = (s.get("pans") or [None] * len(alloted))[i]
                client = None
                if pan:
                    client = ClientDetail.objects.get_or_create(
                        user=self.user, PANNo=pan, defaults={"Name": "BILLING CLIENT", "Group": order.OrderGroup})[0]
                OrderDetail.objects.create(user=self.user, Order=order, AllotedQty=qty, PreOpenPrice=PRE_OPEN,
                                           OrderDetailPANNo=client)

    def reset_prices(self):
        """Back to the canonical state: pre-open 250 everywhere, amounts recalculated by the app."""
        for ipo in (self.bipo, self.sme):
            ipo.PreOpenPrice = PRE_OPEN
            ipo.save()
            OrderDetail.objects.filter(Order__OrderIPOName=ipo).update(PreOpenPrice=PRE_OPEN)
            calculate(ipo.id, self.user)

    # ----------------------------------------------- independent oracle (formulas above)
    detail_amount = staticmethod(oracle.detail_amount)
    order_amount = staticmethod(oracle.order_amount)

    def assert_amounts_match_formulas(self, ipo):
        ipo.refresh_from_db()
        for od in OrderDetail.objects.filter(Order__OrderIPOName=ipo).select_related("Order"):
            self.assertAlmostEqual(od.Amount, self.detail_amount(od, ipo), places=6,
                                   msg=f"{od.Order.OrderCategory} {od.Order.OrderType} rate {od.Order.Rate} "
                                       f"alloted {od.AllotedQty} pre-open {od.PreOpenPrice}")
        for o in Order.objects.filter(OrderIPOName=ipo, OrderCategory__in=["Premium", "CALL", "PUT"]):
            self.assertAlmostEqual(o.Amount, self.order_amount(o, ipo), places=6,
                                   msg=f"{o.OrderCategory} {o.OrderType} strike {o.Method}")

    # --------------------------------------------- what the page should show (from the DB)
    def db_scope(self, ipo, group=None, cat=None, inv=None):
        details = OrderDetail.objects.filter(user=self.user, Order__OrderIPOName=ipo).select_related(
            "Order__OrderGroup", "OrderDetailPANNo")
        orders = Order.objects.filter(user=self.user, OrderIPOName=ipo,
                                      OrderCategory__in=["Premium", "CALL", "PUT"]).select_related("OrderGroup")
        if group:
            details, orders = details.filter(Order__OrderGroup__GroupName=group), orders.filter(OrderGroup__GroupName=group)
        if cat:
            details, orders = details.filter(Order__OrderCategory=cat), orders.filter(OrderCategory=cat)
        if inv:
            details, orders = details.filter(Order__InvestorType=inv), orders.filter(InvestorType=inv)
        return details, orders

    def db_rows(self, ipo, **filters):
        """Counter of the rows the page must show: (group, category, investor, type, rate, pan, pre-open, qty, amount)."""
        show_inv = ipo.IPOType == "MAINBOARD"
        details, orders = self.db_scope(ipo, **filters)
        rows = Counter()
        for od in details:
            o = od.Order
            rows[(o.OrderGroup.GroupName, o.OrderCategory, o.InvestorType if show_inv else "", o.OrderType,
                  float(o.Rate), od.OrderDetailPANNo.PANNo if od.OrderDetailPANNo else "",
                  float(od.PreOpenPrice or ipo.PreOpenPrice),
                  0.0 if od.AllotedQty is None else float(f"{od.AllotedQty:.0f}"), int(float(f"{od.Amount:.0f}")))] += 1
        for o in orders:
            rows[(o.OrderGroup.GroupName, o.OrderCategory, o.InvestorType if show_inv else "", o.OrderType,
                  float(o.Rate), "-", float(ipo.PreOpenPrice), float(f"{o.Quantity:.0f}"),
                  int(float(f"{o.Amount:.0f}")))] += 1
        return rows

    def db_total(self, ipo, **filters):
        details, orders = self.db_scope(ipo, **filters)
        return sum(od.Amount for od in details) + sum(o.Amount for o in orders)

    # ---------------------------------------------------------------------- page helpers
    def billing_url(self, ipo=None, group="All", cat="All", inv="All", query=""):
        ipo = ipo or self.bipo
        if (group, cat, inv) == ("All", "All", "All") and not query.startswith("?page"):
            return f"{self.live_server_url}/{ipo.id}/Billing{query}"
        return f"{self.live_server_url}/{ipo.id}/Billing/{quote(group)}/{quote(cat)}/{quote(inv)}{query}"

    def open_billing(self, ipo=None, group="All", cat="All", inv="All", query=""):
        self.page.goto(self.billing_url(ipo, group, cat, inv, query))
        settle(self.page)
        if self.page.locator("#example").count() == 0:
            self.fail(f"no billing grid at {self.page.url}: {self.page.inner_text('body')[:300]!r}")
        expect(self.page.locator("#example")).to_be_visible()

    def ui(self):
        """Visible headers, rows (dicts by header) and checkbox ids currently in the grid."""
        raw = self.page.evaluate("""() => {
            const dt = $('#example').DataTable();
            const heads = dt.columns(':visible').header().toArray().map(h => h.textContent.trim());
            const rows = [...document.querySelectorAll('#example tbody tr')]
                .filter(tr => !tr.querySelector('td.dataTables_empty'))
                .map(tr => ({cells: [...tr.querySelectorAll('td')].map(td => td.textContent.trim()),
                             id: tr.querySelector('.billing-row-checkbox')?.value ?? null}));
            const foot = [...document.querySelectorAll('#example tfoot th, #example tfoot td')].map(c => c.textContent.trim());
            return {heads, rows, foot};
        }""")
        dicts = []
        for r in raw["rows"]:
            d = dict(zip(raw["heads"], r["cells"]))
            d["_id"] = r["id"]
            dicts.append(d)
        return raw["heads"], dicts, raw["foot"]

    def ui_signature(self, show_inv=True):
        _heads, rows, _foot = self.ui()
        sig = Counter()
        for r in rows:
            alloted = r["Alloted Qty"]
            sig[(r["Group"], r["Order Category"], r.get("Investor Type", "") if show_inv else "", r["Order Type"],
                 float(r["Rate"]), r["PAN No"], float(r["Pre-Open Price"]),
                 0.0 if alloted == "" else float(alloted), int(float(r["Amount"])))] += 1     # not allotted: 0 or blank
        return sig

    def summary(self):
        """(total, ipo price, pre-open price) from the summary card."""
        cells = self.page.locator("#bootstrapdatatable tbody td.tablecenter").all_inner_texts()
        return tuple(float(c.strip().replace(",", "")) for c in cells[:3])

    def choose(self, selector, value):
        """The filter dropdowns POST the form on change."""
        with self.page.expect_navigation():
            self.page.select_option(selector, value)
        settle(self.page)

    def set_page_size(self, size):
        with self.page.expect_navigation():
            self.page.select_option("#Billing_page_size", size)
        settle(self.page)

    def relogin(self):
        self.page.context.clear_cookies()
        self.page.context.add_cookies([{"name": "sessionid", "value": self.client.cookies["sessionid"].value,
                                        "domain": "localhost", "path": "/"}])

    def assert_same_rows(self, on_screen, in_db, label=""):
        if on_screen != in_db:
            self.fail(f"{label} page rows differ from the database\n"
                      f"  only on screen: {sorted((on_screen - in_db).elements(), key=str)}\n"
                      f"  only in DB:     {sorted((in_db - on_screen).elements(), key=str)}")

    def assert_page_matches_db(self, ipo=None, **filters):
        ipo = ipo or self.bipo
        self.assert_same_rows(self.ui_signature(ipo.IPOType == "MAINBOARD"), self.db_rows(ipo, **filters),
                              str(filters))

    # ======================================================================================
    # Section 1: billing calculations
    # ======================================================================================
    def test_kostak_amounts_follow_the_formula(self):
        """BUY / SELL, allotted > 0, 0 and None."""
        self.assert_amounts_match_formulas(self.bipo)
        kostak = {(od.Order.OrderType, od.AllotedQty): od.Amount for od in OrderDetail.objects.filter(
            Order__OrderIPOName=self.bipo, Order__OrderCategory="Kostak", Order__OrderGroup__GroupName=ALPHA
        ).select_related("Order")}
        gain = PRE_OPEN - IPO_PRICE
        self.log(f"Kostak ALPHA: BUY 10 -> {kostak[('BUY', 10.0)]:.0f}, BUY 0 -> {kostak[('BUY', 0.0)]:.0f}, "
                 f"BUY None -> {kostak[('BUY', None)]:.0f}, SELL 10 -> {kostak[('SELL', 10.0)]:.0f}", indent=2)
        self.assertEqual(kostak[("BUY", 10.0)], gain * 10 - 450)        # 1050
        self.assertEqual(kostak[("BUY", 0.0)], -450)                     # quantity 0 still pays the rate
        self.assertEqual(kostak[("BUY", None)], 0)                       # not allotted yet
        self.assertEqual(kostak[("SELL", 10.0)], -(gain * 10 - 455))     # -1045
        self.assertEqual(kostak[("SELL", None)], 0)
        self.open_billing(group=ALPHA, cat="Kostak")
        self.assert_page_matches_db(group=ALPHA, cat="Kostak")

    def test_subject_to_amounts_follow_the_formula(self):
        """Fixed and Premium method, BUY / SELL, and allotted 0."""
        self.assert_amounts_match_formulas(self.bipo)
        gain = PRE_OPEN - IPO_PRICE
        amounts = {(od.Order.OrderType, od.Order.Method, od.Order.Rate, od.AllotedQty): od.Amount
                   for od in OrderDetail.objects.filter(Order__OrderIPOName=self.bipo, Order__OrderCategory="Subject To"
                                                        ).select_related("Order")}
        self.assertEqual(amounts[("BUY", None, 1200.0, 10.0)], gain * 10 - 1200)             # fixed: 300
        self.assertEqual(amounts[("BUY", None, 1200.0, 0.0)], 0)                              # allotted 0 -> nothing
        self.assertEqual(amounts[("BUY", "Premium", 5.0, 10.0)], gain * 10 - 5 * 10)          # premium: R * Q
        self.assertEqual(amounts[("BUY", "Premium", 5.0, 20.0)], gain * 20 - 5 * 20)
        self.assertEqual(amounts[("SELL", None, 1210.0, 10.0)], -(gain * 10 - 1210))
        self.assertEqual(amounts[("SELL", "Premium", 6.0, 10.0)], -(gain * 10 - 6 * 10))
        self.log(f"fixed BUY 300, premium BUY 1450 / 2900, fixed SELL -290, premium SELL -1440 all match", indent=2)
        self.open_billing(cat="Subject To")
        self.set_page_size("All")
        self.assert_page_matches_db(cat="Subject To")

    def test_premium_and_option_amounts_follow_the_formulas(self):
        self.assert_amounts_match_formulas(self.bipo)
        by = {(o.OrderCategory, o.OrderType, o.Method): o.Amount
              for o in Order.objects.filter(OrderIPOName=self.bipo, OrderCategory__in=["Premium", "CALL", "PUT"])}
        self.assertEqual(by[("Premium", "BUY", None)], 13000)       # (250 - (100+20)) * 100
        self.assertEqual(by[("Premium", "SELL", None)], -6000)      # -((250 - (100+30)) * 50)
        # CALL / PUT: pay-off against pre-open 250 with strikes 120 (in the money for CALL) and 200
        expected = {("CALL", "BUY", "120"): 2000, ("CALL", "BUY", "200"): -1000,
                    ("CALL", "SELL", "120"): -2000, ("CALL", "SELL", "200"): 1000,
                    ("PUT", "BUY", "120"): -1000, ("PUT", "BUY", "200"): 4000,
                    ("PUT", "SELL", "120"): 1000, ("PUT", "SELL", "200"): -4000}
        for key, amount in expected.items():
            self.assertEqual(by[key], amount, f"{key}")
        self.log("Premium BUY 13000 / SELL -6000 and all 8 CALL/PUT pay-offs match", indent=2)
        for cat in ("Premium", "CALL", "PUT"):
            with self.subTest(category=cat):
                self.open_billing(cat=cat)
                self.assert_page_matches_db(cat=cat)

    def test_broker_sees_every_group_and_the_total_covers_them_all(self):
        self.open_billing()
        self.assertAlmostEqual(self.summary()[0], self.db_total(self.bipo), places=1)
        groups = {g.strip() for g in self.page.locator("#Groupfilter option").all_inner_texts()}
        self.assertTrue({"All", ALPHA, BETA, SIGMA, PAGER} <= groups, groups)
        per_group = {g: self.db_total(self.bipo, group=g) for g in (ALPHA, BETA, SIGMA, PAGER)}
        self.assertAlmostEqual(sum(per_group.values()), self.db_total(self.bipo), places=6)
        self.log("broker total " + f"{self.db_total(self.bipo):.0f} = " +
                 " + ".join(f"{g.split()[0]} {v:.0f}" for g, v in per_group.items()), indent=2)

    def make_customer(self, group=ALPHA):
        user_model = get_user_model()
        customer, _ = user_model.objects.get_or_create(username=CUSTOMER_USER, defaults={"email": "c@example.com"})
        customer.set_password(PASSWORD)
        customer.Broker_id, customer.Group_id = str(self.user.id), str(self.group(group).id)
        customer.save()
        customer.groups.set([Group.objects.get_or_create(name="Customer")[0]])
        browser = Client()
        self.assertTrue(browser.login(username=CUSTOMER_USER, password=PASSWORD))
        self.page.context.clear_cookies()
        self.page.context.add_cookies([{"name": "sessionid", "value": browser.cookies["sessionid"].value,
                                        "domain": "localhost", "path": "/"}])

    @unittest.expectedFailure
    def test_client_user_only_sees_their_own_group(self):
        """KNOWN APP BUG: the Billing page returns HTTP 500 for a client (non-broker) user: for them `entry` is a
        queryset of OrderDetail objects, but the loop reads row["id"] like the broker's .values() dicts
        (views.py ~9574, TypeError: 'OrderDetail' object is not subscriptable). The filtered URL
        /Billing/All/All/All is Broker-only, so a client gets 'You are not authorized' there.
        Asserts the CORRECT behaviour: a client of ALPHA TRADERS must not see BETA INVESTORS / SIGMA CAPITAL
        data - rows, groups or totals. 'Unexpected success' once the page works for clients."""
        self.make_customer(ALPHA)
        try:
            self.open_billing()
            self.set_page_size("All")
            _heads, rows, _foot = self.ui()
            self.log(f"client of {ALPHA} sees {len(rows)} rows in groups {sorted({r['Group'] for r in rows})}", indent=2)
            self.assertEqual({r["Group"] for r in rows}, {ALPHA})
            self.assert_same_rows(self.ui_signature(), self.db_rows(self.bipo, group=ALPHA), 'client of ALPHA')
            self.assertAlmostEqual(self.summary()[0], self.db_total(self.bipo, group=ALPHA), places=1)
            options = {g.strip() for g in self.page.locator("#Groupfilter option").all_inner_texts()}
            self.assertEqual(options - {"All"}, {ALPHA})
            body = self.page.content()
            for other in (BETA, SIGMA, PAGER):
                self.assertNotIn(other, body)
        finally:
            self.relogin()

    # ======================================================================================
    # Section 2: filters
    # ======================================================================================
    def test_group_filter(self):
        self.open_billing()
        for group in (ALPHA, BETA, SIGMA, "All"):
            with self.subTest(group=group):
                self.choose("#Groupfilter", group)
                self.set_page_size("All")
                _h, rows, _f = self.ui()
                if group != "All":
                    self.assertEqual({r["Group"] for r in rows}, {group})
                self.assert_page_matches_db(group=None if group == "All" else group)
                expect(self.page.locator("#Groupfilter")).to_have_value(group)
        self.assertEqual(self.js_errors, [])

    def test_order_category_filter(self):
        self.open_billing()
        for cat in ("Kostak", "Subject To", "Premium", "CALL", "PUT"):
            with self.subTest(category=cat):
                self.choose("select[name=IPOTypefilter]", cat)
                self.set_page_size("All")
                _h, rows, _f = self.ui()
                self.assertGreater(len(rows), 0)
                self.assertEqual({r["Order Category"] for r in rows}, {cat})
                self.assert_page_matches_db(cat=cat)

    def test_investor_type_filter_on_a_mainboard_ipo(self):
        self.open_billing()
        for inv in ("RETAIL", "SHNI", "BHNI"):
            with self.subTest(investor=inv):
                self.choose("select[name=InvestorTypeFilter]", inv)
                self.set_page_size("All")
                _h, rows, _f = self.ui()
                self.assertGreater(len(rows), 0)
                self.assertEqual({r["Investor Type"] for r in rows}, {inv})
                self.assert_page_matches_db(inv=inv)

    def test_investor_type_filter_is_hidden_on_an_sme_ipo(self):
        self.open_billing(self.sme)
        self.assertEqual(self.page.locator("select[name=InvestorTypeFilter]").count(), 0)
        heads, rows, _f = self.ui()
        self.assertNotIn("Investor Type", heads)
        self.log(f"SME IPO: no investor filter, columns = {heads}", indent=2)
        self.assert_same_rows(self.ui_signature(show_inv=False), self.db_rows(self.sme), 'SME')
        # an investor segment typed into the URL of an SME IPO still works (it filters on RETAIL)
        response = self.page.goto(self.billing_url(self.sme, "All", "All", "RETAIL"))
        settle(self.page)
        self.assertEqual(response.status, 200)
        self.assertEqual({r["Order Category"] for r in self.ui()[1]}, {"Kostak"})

    def test_combined_group_category_investor_filters_and_search(self):
        combos = [(ALPHA, "Kostak", "RETAIL"), (BETA, "Subject To", "SHNI"), (BETA, "Kostak", "BHNI"),
                  (SIGMA, "Kostak", "SHNI"), (BETA, "Kostak", "RETAIL")]       # the last one has no orders
        for group, cat, inv in combos:
            with self.subTest(group=group, category=cat, investor=inv):
                self.open_billing()
                self.choose("#Groupfilter", group)
                self.choose("select[name=IPOTypefilter]", cat)
                self.choose("select[name=InvestorTypeFilter]", inv)
                want = self.db_rows(self.bipo, group=group, cat=cat, inv=inv)
                self.log(f"{group} + {cat} + {inv} -> {sum(want.values())} rows", indent=2)
                self.assertEqual(self.ui_signature(), want)
                if not want:
                    expect(self.page.locator("#example td.dataTables_empty")).to_have_text("No data available")
                    self.assertAlmostEqual(self.summary()[0], 0.0, places=2)
        self.assertEqual(self.js_errors, [])

    def test_search_box_matches_without_duplicate_rows(self):
        self.open_billing()
        search = self.page.locator("#example_filter input")

        def search_for(term):
            with self.page.expect_navigation():                    # fires ~500 ms after typing stops
                search.fill(term)
            settle(self.page)

        search_for("ABCDE1234F")
        _h, rows, _f = self.ui()
        self.assertEqual([(r["Group"], r["PAN No"]) for r in rows], [(ALPHA, "ABCDE1234F")])
        search_for("SIGMA")
        self.assertEqual({r["Group"] for r in self.ui()[1]}, {SIGMA})
        self.assert_same_rows(self.ui_signature(), self.db_rows(self.bipo, group=SIGMA), 'SIGMA search')
        search_for("FGHIJ5678K")
        self.assertEqual([r["PAN No"] for r in self.ui()[1]], ["FGHIJ5678K"])
        search_for("")
        self.assertEqual(sum(self.ui_signature().values()), 50)     # back to a full first page
        self.assertEqual(self.js_errors, [])

    def test_the_all_filter_url_shows_the_same_rows_as_the_plain_billing_url(self):
        self.open_billing()
        self.set_page_size("All")
        plain = self.ui_signature()
        self.page.goto(self.billing_url(self.bipo, "All", "All", "All", query="?page=1"))
        settle(self.page)
        self.set_page_size("All")
        self.assertEqual(self.ui_signature(), plain)
        self.assert_same_rows(plain, self.db_rows(self.bipo), 'plain URL')

    @unittest.expectedFailure
    def test_rate_segment_in_the_url_filters_by_rate(self):
        """KNOWN GAP: /<IPOid>/Billing/<group>/<category>/<investor>/<rate> accepts a rate, but the view never
        uses it and the page has no rate dropdown, so every rate is shown. Asserts the CORRECT behaviour."""
        self.page.goto(f"{self.billing_url(self.bipo, ALPHA, 'Kostak', 'RETAIL')}/450")
        settle(self.page)
        _h, rows, _f = self.ui()
        self.assertGreater(len(rows), 0)
        self.assertEqual({float(r["Rate"]) for r in rows}, {450.0})

    # ======================================================================================
    # Section 3: exports
    # ======================================================================================
    def export(self, kind):
        self.page.get_by_role("button", name="Download").click()
        self.page.locator(".dt-button-collection .dt-button", has_text=kind).click()
        popup = self.page.locator(".swal2-popup")
        expect(popup).to_be_visible()
        with self.page.expect_download() as info:
            popup.locator(".swal2-confirm").click()
        download = info.value
        path = os.path.join(self.tmpdir, download.suggested_filename)
        download.save_as(path)
        return download.suggested_filename, path

    def read_excel(self, path):
        rows = [list(r) for r in openpyxl.load_workbook(path).active.iter_rows(values_only=True)]
        head_idx = next(i for i, r in enumerate(rows) if "Group" in r)
        foot_idx = next(i for i, r in enumerate(rows) if i > head_idx and r and r[0] == "Total")
        norm = lambda r: ["" if c is None else str(c).strip() for c in r]
        return norm(rows[head_idx]), [norm(r) for r in rows[head_idx + 1:foot_idx]], norm(rows[foot_idx])

    def test_excel_export_filename_columns_rows_and_total(self):
        cases = [((), "{ipo} - Client Wise Billing"),
                 ((ALPHA,), "{ipo} - Client Wise Billing - ALPHA TRADERS"),
                 ((ALPHA, "Kostak"), "{ipo} - Client Wise Billing - ALPHA TRADERS - Kostak"),
                 ((ALPHA, "Kostak", "RETAIL"), "{ipo} - Client Wise Billing - ALPHA TRADERS - Kostak - RETAIL")]
        with tempfile.TemporaryDirectory() as tmp:
            self.tmpdir = tmp
            for filters, stem in cases:
                with self.subTest(filters=" / ".join(filters) or "none"):
                    self.open_billing(group=filters[0] if filters else "All",
                                      cat=filters[1] if len(filters) > 1 else "All",
                                      inv=filters[2] if len(filters) > 2 else "All")
                    heads, ui_rows, foot = self.ui()
                    name, path = self.export("Excel")
                    self.log(f"exported '{name}' ({len(ui_rows)} rows on screen)", indent=2)
                    self.assertEqual(name, stem.format(ipo=BILL_IPO) + ".xlsx")
                    x_heads, x_rows, x_foot = self.read_excel(path)
                    self.assertEqual(x_heads, heads, "Excel columns differ from the grid")
                    self.assertEqual(len(x_rows), len(ui_rows), "Excel row count differs from the grid")
                    wanted = Counter((r["Group"], r["Order Category"], r["Order Type"], float(r["Rate"]),
                                      int(float(r["Amount"]))) for r in ui_rows)
                    got = Counter((r[x_heads.index("Group")], r[x_heads.index("Order Category")],
                                   r[x_heads.index("Order Type")], float(r[x_heads.index("Rate")]),
                                   int(float(r[x_heads.index("Amount")]))) for r in x_rows)
                    self.assertEqual(got, wanted)
                    total_on_screen = next(c for c in reversed(foot) if c != "")
                    self.assertEqual(int(float(x_foot[-1])), int(float(total_on_screen)),
                                     "Excel Total row differs from the Total shown under the grid")
                    self.assertEqual(int(float(total_on_screen)), sum(int(float(r["Amount"])) for r in ui_rows),
                                     "the Total under the grid must equal the sum of the rows shown")
        self.assertEqual(self.js_errors, [])

    def test_datatables_pdf_export(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.tmpdir = tmp
            self.open_billing(group=ALPHA)
            name, path = self.export("PDF")
            self.assertEqual(name, f"{BILL_IPO} - Client Wise Billing - {ALPHA}.pdf")
            self.assertEqual(open(path, "rb").read(5), b"%PDF-")

    def test_server_csv_export_route(self):
        """The 'Download File' modal that posts here is commented out of the template, but the route is live."""
        response = self.page.request.get(f"{self.live_server_url}/{self.bipo.id}/download-Billing/{quote(ALPHA)}/Kostak/RETAIL")
        self.assertEqual(response.status, 200)
        self.assertIn("text/csv", response.headers["content-type"])
        self.assertIn(f'filename="{BILL_IPO}-Billing.csv"', response.headers["content-disposition"])
        rows = list(csv.reader(io.StringIO(response.text())))
        self.assertEqual(rows[0][:4], ["Group", "Order Category", "Investor Type", "Order Type"])
        details, orders = self.db_scope(self.bipo, group=ALPHA, cat="Kostak", inv="RETAIL")
        self.assertEqual(len(rows) - 1, details.count() + orders.filter(OrderCategory="Premium").count())
        self.assertEqual({r[0] for r in rows[1:]}, {ALPHA})

    def test_server_pdf_export_route(self):
        response = self.page.request.get(f"{self.live_server_url}/{self.bipo.id}/download-Billing-Pdf/All/All/All")
        self.assertEqual(response.status, 200)
        self.assertEqual(response.headers["content-type"], "application/pdf")
        self.assertRegex(response.headers["content-disposition"], rf'attachment; filename="{re.escape(BILL_IPO)}-Billing\.PDF"')
        self.assertEqual(response.body()[:5], b"%PDF-")
        if shutil.which("pdftotext"):
            with tempfile.NamedTemporaryFile(suffix=".pdf") as f:
                f.write(response.body())
                f.flush()
                text = subprocess.run(["pdftotext", "-layout", f.name, "-"], capture_output=True, text=True).stdout
            self.assertIn(BILL_IPO, text)
            self.assertIn("TOTAL", text)
            self.assertIn(f"{self.db_total(self.bipo):.0f}", text)

    def test_groupwise_export_route(self):
        response = self.page.request.get(f"{self.live_server_url}/download-Groupwise")
        self.assertEqual(response.status, 200)
        self.assertIn('filename="GroupWiseDashboard.xlsx"', response.headers["content-disposition"])
        workbook = openpyxl.load_workbook(io.BytesIO(response.body()))        # the body really is an .xlsx
        self.assertIn("Total", [str(c.value) for c in workbook.active["A"]])

    @unittest.expectedFailure
    def test_groupwise_export_declares_an_excel_content_type(self):
        """KNOWN APP BUG: /download-Groupwise sends an .xlsx body but declares Content-Type text/csv."""
        response = self.page.request.get(f"{self.live_server_url}/download-Groupwise")
        self.assertIn("spreadsheetml", response.headers["content-type"])

    # ======================================================================================
    # Section 4: summary card and pre-open price edits
    # ======================================================================================
    def test_summary_card_matches_the_database(self):
        self.open_billing()
        total, ipo_price, pre_open = self.summary()
        self.assertAlmostEqual(total, self.db_total(self.bipo), places=1)
        self.assertEqual((ipo_price, pre_open), (self.bipo.IPOPrice, self.bipo.PreOpenPrice))
        self.open_billing(group=BETA)
        self.assertAlmostEqual(self.summary()[0], self.db_total(self.bipo, group=BETA), places=1)
        self.log(f"IPO price {ipo_price}, pre-open {pre_open}, total {total:.0f}", indent=2)

    def open_global_edit(self):
        self.page.locator("i.fa-pencil-square-o[data-target='#editbtn']").click()
        modal = self.page.locator("#editbtn")
        expect(modal).to_have_class(re.compile(r"\bshow\b"))
        return modal

    def submit_modal(self, modal, value):
        modal.locator("input[name=PreOpenPrice]").fill(value)
        with self.page.expect_navigation() as nav:
            modal.get_by_role("button", name="Submit").click()
        settle(self.page)
        return nav.value

    def test_global_preopen_price_update_recalculates_everything(self):
        self.open_billing()
        before = self.summary()[0]
        modal = self.open_global_edit()
        expect(modal.locator(".modal-title")).to_have_text("Pre-Open Price Edit (All Records)")
        self.log("setting the pre-open price of the whole IPO to 185.50", indent=2)
        self.submit_modal(modal, "185.50")

        self.bipo.refresh_from_db()
        self.assertEqual(self.bipo.PreOpenPrice, 185.5)
        self.assertEqual(set(OrderDetail.objects.filter(Order__OrderIPOName=self.bipo)
                             .values_list("PreOpenPrice", flat=True)), {185.5})
        self.assert_amounts_match_formulas(self.bipo)                 # Kostak, Subject To, Premium, CALL and PUT
        self.assertIn("/Billing", self.page.url)
        total, _price, pre_open = self.summary()
        self.assertEqual(pre_open, 185.5)
        self.assertAlmostEqual(total, self.db_total(self.bipo), places=1)
        self.assertNotAlmostEqual(total, before, places=1)
        self.log(f"total moved from {before:.0f} to {total:.0f}", indent=2)
        self.set_page_size("All")
        self.assert_same_rows(self.ui_signature(), self.db_rows(self.bipo), 'all rows')
        self.assertEqual(self.js_errors, [])

    def test_bulk_preopen_update_changes_only_the_selected_rows(self):
        self.open_billing(group=ALPHA, cat="Kostak", inv="RETAIL")
        _h, rows, _f = self.ui()
        chosen = [r["_id"] for r in rows[:2]]
        self.assertEqual(len(chosen), 2)
        for box_id in chosen:
            self.page.locator(f".billing-row-checkbox[value='{box_id}']").check()
        modal = self.open_global_edit()
        expect(modal.locator(".modal-title")).to_have_text("Bulk Pre-Open Price Edit (2 records)")
        self.assertEqual(self.page.locator("#top-selected-ids-input").input_value(), ",".join(chosen))
        self.log(f"bulk edit of details {chosen} -> 185.50", indent=2)
        self.submit_modal(modal, "185.50")

        prices = dict(OrderDetail.objects.filter(Order__OrderIPOName=self.bipo).values_list("id", "PreOpenPrice"))
        for od_id, price in prices.items():
            self.assertEqual(price, 185.5 if str(od_id) in chosen else PRE_OPEN, f"detail {od_id}")
        self.bipo.refresh_from_db()
        self.assertEqual(self.bipo.PreOpenPrice, PRE_OPEN, "a bulk edit must not change the IPO-level price")
        self.assert_amounts_match_formulas(self.bipo)
        self.assertAlmostEqual(self.summary()[0], self.db_total(self.bipo, group=ALPHA, cat="Kostak", inv="RETAIL"), places=1)
        self.assertEqual(self.js_errors, [])

    def test_preopen_price_boundaries_empty_and_non_numeric_input(self):
        self.open_billing()
        modal = self.open_global_edit()
        field = modal.locator("input[name=PreOpenPrice]")
        # characters other than digits and "." are stripped as you type
        for typed in ("abc", "12abc.5", "1,2;3"):
            field.fill(typed)
            self.assertRegex(field.input_value(), r"^[0-9.]*$", f"typed {typed!r}")
        self.assertEqual((field.fill("abc"), field.input_value())[1], "")
        # an empty value falls back to 0.0 - no server error
        self.log("submitting an empty price -> falls back to 0.0", indent=2)
        response = self.submit_modal(modal, "")
        self.assertEqual(response.status, 200)
        self.bipo.refresh_from_db()
        self.assertEqual(self.bipo.PreOpenPrice, 0.0)
        self.assertEqual(set(OrderDetail.objects.filter(Order__OrderIPOName=self.bipo)
                             .values_list("PreOpenPrice", flat=True)), {0.0})
        self.assert_amounts_match_formulas(self.bipo)
        self.assertEqual(self.summary()[2], 0.0)
        # ... and a normal value afterwards works again
        modal = self.open_global_edit()
        self.submit_modal(modal, "250")
        self.bipo.refresh_from_db()
        self.assertEqual(self.bipo.PreOpenPrice, 250.0)
        self.assertEqual(self.js_errors, [])

    # ======================================================================================
    # Section 5: in-grid behaviour
    # ======================================================================================
    def test_select_all_and_selected_ids(self):
        self.open_billing(group=ALPHA, cat="Kostak")
        boxes = self.page.locator(".billing-row-checkbox")
        count = boxes.count()
        self.assertGreater(count, 1)
        select_all = self.page.locator("#select-all-billing")
        select_all.check()
        self.assertEqual(self.page.locator(".billing-row-checkbox:checked").count(), count)
        select_all.uncheck()
        self.assertEqual(self.page.locator(".billing-row-checkbox:checked").count(), 0)

        select_all.check()
        ids = [r["_id"] for r in self.ui()[1]]
        self.page.locator("#example tbody a[data-target='#edit-modal']").first.click()
        modal = self.page.locator("#edit-modal")
        expect(modal).to_have_class(re.compile(r"\bshow\b"))
        self.assertEqual(self.page.locator("#selected-ids-input").input_value(), ",".join(ids))
        expect(modal.locator(".modal-title")).to_have_text(f"Bulk Pre-Open Price Edit ({count} rows)")
        modal.locator("button.close").click()
        expect(modal).to_be_hidden()
        # one box unticked -> the list shrinks accordingly
        self.page.locator(f".billing-row-checkbox[value='{ids[0]}']").uncheck()
        self.page.locator("#example tbody a[data-target='#edit-modal']").first.click()
        expect(modal).to_have_class(re.compile(r"\bshow\b"))
        self.assertEqual(self.page.locator("#selected-ids-input").input_value(), ",".join(ids[1:]))

    def test_single_row_preopen_edit_updates_one_row_and_keeps_the_page(self):
        self.open_billing(group=PAGER, query="?page=2")
        self.assertIn("page=2", self.page.url)
        link = self.page.locator("#example tbody a[data-target='#edit-modal']").first
        row_id = link.get_attribute("data-id")
        modal = self.page.locator("#edit-modal")
        link.click()
        expect(modal).to_have_class(re.compile(r"\bshow\b"))
        self.assertEqual(self.page.locator("#preopen-price").input_value(), "250.0")
        self.assertIn("page=2", self.page.locator("#edit-form").get_attribute("action"))
        self.assertEqual(self.page.locator("#selected-ids-input").input_value(), "")
        self.page.locator("#preopen-price").fill("190")
        self.log(f"editing detail {row_id} on page 2 -> 190", indent=2)
        with self.page.expect_navigation():
            self.page.locator("#edit-form button[type=submit]").click()
        settle(self.page)

        self.assertIn("page=2", self.page.url, "the user must stay on page 2")
        expect(self.page.locator(".pagination .page-item.active")).to_have_text("2")
        edited = OrderDetail.objects.get(id=row_id)
        self.assertEqual(edited.PreOpenPrice, 190.0)
        self.assertEqual(edited.Amount, (190.0 - IPO_PRICE) * 1 - 111)                       # -21
        others = OrderDetail.objects.filter(Order__OrderGroup__GroupName=PAGER).exclude(id=row_id)
        self.assertEqual(set(others.values_list("PreOpenPrice", flat=True)), {PRE_OPEN})
        self.assertEqual(set(others.values_list("Amount", flat=True)), {(PRE_OPEN - IPO_PRICE) - 111})   # 39
        self.assert_amounts_match_formulas(self.bipo)

    def column(self, header):
        return [float(r[header]) for r in self.ui()[1]]

    def click_header(self, header):
        with self.page.expect_navigation():
            self.page.locator(".dataTables_scrollHeadInner thead tr:not(.filters) th").filter(
                has_text=re.compile(rf"^{re.escape(header)}$")).first.click()
        settle(self.page)

    def test_sorting_numeric_columns_is_numeric_not_alphabetical(self):
        self.open_billing(group=ALPHA)
        for header in ("Rate", "Amount"):
            with self.subTest(column=header):
                self.click_header(header)
                ascending = self.column(header)
                self.assertEqual(ascending, sorted(ascending), f"{header} ascending")
                self.assertGreater(len(set(ascending)), 3)
                self.click_header(header)
                descending = self.column(header)
                self.assertEqual(descending, sorted(descending, reverse=True), f"{header} descending")
        # the same values sorted as TEXT would be in a different order -> proves the check can tell them apart
        rates = [str(int(r)) for r in self.column("Rate")]
        self.assertNotEqual(sorted(rates), sorted(rates, key=float))

    def test_sorting_pre_open_price_numerically(self):
        details = OrderDetail.objects.filter(Order__OrderIPOName=self.bipo, Order__OrderGroup__GroupName=ALPHA,
                                             Order__OrderCategory="Kostak", Order__OrderType="BUY")
        for od, price in zip(details, (95, 1000, 250)):                 # "1000" < "250" < "95" as text
            OrderDetail.objects.filter(id=od.id).update(PreOpenPrice=price)
        self.open_billing(group=ALPHA, cat="Kostak")
        self.click_header("Pre-Open Price")
        prices = self.column("Pre-Open Price")
        self.assertEqual(prices, sorted(prices))
        self.assertEqual(prices[0], 95.0)
        self.assertEqual(prices[-1], 1000.0)

    def test_sorting_alloted_qty_with_unallotted_applications(self):
        """Blank / 0 first, then ascending numerically; descending reverses it."""
        self.open_billing(group=ALPHA, cat="Kostak")
        self.click_header("Alloted Qty")
        ascending = [0.0 if v == "" else float(v) for v in (r["Alloted Qty"] for r in self.ui()[1])]
        self.assertEqual(ascending, sorted(ascending))
        self.assertGreater(len(set(ascending)), 1)
        self.click_header("Alloted Qty")
        descending = [0.0 if v == "" else float(v) for v in (r["Alloted Qty"] for r in self.ui()[1])]
        self.assertEqual(descending, sorted(descending, reverse=True))

    @unittest.expectedFailure
    def test_unallotted_applications_look_the_same_on_both_billing_code_paths(self):
        """KNOWN APP INCONSISTENCY: an application with no allotment shows '0' on the filtered URL
        (FileterBilling) but is BLANK on the page you get by filtering or sorting from the grid (Billing).
        Asserts the CORRECT behaviour (identical); 'unexpected success' once the two views agree."""
        self.open_billing(group=ALPHA, cat="Kostak")
        via_url = sorted(r["Alloted Qty"] for r in self.ui()[1])
        self.click_header("Rate")
        via_form = sorted(r["Alloted Qty"] for r in self.ui()[1])
        self.assertEqual(via_url, via_form)
