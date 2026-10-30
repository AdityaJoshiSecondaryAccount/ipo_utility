"""
Back Up page E2E (/BackUp, templates/Backup.html; "Back Up" in the menu).

The page lists every IPO of the broker (newest first, server pages of 50 / 100 / 200 / 250 / All, kept in the session)
with a Backup button per IPO, plus "All IPOs Backup" and "Accounting Backup". There is NO restore / import: the page
only downloads.

  Backup (GET /<IPOid>/Backup/)   one .xlsx: Sheet1 = the IPO's orders (Group, OrderType, Order Category, [Investor
                                  Type - MAINBOARD only], Qty, Rate, Amount, Order Date, Order Time);
                                  Sheet2 = one row per application (OrderDetail: Group, Category, [Investor], Type, Rate,
                                  AllotedQty, Pre-Open Price, Amount, PAN No, Client Name) followed by the Premium orders.
  All IPOs Backup (POST /AllIpoBackup)  a .zip with one such .xlsx per IPO.
  Accounting Backup (GET /AccountingBackup/)  one .xlsx: "Accounting" (IPO, Group, type, amount, remark, date time)
                                  and "Accounting Logs" (the audit log of edits / deletes / restores).

Runs on its OWN IPOs (BKP MAIN IPO - MAINBOARD, BKP SME IPO - SME, BKP EMPTY IPO - no orders), groups (BKP ONE / TWO),
clients and accounting rows, rebuilt in setUp; IPOs a single test creates are removed again in that test. Every file
is opened and compared row by row with the database; the billing amounts are also re-computed with test_cases/oracle.py.
"""
import io
import os
import re
import tempfile
import unittest
import zipfile
from collections import Counter
from datetime import date, datetime, time, timedelta

import openpyxl
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.test import Client
from django.utils import timezone
from playwright.sync_api import expect

from home.models import Accounting, AccountingAuditLog, ClientDetail, CurrentIpoName, GroupDetail, Order, OrderDetail
from home.views import calculate
from test_cases import oracle
from test_cases.base import PersistentLiveServerTestCase, pause_for_inspection, settle

MAIN_IPO, SME_IPO, EMPTY_IPO = "BKP MAIN IPO", "BKP SME IPO", "BKP EMPTY IPO"
ONE, TWO = "BKP ONE", "BKP TWO"
IPO_PRICE, PRE_OPEN = 100.0, 250.0
REMARK = "BKP-ledger"
PANS = {"A": ("BKPAA1234A", "BKP CLIENT A"), "B": ("BKPAB1234B", "BKP CLIENT B"), "C": ("BKPAC1234C", "BKP CLIENT C")}

MAIN_SPECS = [
    dict(group=ONE, type="BUY", cat="Kostak", inv="RETAIL", rate=450, alloted=[10, 0, None], pans=["BKPAA1234A", "BKPAB1234B", None]),
    dict(group=ONE, type="SELL", cat="Kostak", inv="SHNI", rate=600, alloted=[5], pans=["BKPAC1234C"]),
    dict(group=TWO, type="BUY", cat="Subject To", inv="RETAIL", rate=1200, alloted=[10, 10]),
    dict(group=TWO, type="BUY", cat="Subject To", inv="BHNI", rate=5, method="Premium", alloted=[20]),
    dict(group=TWO, type="BUY", cat="Premium", inv="PREMIUM", rate=20, qty=100),
    dict(group=ONE, type="SELL", cat="Premium", inv="PREMIUM", rate=30, qty=40),
    dict(group=ONE, type="BUY", cat="CALL", inv="OPTIONS", rate=10, qty=100, method="120"),
]
SME_SPECS = [
    dict(group=ONE, type="BUY", cat="Kostak", inv="RETAIL", rate=300, alloted=[5, None]),
    dict(group=TWO, type="SELL", cat="Kostak", inv="RETAIL", rate=310, alloted=[2]),
    dict(group=TWO, type="BUY", cat="Premium", inv="PREMIUM", rate=10, qty=20),
]
ORDER_HEAD = {"MAINBOARD": ["Group", "OrderType", "Order Category", "Investor Type", "Qty", "Rate", "Amount", "Order Date", "Order Time"],
              "SME": ["Group", "OrderType", "Order Category", "Qty", "Rate", "Amount", "Order Date", "Order Time"]}
DETAIL_HEAD = {"MAINBOARD": ["Group", "Order Category", "Investor Type", "Order Type", "Rate", "AllotedQty", "Pre-Open Price", "Amount", "PAN No", "Client Name"],
               "SME": ["Group", "Order Category", "Order Type", "Rate", "AllotedQty", "Pre-Open Price", "Amount Diff.", "PAN No", "Client Name"]}

ROWS_JS = """() => [...document.querySelectorAll('#example tbody tr')].filter(tr => !tr.querySelector('td.dataTables_empty'))
    .map(tr => [...tr.children].map(c => c.textContent.replace(/\\s+/g, ' ').trim()))"""


def cell(v):
    """One spreadsheet value in a comparable form."""
    if isinstance(v, datetime):
        return v.date().isoformat() if v.time() == time(0) else v.isoformat(sep=" ")
    if isinstance(v, (date, time)):
        return v.isoformat()
    if isinstance(v, (int, float)):
        return round(float(v), 2)
    if v in (None, ""):
        return None
    return str(v).strip()


def db(v):
    """One database value in the same form."""
    if v is None:
        return None
    if isinstance(v, (date, time)):
        return v.isoformat()
    try:
        return round(float(v), 2)
    except (TypeError, ValueError):
        return str(v)


def sheet(workbook, name):
    rows = [[cell(c) for c in r] for r in workbook[name].iter_rows(values_only=True)]
    return rows[0], [r for r in rows[1:] if any(c is not None for c in r)]


class BackupPlaywrightTestCase(PersistentLiveServerTestCase):
    PAGE_NAME = "backup"
    _ran_any = False

    # ------------------------------------------------------------------ setup / teardown
    @classmethod
    def tearDownClass(cls):
        cls.log_summary()
        if cls._ran_any:
            try:
                cls.page.goto(f"{cls.live_server_url}/BackUp")
                settle(cls.page)
            except Exception:
                pass
            pause_for_inspection(cls.PAGE_NAME)
        super().tearDownClass()

    def setUp(self):
        super().setUp()
        type(self)._ran_any = True
        base = dict(IPOPrice=IPO_PRICE, PreOpenPrice=PRE_OPEN, LotSizeRetail=100, LotSizeSHNI=1400, LotSizeBHNI=7000, Active=True)
        self.ipos = {
            "main": CurrentIpoName.objects.get_or_create(user=self.user, IPOName=MAIN_IPO, defaults=dict(base, IPOType="MAINBOARD"))[0],
            "sme": CurrentIpoName.objects.get_or_create(user=self.user, IPOName=SME_IPO, defaults=dict(base, IPOType="SME"))[0],
            "empty": CurrentIpoName.objects.get_or_create(user=self.user, IPOName=EMPTY_IPO, defaults=dict(base, IPOType="MAINBOARD"))[0],
        }
        self.groups = {n: GroupDetail.objects.get_or_create(user=self.user, GroupName=n)[0] for n in (ONE, TWO)}
        for pan, name in PANS.values():
            ClientDetail.objects.get_or_create(user=self.user, PANNo=pan, defaults={"Name": name, "Group": self.groups[ONE]})
        self.reset()
        self.js_errors = []
        self.page.on("pageerror", lambda e: self.js_errors.append(f"{e.message}\n{e.stack}"))
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)

    def reset(self):
        ipos = list(self.ipos.values())
        OrderDetail.objects.filter(user=self.user, Order__OrderIPOName__in=ipos).delete()
        Order.objects.filter(user=self.user, OrderIPOName__in=ipos).delete()
        oracle.seed_specs(self.user, self.ipos["main"], MAIN_SPECS, PRE_OPEN)
        oracle.seed_specs(self.user, self.ipos["sme"], SME_SPECS, PRE_OPEN)
        for key in ("main", "sme"):
            calculate(self.ipos[key].id, self.user)
        Accounting.objects.filter(user=self.user, remark__startswith=REMARK).delete()
        now = timezone.now().replace(microsecond=0)
        mk = lambda **kw: Accounting.objects.create(user=self.user, date_time=now - timedelta(minutes=len(self.ledger)), **kw)
        self.ledger = []
        self.ledger.append(mk(ipo=self.ipos["main"], group=self.groups[ONE], amount=1500, amount_type="credit", remark=f"{REMARK} paid"))
        self.ledger.append(mk(ipo=self.ipos["sme"], group=self.groups[TWO], amount="250.50", amount_type="debit", remark=f"{REMARK} received"))
        self.ledger.append(mk(ipo=None, group=self.groups[ONE], amount=75, amount_type="debit", remark=f"{REMARK} jv", jv=True))
        self.deleted = mk(ipo=self.ipos["main"], group=self.groups[TWO], amount=999, amount_type="credit", remark=f"{REMARK} removed",
                          is_deleted=True, deleted_at=now)
        self.log_entry = AccountingAuditLog.objects.create(user=self.user, accounting=self.ledger[0], action="EDIT",
                                                           changes={"amount": {"old": "1000.00", "new": "1500.00"}})

    # ---------------------------------------------------------------------- helpers
    def open(self):
        # Below ~1000px the DataTables search box lies on top of "Accounting Backup" (headless Chromium is 800px wide),
        # so test at a laptop width; a maximized window is left alone.
        if self.page.evaluate("innerWidth") < 1200:
            self.page.set_viewport_size({"width": 1366, "height": 768})
        self.page.goto(f"{self.live_server_url}/BackUp")
        settle(self.page)
        expect(self.page.locator("h3.iponame")).to_have_text("BACK UP")
        self.page.wait_for_selector("#example tbody tr")

    def rows(self):
        return self.page.evaluate(ROWS_JS)

    def save(self, download):
        path = os.path.join(self.tmp.name, f"{len(os.listdir(self.tmp.name))}-{download.suggested_filename}")
        download.save_as(path)
        return path

    def download_backup_of(self, key):
        """Click the IPO's Backup button on the page; returns (suggested file name, workbook)."""
        ipo = self.ipos[key]
        row = self.page.locator("#example tbody tr", has=self.page.locator("th", has_text=re.compile(rf"^{re.escape(ipo.IPOName)}$")))
        with self.page.expect_download() as info:
            row.get_by_role("button", name="Backup").click()
        return info.value.suggested_filename, openpyxl.load_workbook(self.save(info.value))

    def set_page_size(self, value):
        with self.page.expect_navigation():
            self.page.select_option("#Backup_page_size", value)
        settle(self.page)

    def expected_orders(self, ipo):
        mainboard = ipo.IPOType == "MAINBOARD"
        out = Counter()
        for o in Order.objects.filter(user=self.user, OrderIPOName=ipo).select_related("OrderGroup"):
            row = [o.OrderGroup.GroupName, o.OrderType, o.OrderCategory] + ([o.InvestorType] if mainboard else []) + \
                  [db(o.Quantity), db(o.Rate), db(o.Amount), db(o.OrderDate), db(o.OrderTime)]
            out[tuple(row)] += 1
        return out

    def expected_details(self, ipo):
        mainboard = ipo.IPOType == "MAINBOARD"
        out = Counter()
        for od in OrderDetail.objects.filter(user=self.user, Order__OrderIPOName=ipo).select_related("Order__OrderGroup", "OrderDetailPANNo"):
            o, c = od.Order, od.OrderDetailPANNo
            row = [o.OrderGroup.GroupName, o.OrderCategory] + ([o.InvestorType] if mainboard else []) + \
                  [o.OrderType, db(o.Rate), db(od.AllotedQty), db(od.PreOpenPrice), db(od.Amount), c.PANNo if c else None, c.Name if c else None]
            out[tuple(row)] += 1
        for o in Order.objects.filter(user=self.user, OrderIPOName=ipo, OrderCategory="Premium").select_related("OrderGroup"):
            row = [o.OrderGroup.GroupName, o.OrderCategory] + ([o.InvestorType] if mainboard else []) + \
                  [o.OrderType, db(o.Rate), db(o.Quantity), db(ipo.PreOpenPrice), db(o.Amount), None, None]
            out[tuple(row)] += 1
        return out

    def assert_ipo_workbook(self, wb, key):
        ipo = self.ipos[key]
        ipo.refresh_from_db()
        self.assertEqual(wb.sheetnames, ["Sheet1", "Sheet2"])
        head, data = sheet(wb, "Sheet1")
        self.assertEqual(head, ORDER_HEAD[ipo.IPOType])
        self.assertEqual(Counter(tuple(r) for r in data), self.expected_orders(ipo), f"{ipo.IPOName}: Sheet1 (orders)")
        head, data = sheet(wb, "Sheet2")
        has_details = OrderDetail.objects.filter(user=self.user, Order__OrderIPOName=ipo).exists()
        want_head = DETAIL_HEAD[ipo.IPOType] if has_details else DETAIL_HEAD[ipo.IPOType][:-2]
        self.assertEqual(head, want_head)
        width = len(want_head)
        self.assertEqual(Counter(tuple((r + [None] * width)[:width]) for r in data),
                         Counter(tuple(r[:width]) for r in self.expected_details(ipo).elements()), f"{ipo.IPOName}: Sheet2 (applications)")
        return data

    # ====================================================================== 1. the page
    def test_the_page_lists_every_ipo_newest_first_with_a_backup_button(self):
        self.open()
        everything = list(CurrentIpoName.objects.filter(user=self.user).order_by("-id"))
        self.assertLessEqual(len(everything), 50, "this check expects one page")
        rows = self.rows()
        self.assertEqual([r[1] for r in rows], [i.IPOName for i in everything])
        self.assertEqual([r[0] for r in rows], [str(n) for n in range(1, len(everything) + 1)])
        for key, ipo in self.ipos.items():
            button = self.page.locator("#example tbody tr", has=self.page.locator("th", has_text=re.compile(rf"^{ipo.IPOName}$"))) \
                .get_by_role("button", name="Backup")
            expect(button).to_have_attribute("onclick", re.compile(rf"/{ipo.id}/Backup/"))
        expect(self.page.get_by_role("button", name=re.compile("All IPOs Backup"))).to_be_visible()
        expect(self.page.get_by_role("link", name=re.compile("Accounting Backup"))).to_have_attribute("href", "/AccountingBackup/")
        expect(self.page.get_by_text(f"Showing 1 to {len(everything)} of {len(everything)} Rows")).to_be_visible()
        self.assertEqual(self.js_errors, [])

    def test_the_search_box_narrows_the_list(self):
        self.open()
        self.page.locator("#example_filter input").fill("BKP SME")
        self.page.wait_for_timeout(300)
        self.assertEqual([r[1] for r in self.rows()], [SME_IPO])
        self.page.locator("#example_filter input").fill("")
        self.page.wait_for_timeout(300)
        self.assertGreaterEqual(len(self.rows()), 3)

    def test_only_a_logged_in_broker_can_back_up(self):
        anonymous = Client()
        for url in ("/BackUp", f"/{self.ipos['main'].id}/Backup/", "/AccountingBackup/"):
            with self.subTest(url=url):
                response = anonymous.get(url)
                self.assertEqual(response.status_code, 302)
                self.assertIn("login", response["Location"].lower())
        # another broker cannot download this broker's IPO
        other, _ = get_user_model().objects.get_or_create(username="bkp_other_broker", defaults={"email": "bkp@example.com"})
        other.groups.add(Group.objects.get(name="Broker"))
        client = Client(raise_request_exception=False)
        client.force_login(other)
        response = client.get(f"/{self.ipos['main'].id}/Backup/")
        self.assertNotEqual(response.status_code, 200)
        self.assertNotIn(b"BKP ONE", response.content)

    # ====================================================================== 2. one IPO
    def test_the_backup_of_a_mainboard_ipo_holds_every_order_and_application(self):
        self.open()
        name, wb = self.download_backup_of("main")
        self.assertRegex(name, rf"^{MAIN_IPO} \d{{2}}-\d{{2}}-\d{{4}}- \d{{2}}-\d{{2}} ?\.xlsx$")
        data = self.assert_ipo_workbook(wb, "main")
        # the amounts are the right billing figures, not just copied blanks
        ipo = self.ipos["main"]
        details = OrderDetail.objects.filter(user=self.user, Order__OrderIPOName=ipo).select_related("Order", "OrderDetailPANNo")
        want = Counter((d.OrderDetailPANNo.PANNo if d.OrderDetailPANNo else None, round(oracle.detail_amount(d, ipo), 2)) for d in details)
        want += Counter((None, round(oracle.order_amount(o, ipo), 2)) for o in Order.objects.filter(user=self.user, OrderIPOName=ipo, OrderCategory="Premium"))
        self.assertEqual(Counter((r[8], r[7] or 0.0) for r in data), want)
        # PANs and names arrive with their application
        self.assertIn(("BKPAA1234A", "BKP CLIENT A"), {(r[8], r[9]) for r in data})
        self.assertEqual(self.js_errors, [])

    def test_the_backup_of_an_sme_ipo_has_no_investor_type_column(self):
        self.open()
        name, wb = self.download_backup_of("sme")
        self.assertTrue(name.startswith(SME_IPO) and name.endswith(".xlsx"), name)
        self.assertNotIn("Investor Type", sheet(wb, "Sheet1")[0])
        self.assertNotIn("Investor Type", sheet(wb, "Sheet2")[0])
        self.assertEqual(sheet(wb, "Sheet2")[0][6], "Amount Diff.")
        self.assert_ipo_workbook(wb, "sme")

    def test_the_backup_of_an_ipo_without_orders_has_only_headers(self):
        self.open()
        _, wb = self.download_backup_of("empty")
        head, data = sheet(wb, "Sheet1")
        self.assertEqual((head, data), (ORDER_HEAD["MAINBOARD"], []))
        head, data = sheet(wb, "Sheet2")
        self.assertEqual((head, data), (DETAIL_HEAD["MAINBOARD"][:-2], []))

    def test_a_backup_follows_later_changes(self):
        """A second backup is taken from the database as it is now, not cached."""
        o = Order.objects.filter(user=self.user, OrderIPOName=self.ipos["main"], OrderCategory="Premium", OrderType="BUY").first()
        Order.objects.filter(id=o.id).update(Rate=21)
        OrderDetail.objects.filter(user=self.user, Order__OrderIPOName=self.ipos["main"], OrderDetailPANNo__PANNo="BKPAB1234B").update(AllotedQty=7)
        calculate(self.ipos["main"].id, self.user)
        self.open()
        _, wb = self.download_backup_of("main")
        data = self.assert_ipo_workbook(wb, "main")
        self.assertIn(7.0, [r[5] for r in data if r[8] == "BKPAB1234B"])

    def test_the_backup_files_are_not_left_on_the_server(self):
        before = set(os.listdir("."))
        self.open()
        self.download_backup_of("main")
        with self.page.expect_download():
            self.page.get_by_role("button", name=re.compile("All IPOs Backup")).click()
        self.assertEqual({f for f in set(os.listdir(".")) - before if f.endswith((".xlsx", ".zip")) or f == "path"}, set())

    @unittest.expectedFailure
    def test_an_ipo_name_with_a_slash_can_be_backed_up(self):
        """KNOWN BUG: the backup is first written to a file named after the IPO, so a "/" in the name crashes it (500)."""
        ipo = CurrentIpoName.objects.create(user=self.user, IPOName="BKP A/B IPO", IPOType="MAINBOARD", IPOPrice=IPO_PRICE,
                                            PreOpenPrice=PRE_OPEN, LotSizeRetail=100, LotSizeSHNI=1400, LotSizeBHNI=7000)
        try:
            response = self.page.request.get(f"{self.live_server_url}/{ipo.id}/Backup/")
            self.assertEqual(response.status, 200)
        finally:
            ipo.delete()

    @unittest.expectedFailure
    def test_an_ipo_name_is_shown_as_text_not_html(self):
        """KNOWN BUG: the list is built as raw HTML without escaping, so markup in an IPO name is rendered (stored XSS)."""
        ipo = CurrentIpoName.objects.create(user=self.user, IPOName="BKP <i>SLANTED</i> IPO", IPOType="MAINBOARD", IPOPrice=IPO_PRICE,
                                            PreOpenPrice=PRE_OPEN, LotSizeRetail=100, LotSizeSHNI=1400, LotSizeBHNI=7000)
        try:
            self.open()
            self.assertEqual(self.page.locator("#example tbody i", has_text="SLANTED").count(), 0)
            expect(self.page.locator("#example tbody").get_by_text("BKP <i>SLANTED</i> IPO")).to_be_visible()
        finally:
            ipo.delete()

    # ====================================================================== 3. all IPOs
    def download_all(self):
        self.open()
        with self.page.expect_download() as info:
            self.page.get_by_role("button", name=re.compile("All IPOs Backup")).click()
        return info.value.suggested_filename, zipfile.ZipFile(self.save(info.value))

    def test_the_all_ipos_backup_is_a_zip_with_one_workbook_per_ipo(self):
        name, archive = self.download_all()
        self.assertRegex(name, r"^backup_files-\d{2}-\d{2}-\d{4}- \d{2}-\d{2}\.zip$")
        workbooks = [n for n in archive.namelist() if n.endswith(".xlsx")]
        everything = CurrentIpoName.objects.filter(user=self.user)
        self.assertEqual(len(workbooks), everything.count())
        for ipo in everything:
            self.assertEqual(sum(n.startswith(f"{self.user}-{ipo.IPOName}-") for n in workbooks), 1, ipo.IPOName)
        for key in ("main", "sme", "empty"):
            with self.subTest(ipo=key):
                member = next(n for n in workbooks if n.startswith(f"{self.user}-{self.ipos[key].IPOName}-"))
                wb = openpyxl.load_workbook(io.BytesIO(archive.read(member)))
                if key == "empty":
                    self.assertEqual(sheet(wb, "Sheet1")[1], [])
                else:
                    self.assert_ipo_workbook(wb, key)

    @unittest.expectedFailure
    def test_the_all_ipos_zip_holds_only_the_workbooks(self):
        """KNOWN BUG: the zip is written inside the folder it zips, so it also packs a (partial) copy of itself."""
        _, archive = self.download_all()
        self.assertEqual([n for n in archive.namelist() if not n.endswith(".xlsx")], [])

    # ====================================================================== 4. accounting
    def download_accounting(self):
        self.open()
        with self.page.expect_download() as info:
            self.page.get_by_role("link", name=re.compile("Accounting Backup")).click()
        return info.value.suggested_filename, openpyxl.load_workbook(self.save(info.value))

    @staticmethod
    def ledger_row(e):
        e = Accounting.objects.get(id=e.id)                                    # the stored amount, e.g. 1500.00
        return [e.ipo.IPOName if e.ipo else "JV", e.group.GroupName if e.group else None, e.amount_type.upper(),
                f"{e.amount}", e.remark, timezone.localtime(e.date_time).strftime("%d-%m-%y %H:%M:%S")]

    def test_the_accounting_backup_holds_the_ledger_and_its_audit_log(self):
        name, wb = self.download_accounting()
        self.assertRegex(name, r"^Accounting-Backup \d{2}-\d{2}-\d{2}_\d{2}-\d{2}-\d{2}\.xlsx$")
        self.assertEqual(wb.sheetnames, ["Accounting", "Accounting Logs"])
        head, data = sheet(wb, "Accounting")
        self.assertEqual(head, ["IPO Name", "Group Name", "Amount Type", "Amount", "Remark", "Date Time"])
        mine = [r for r in data if r[4] and str(r[4]).startswith(REMARK) and r[4] != f"{REMARK} removed"]
        as_text = lambda r: [None if v is None else (f"{v:.2f}" if isinstance(v, float) else str(v)) for v in r]
        self.assertEqual(Counter(tuple(as_text(r)) for r in mine), Counter(tuple(self.ledger_row(e)) for e in self.ledger))
        self.assertEqual(len(data), Accounting.objects.filter(user=self.user).count(), "every ledger row of the broker")
        head, logs = sheet(wb, "Accounting Logs")
        self.assertEqual(head, ["Timestamp", "Transaction", "Action", "Changes", "Accounting ID"])
        mine = [r for r in logs if r[4] == float(self.ledger[0].id)]
        self.assertEqual(len(mine), 1)
        self.assertEqual(mine[0][1:4], [f"{MAIN_IPO} / {ONE}", "EDIT", "amount: 1000.00 -> 1500.00"])
        self.assertEqual(len(logs), AccountingAuditLog.objects.filter(user=self.user).count())

    def test_the_accounting_backup_names_a_deleted_ipo_or_group(self):
        ipo = CurrentIpoName.objects.create(user=self.user, IPOName="BKP GONE IPO", IPOType="MAINBOARD", IPOPrice=IPO_PRICE,
                                            PreOpenPrice=PRE_OPEN, LotSizeRetail=100, LotSizeSHNI=1400, LotSizeBHNI=7000)
        entry = Accounting.objects.create(user=self.user, ipo=ipo, ipo_name=ipo.IPOName, group=self.groups[TWO], amount=10,
                                          amount_type="credit", remark=f"{REMARK} gone", date_time=timezone.now())
        ipo.delete()
        entry.refresh_from_db()
        self.assertIsNone(entry.ipo)
        _, wb = self.download_accounting()
        row = next(r for r in sheet(wb, "Accounting")[1] if r[4] == f"{REMARK} gone")
        self.assertEqual(row[:2], ["BKP GONE IPO (Deleted)", TWO])

    @unittest.expectedFailure
    def test_a_deleted_ledger_entry_is_not_backed_up_as_a_live_one(self):
        """KNOWN BUG: soft-deleted entries are exported exactly like live ones - nothing in the file tells them apart."""
        _, wb = self.download_accounting()
        head, data = sheet(wb, "Accounting")
        removed = [r for r in data if r[4] == f"{REMARK} removed"]
        self.assertTrue(not removed or any("delet" in str(v).lower() for v in head + removed[0]))

    # ====================================================================== 5. page size and pages
    def fillers(self):
        count = max(0, 56 - CurrentIpoName.objects.filter(user=self.user).count())
        made = [CurrentIpoName.objects.create(user=self.user, IPOName=f"BKP FILL {i:02d}", IPOType="SME", IPOPrice=IPO_PRICE,
                                              PreOpenPrice=PRE_OPEN, LotSizeRetail=100, LotSizeSHNI=1400, LotSizeBHNI=7000)
                for i in range(count)]
        self.addCleanup(lambda: CurrentIpoName.objects.filter(id__in=[i.id for i in made]).delete())
        return CurrentIpoName.objects.filter(user=self.user).count()

    def test_pages_of_fifty_and_the_page_links(self):
        total = self.fillers()
        self.addCleanup(lambda: self.client.post("/BackUp", {"Backup_page_size": "50"}))
        self.open()
        self.set_page_size("50")
        rows = self.rows()
        self.assertEqual(len(rows), 50)
        expect(self.page.get_by_text(f"Showing 1 to 50 of {total} Rows")).to_be_visible()
        with self.page.expect_navigation():
            self.page.locator('a.page-link[href="/BackUp?page=2"]').first.click()
        settle(self.page)
        rows = self.rows()
        self.assertEqual([r[0] for r in rows], [str(n) for n in range(51, total + 1)])
        newest_first = list(CurrentIpoName.objects.filter(user=self.user).order_by("-id").values_list("IPOName", flat=True))
        self.assertEqual([r[1] for r in rows], newest_first[50:])

    def test_a_bigger_page_size_sends_every_ipo_and_is_remembered(self):
        total = self.fillers()
        self.addCleanup(lambda: self.client.post("/BackUp", {"Backup_page_size": "50"}))
        self.open()
        for size in ("100", "All"):
            with self.subTest(size=size):
                self.set_page_size(size)
                expect(self.page.get_by_text(f"Showing 1 to {total} of {total} Rows")).to_be_visible()
                self.assertEqual(self.page.input_value("#Backup_page_size"), size)
                self.assertEqual(self.page.evaluate("$('#example').DataTable().rows().count()"), total, "every IPO is in the page")
                self.open()                                                    # a plain visit keeps the size
                expect(self.page.get_by_text(f"Showing 1 to {total} of {total} Rows")).to_be_visible()
        self.assertEqual(self.js_errors, [])

    @unittest.expectedFailure
    def test_a_bigger_page_size_shows_every_ipo_on_screen(self):
        """KNOWN BUG: the table length is parseInt(size, 50) (radix 50 = NaN), so it stays at 50 rows and, with its pager
        removed, IPOs 51+ of a 100 / All page can never be reached."""
        total = self.fillers()
        self.addCleanup(lambda: self.client.post("/BackUp", {"Backup_page_size": "50"}))
        self.open()
        self.set_page_size("100")
        self.assertEqual(len(self.rows()), total)


