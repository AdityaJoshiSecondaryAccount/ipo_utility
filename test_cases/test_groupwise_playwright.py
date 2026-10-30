"""
Group-Wise Dashboard E2E (/GroupWiseDashboard, templates/GroupWiseDashboard.html).

Runs on its OWN IPOs and groups (GWD ...) and re-seeds their orders / ledger entries before every test,
so nothing it does leaks into the other pages. The dashboard lists EVERY group and IPO of the broker,
so (like the other pages) the checks are made on this file's groups by name, and anything that is about
the whole table (row counts, sort order) is compared with an independent recomputation over all groups.

What a row shows (home/views.py GroupWiseDashboard):
  IPO column      Total = sum of the group's orders in that IPO          Collection = net ledger (credit - debit)
                  Due = Total - Collection                               (IPOs with PreOpenPrice == IPOPrice count as 0)
  JV              minus the group's net JV ledger (credit - debit, jv=True)      -> shown with the sign flipped
  Total           sum of the IPO Totals
  Collection      "New Collection" = sum of the IPO Collections + net JV         (header text is just "Collection")
  Due Amount      Total - New Collection; the sign picks the due-positive / due-negative / due-zero class
Server side the same figures are annotated on the queryset (total_orders, total_jv, new_collection,
due_amount = Round(...)); they drive the zero-due filter and every sort.

Where the page differs from the wording of the task this file was written from (all tests assert the page):
  * the zero-due switch is "Show Zero Due Account": UNCHECKED hides zero-due groups, CHECKED shows them;
  * search / page size / sort / zero-due all reload the table with an AJAX GET (no form submit);
  * page-size options are 50 / 100 / 200 / All;
  * the Excel button is client-side SheetJS (static/JS/xlsx.full.min.js) built from the on-screen table; it
    never calls /download-Groupwise (that legacy endpoint is smoke-tested separately);
  * payments are made with the three modals users can reach: "Payment" (#transactionModal), "JV Payment"
    (#combinedTransferModal, a balanced transfer between two groups) and "Add Multiple IPO Payments" (double-click
    a Due Amount cell). #transactionModal1 / #transactionModal2 / #edit are in the DOM but nothing opens them, so
    they are not tested. All three modals refuse a payment of the wrong type or above the outstanding balance.
  * the entries those tests add are PIPELINE data (see test_cases/pipeline_state.py), kept for the Accounting page.
"""
import datetime
import io
import math
import os
import re
import tempfile
import unittest
from decimal import Decimal
from urllib.parse import parse_qs, urlparse

import openpyxl
from django.utils import timezone
from playwright.sync_api import expect

from home.models import Accounting, CurrentIpoName, GroupDetail, Order
from test_cases import pipeline_state
from test_cases.base import PersistentLiveServerTestCase, pause_for_inspection, settle

SETTLED_IPO, FLAT_IPO = "GWD SETTLED IPO", "GWD FLAT IPO"
ALPHA, BETA, DELTA = "GWD ALPHA TRADERS", "GWD BETA INVESTORS", "GWD DELTA CAPITAL"
PROBE = "GWD FLAT PROBE"
MAIN = (ALPHA, BETA, DELTA)
SEED_DATE, SEED_TIME = datetime.date(2026, 1, 20), datetime.time(10, 0, 0)

# Pipeline data for the payment modals (Section 5): created once per database and NEVER reset, because the
# entries the modals add are what the Accounting page has to show. orders are (ipo key, amount).
PAYFLOW_IPOS = {"one": "PAYFLOW IPO ONE", "two": "PAYFLOW IPO TWO"}
PAYFLOW_GROUPS = {
    "PAYFLOW SINGLE": [("one", 10000.0)],                      # owes 10,000 -> Add Payment
    "PAYFLOW MULTI": [("one", 6000.0), ("two", 4000.0)],       # owes 6,000 + 4,000 -> Add Multiple IPO Payments
    "PAYFLOW FROM": [("one", 7000.0)],                         # owes 7,000 -> pays ...
    "PAYFLOW TO": [("two", -5000.0)],                          # ... to a group the broker owes 5,000
    "PAYFLOW WITH JV": [("one", 5000.0)],                      # owes 5,000 and holds a JV credit of 800
}
PAYFLOW_SEED_JV = {"PAYFLOW WITH JV": ("credit", "800.00")}    # seeded with the group, not created by a test

# group -> orders (settled IPO, flat IPO), ledger entries (type, amount, ipo, jv)
#   ALPHA  owes the broker:      orders 15000.60, collected 4000.25 - 500.00, JV credit 1000       -> due  10500.35
#   BETA   is owed by the broker: orders -8000,   collected 1000                                   -> due  -9000
#   DELTA  is square:            orders 6000,    collected 6000                                    -> due      0
#   PROBE  only touches the flat IPO (its orders and ledger must be ignored)                       -> due      0
FIXTURE = {
    ALPHA: dict(settled=[12000.40, 3000.20], flat=[5000.0],
                ledger=[("credit", "4000.25", "settled", False), ("debit", "500.00", "settled", False),
                        ("credit", "1000.00", None, True)]),
    BETA: dict(settled=[-8000.0], flat=[-1500.0], ledger=[("credit", "1000.00", "settled", False)]),
    DELTA: dict(settled=[6000.0], flat=[700.0], ledger=[("credit", "6000.00", "settled", False)]),
    PROBE: dict(settled=[], flat=[300.0], ledger=[("credit", "250.00", "flat", False)]),
}

HEADERS_JS = """() => [...document.querySelectorAll('#example thead th')].map(th => ({
    text: th.textContent.replace(/\\s+/g, ' ').trim(),
    title: th.querySelector('span[title]')?.getAttribute('title') || '',
    visible: th.offsetParent !== null && getComputedStyle(th).display !== 'none',
    icon: th.querySelector('.sort-icon')?.className || ''}))"""

ROWS_JS = """() => [...document.querySelectorAll('#example tbody tr')].map(tr => {
    const kids = [...tr.children], n = kids.length, ipo = {};
    tr.querySelectorAll('.amount-cell').forEach(td => {
        ipo[td.dataset.ipoName] = {
            total: td.querySelector('.cell-total').textContent.trim(),
            collection: td.querySelector('.cell-collection').textContent.trim(),
            due: td.querySelector('.cell-due').textContent.trim(),
            exactDue: td.dataset.exactDue};
    });
    return {name: kids[0].textContent.trim(), ipo,
            jvText: kids[n - 4].textContent.trim(), jvAttr: kids[n - 4].dataset.jvAmount,
            total: kids[n - 3].textContent.trim(), newColl: kids[n - 2].textContent.trim(),
            dueText: kids[n - 1].querySelector('span').textContent.trim(), dueAttr: kids[n - 1].dataset.dueAmount,
            dueClass: kids[n - 1].className,
            shown: kids.filter(k => getComputedStyle(k).display !== 'none').length};
})"""

FOOTER_JS = """() => [...document.querySelectorAll('#example tfoot tr')[0].children]
    .map(c => ({text: c.textContent, visible: getComputedStyle(c).display !== 'none'}))"""

OPEN_MODAL_JS = """id => $(document.getElementById(id)).modal('show')"""

LOCAL_NOW_JS = """sel => { const e = document.querySelector(sel); e.removeAttribute('readonly');
    e.value = new Date(Date.now() - new Date().getTimezoneOffset() * 60000).toISOString().slice(0, 19); }"""


def fmt(x):
    """The page formats every amount with '{:.0f}'."""
    x = float(x)
    return "0" if round(x) == 0 else f"{x:.0f}"


def sql_round(x):
    """ROUND(x) as the database does it: half away from zero."""
    return math.copysign(math.floor(abs(x) + 0.5), x)


def indian(n):
    """toLocaleString('en-IN') for an integer."""
    n = int(n)
    sign, digits = ("-" if n < 0 else ""), str(abs(n))
    if len(digits) > 3:
        head, tail, parts = digits[:-3], digits[-3:], []
        while len(head) > 2:
            parts.insert(0, head[-2:])
            head = head[:-2]
        if head:
            parts.insert(0, head)
        digits = ",".join(parts) + "," + tail
    return sign + digits


def js_round(x):
    """Math.round: halves go towards +infinity."""
    return int(math.floor(x + 0.5))


def number(text):
    return float(re.search(r"-?\d+(?:\.\d+)?", text.replace(",", "")).group())


class GroupwisePlaywrightTestCase(PersistentLiveServerTestCase):
    PAGE_NAME = "groupwise"
    _ran_any = False

    # ------------------------------------------------------------------ setup / teardown
    @classmethod
    def tearDownClass(cls):
        cls.log_summary()
        if cls._ran_any:
            try:
                cls.page.goto(f"{cls.live_server_url}/GroupWiseDashboard?GWD_page_size=All&show_zero_due=true")
                settle(cls.page)
            except Exception:
                pass
            pause_for_inspection(cls.PAGE_NAME)
        super().tearDownClass()

    def setUp(self):
        super().setUp()
        type(self)._ran_any = True
        self.clear_local_storage()
        base = dict(IPOType="MAINBOARD", IPOPrice=100.0, LotSizeRetail=100, LotSizeSHNI=1400, LotSizeBHNI=7000,
                    Active=True)
        self.settled = CurrentIpoName.objects.get_or_create(
            user=self.user, IPOName=SETTLED_IPO, defaults=dict(base, PreOpenPrice=250.0))[0]
        self.flat = CurrentIpoName.objects.get_or_create(
            user=self.user, IPOName=FLAT_IPO, defaults=dict(base, PreOpenPrice=100.0))[0]
        self.assertNotEqual(self.settled.IPOPrice, self.settled.PreOpenPrice)
        self.assertEqual(self.flat.IPOPrice, self.flat.PreOpenPrice)
        self.reset_ledger()
        self.page.context.grant_permissions(["clipboard-read", "clipboard-write"], origin=self.live_server_url)

    def tearDown(self):
        self.clear_local_storage()
        super().tearDown()

    def clear_local_storage(self):
        try:
            if self.page.url.startswith(self.live_server_url):
                self.page.evaluate("localStorage.clear()")
        except Exception:
            pass

    def reset_ledger(self):
        """Canonical fixture: delete this file's orders / ledger entries and write them again."""
        self.groups = {name: GroupDetail.objects.get_or_create(user=self.user, GroupName=name)[0] for name in FIXTURE}
        mine = list(self.groups.values())
        Accounting.objects.filter(user=self.user, group__in=mine).delete()
        Order.objects.filter(user=self.user, OrderGroup__in=mine).delete()
        GroupDetail.objects.filter(pk__in=[g.pk for g in mine]).update(Collection=0)
        ipos = {"settled": self.settled, "flat": self.flat, None: None}
        for name, spec in FIXTURE.items():
            for key in ("settled", "flat"):
                for amount in spec[key]:
                    Order.objects.create(
                        user=self.user, OrderGroup=self.groups[name], OrderIPOName=ipos[key],
                        OrderType="BUY" if amount >= 0 else "SELL", Rate=1.0, Quantity=1.0,
                        OrderCategory="Premium", Amount=amount, OrderDate=SEED_DATE, OrderTime=SEED_TIME)
            for amount_type, amount, ipo_key, jv in spec["ledger"]:
                self.add_entry(self.groups[name], amount_type, amount, ipos[ipo_key], jv)

    def add_entry(self, group, amount_type, amount, ipo=None, jv=False):
        return Accounting.objects.create(
            user=self.user, group=group, ipo=ipo, amount=Decimal(str(amount)), amount_type=amount_type,
            remark="seed", date_time=timezone.now(), jv=jv)

    # -------------------------------------------------------- independent expectation
    def ledger(self):
        """
        {group name: figures} recomputed from the raw rows, straight from the rules in the task:
        IPOs with PreOpenPrice == IPOPrice are ignored (orders AND ledger), deleted entries are ignored.
        """
        ipos = list(CurrentIpoName.objects.filter(user=self.user))
        orders, acc, jv = {}, {}, {}
        for o in Order.objects.filter(user=self.user):
            orders[(o.OrderGroup_id, o.OrderIPOName_id)] = orders.get((o.OrderGroup_id, o.OrderIPOName_id), 0.0) + o.Amount
        for a in Accounting.objects.filter(user=self.user, is_deleted=False):
            signed = float(a.amount) if a.amount_type == "credit" else -float(a.amount)
            if a.jv:
                jv[a.group_id] = jv.get(a.group_id, 0.0) + signed
            if a.ipo_id:
                acc[(a.group_id, a.ipo_id)] = acc.get((a.group_id, a.ipo_id), 0.0) + signed
        out = {}
        for g in GroupDetail.objects.filter(user=self.user):
            cells, total, collected = {}, 0.0, 0.0
            for ipo in ipos:
                if ipo.IPOPrice == ipo.PreOpenPrice:
                    cells[ipo.IPOName] = dict(orders=0.0, collection=0.0, due=0.0, net=0.0)
                    continue
                o, c = orders.get((g.id, ipo.id), 0.0), acc.get((g.id, ipo.id), 0.0)
                cells[ipo.IPOName] = dict(orders=o, collection=c, due=o - c, net=sql_round(o - c))
                total += o
                collected += c
            j = jv.get(g.id, 0.0)
            out[g.GroupName] = dict(cells=cells, jv=j, total=total, new_collection=collected + j,
                                    due=total - (collected + j), due_rounded=sql_round(total - (collected + j)))
        return out

    def expected_group_count(self, show_zero, search=""):
        return sum(1 for name, fig in self.ledger().items()
                   if search.lower() in name.lower() and (show_zero or fig["due_rounded"] != 0))

    # ---------------------------------------------------------------------- page helpers
    def dashboard_url(self, show_zero=True, size="All", query=None):
        params = {"GWD_page_size": size, "show_zero_due": "true" if show_zero else "false"}
        if query:
            params["SearchQuery"] = query
        return f"{self.live_server_url}/GroupWiseDashboard?" + "&".join(f"{k}={v}" for k, v in params.items())

    def open_dashboard(self, show_zero=True, size="All", query=None):
        self.page.goto(self.dashboard_url(show_zero, size, query))
        settle(self.page)
        expect(self.page.locator("#example")).to_be_visible()
        expect(self.page.locator("#columnToggleMenu .col-toggle-btn").first).to_be_attached()

    def rows(self):
        return self.page.evaluate(ROWS_JS)

    def row(self, name):
        found = [r for r in self.rows() if r["name"] == name]
        self.assertEqual(len(found), 1, f"{name!r} should be listed exactly once, found {len(found)}")
        return found[0]

    def names(self):
        return [r["name"] for r in self.rows()]

    def headers(self):
        return self.page.evaluate(HEADERS_JS)

    def visible_headers(self):
        return [h["text"] for h in self.headers() if h["visible"]]

    def header_index(self, label):
        for i, h in enumerate(self.headers()):
            if h["title"] == label or (not h["title"] and h["text"] == label):
                return i
        self.fail(f"no column header {label!r} in {[h['text'] for h in self.headers()]}")

    def ajax_reload(self, action):
        """Run `action` (which makes the page re-load its table by AJAX) and return the request's query dict."""
        with self.page.expect_response(
                lambda r: "/GroupWiseDashboard" in r.url and r.request.resource_type in ("xhr", "fetch"),
                timeout=10000) as info:
            action()
        expect(self.page.locator("#example")).to_have_css("opacity", "1")
        return {k: v[0] for k, v in parse_qs(urlparse(info.value.url).query, keep_blank_values=True).items()}

    def switch(self, input_id):
        return self.page.locator(f"#{input_id} + .custom-switch-track")

    def click_header(self, label):
        index = self.header_index(label)
        return index, self.ajax_reload(
            lambda: self.page.locator("#example thead th").nth(index).click(position={"x": 6, "y": 6}))

    def footer_text(self):
        return self.page.inner_text("#paginationRow p").strip()

    def open_columns_menu(self):
        if not self.page.locator("#columnToggleMenu").is_visible():
            self.page.click("#columnsDropdown")
        expect(self.page.locator("#columnToggleMenu")).to_be_visible()

    def open_modal(self, modal_id):
        self.page.evaluate(OPEN_MODAL_JS, modal_id)
        expect(self.page.locator(f"#{modal_id}")).to_have_class(re.compile(r"\bshow\b"))

    # -------------------------------------------------------- row assertions
    def assert_row(self, name):
        """Every figure of `name`'s row against the independent ledger."""
        exp, row = self.ledger()[name], self.row(name)
        for ipo_name, cell in exp["cells"].items():
            shown = row["ipo"][ipo_name]
            with self.subTest(group=name, ipo=ipo_name):
                self.assertEqual(shown["total"], f"Total: {fmt(cell['orders'])}")
                self.assertEqual(shown["collection"], f"Collection: {fmt(cell['collection'])}")
                self.assertEqual(shown["due"], f"Due: {fmt(cell['due'])}")
                self.assertAlmostEqual(float(shown["exactDue"]), cell["due"], places=2)
        with self.subTest(group=name, columns="JV / Total / Collection / Due Amount"):
            self.assertAlmostEqual(float(row["jvAttr"]), -exp["jv"], places=2)
            self.assertEqual(row["jvText"], fmt(-exp["jv"]))
            self.assertEqual(row["total"], fmt(exp["total"]))
            self.assertEqual(row["newColl"], fmt(exp["new_collection"]))
            self.assertEqual(row["dueText"], fmt(exp["due"]))
            self.assertAlmostEqual(float(row["dueAttr"]), exp["due"], places=2)
            want = "due-positive" if round(exp["due"], 2) > 0 else "due-negative" if round(exp["due"], 2) < 0 else "due-zero"
            self.assertIn(want, row["dueClass"].split())
        return exp, row

    # ====================================================================================
    # Section 1: calculations
    # ====================================================================================
    def test_fixture_matches_the_hand_calculated_figures(self):
        """Guards the oracle itself: the figures in the FIXTURE comment, worked out by hand."""
        led = self.ledger()
        settled, flat = led[ALPHA]["cells"][SETTLED_IPO], led[ALPHA]["cells"][FLAT_IPO]
        self.assertAlmostEqual(settled["orders"], 15000.60, places=2)
        self.assertAlmostEqual(settled["collection"], 3500.25, places=2)
        self.assertAlmostEqual(settled["due"], 11500.35, places=2)
        self.assertEqual((flat["orders"], flat["collection"], flat["due"]), (0.0, 0.0, 0.0))
        self.assertAlmostEqual(led[ALPHA]["jv"], 1000.0, places=2)
        self.assertAlmostEqual(led[ALPHA]["new_collection"], 4500.25, places=2)
        self.assertAlmostEqual(led[ALPHA]["due"], 10500.35, places=2)
        self.assertEqual(led[ALPHA]["due_rounded"], 10500)
        self.assertEqual(led[BETA]["due_rounded"], -9000)
        self.assertEqual(led[DELTA]["due_rounded"], 0)
        self.assertEqual(led[PROBE]["due_rounded"], 0)

    def test_every_ipo_cell_is_orders_minus_collection(self):
        """Per IPO and group: Total = order sum, Collection = credit - debit, Due = Total - Collection."""
        self.open_dashboard()
        for name in MAIN:
            self.log(f"{name}", indent=2)
            exp, row = self.assert_row(name)
        alpha = self.row(ALPHA)["ipo"][SETTLED_IPO]
        self.assertEqual((alpha["total"], alpha["collection"], alpha["due"]),
                         ("Total: 15001", "Collection: 3500", "Due: 11500"))
        self.assertAlmostEqual(float(alpha["exactDue"]), 11500.35, places=2)   # the cell keeps the exact 2dp figure

    def test_total_jv_new_collection_and_due_columns(self):
        self.open_dashboard()
        row = self.row(ALPHA)
        self.assertEqual(row["jvText"], "-1000")                    # JV is displayed with the sign flipped
        self.assertAlmostEqual(float(row["jvAttr"]), -1000.0, places=2)
        self.assertEqual(row["total"], "15001")                     # 15000.60: the flat IPO's 5000 is not part of it
        self.assertEqual(row["newColl"], "4500")                    # 3500.25 + JV 1000
        self.assertEqual(row["dueText"], "10500")                   # 15000.60 - 4500.25 = 10500.35
        self.assertAlmostEqual(float(row["dueAttr"]), 10500.35, places=2)
        for name in (BETA, DELTA):
            self.assert_row(name)

    def test_due_sign_picks_the_status_class(self):
        """Positive: the group owes the broker; negative: the broker owes the group; zero: square."""
        self.open_dashboard()
        classes = {name: self.row(name)["dueClass"].split() for name in MAIN}
        self.assertIn("due-positive", classes[ALPHA])
        self.assertIn("due-negative", classes[BETA])
        self.assertIn("due-zero", classes[DELTA])
        self.assertEqual(self.row(BETA)["dueText"], "-9000")
        self.assertEqual(self.row(DELTA)["dueText"], "0")
        self.assertIn("zero-due-row", self.page.locator(f"#example tbody tr:has(th:text-is('{DELTA}'))").get_attribute("class"))

    def test_ipo_priced_at_pre_open_is_left_out_of_orders(self):
        """PreOpenPrice == IPOPrice: that IPO's orders count for nothing in its cell, Total and Due."""
        self.open_dashboard()
        for name in MAIN:
            cell = self.row(name)["ipo"][FLAT_IPO]
            self.assertEqual((cell["total"], cell["collection"], cell["due"]), ("Total: 0", "Collection: 0", "Due: 0"))
        # ALPHA has a 5000 order in the flat IPO; it is in neither Total nor Due
        self.assertEqual(self.row(ALPHA)["total"], "15001")
        # ... until that IPO's pre-open price differs from the IPO price, when it counts in full
        CurrentIpoName.objects.filter(pk=self.flat.pk).update(PreOpenPrice=120.0)
        try:
            self.open_dashboard()
            self.assertEqual(self.row(ALPHA)["ipo"][FLAT_IPO]["total"], "Total: 5000")
            self.assertEqual(self.row(ALPHA)["total"], "20001")
            self.assert_row(ALPHA)
        finally:
            CurrentIpoName.objects.filter(pk=self.flat.pk).update(PreOpenPrice=100.0)

    @unittest.expectedFailure
    def test_ledger_entries_of_a_pre_open_priced_ipo_are_left_out_of_the_row(self):
        """
        KNOWN APP BUG. The queryset annotation (zero-due filter, sorting) ignores ledger entries of an IPO
        priced at pre-open, but the row builder only zeroes the ORDERS: the flat IPO cell's Collection and the
        row's New Collection / Due Amount still include the entry. PROBE (flat order 300, flat credit 250)
        should be a zero row; the page shows New Collection 250 and Due -250 while filtering/sorting it as 0.
        """
        self.open_dashboard()
        row = self.row(PROBE)
        self.assertEqual(self.ledger()[PROBE]["due_rounded"], 0)
        self.assertEqual(row["ipo"][FLAT_IPO]["collection"], "Collection: 0")
        self.assertEqual(row["newColl"], "0")
        self.assertEqual(row["dueText"], "0")

    def test_server_annotations_match_the_independent_ledger(self):
        """total_orders / total_jv / new_collection / due_amount, read from the view's own context."""
        response = self.client.get("/GroupWiseDashboard", {"GWD_page_size": "All", "show_zero_due": "true"})
        self.assertEqual(response.status_code, 200)
        annotated = {g.GroupName: g for g in response.context["page_obj"]}
        led = self.ledger()
        for name in (*MAIN, PROBE):
            with self.subTest(group=name):
                g, exp = annotated[name], led[name]
                self.assertAlmostEqual(float(g.total_orders), exp["total"], places=2)
                self.assertAlmostEqual(float(g.total_jv), exp["jv"], places=2)
                self.assertAlmostEqual(float(g.new_collection), exp["new_collection"], places=2)
                self.assertEqual(float(g.due_amount), exp["due_rounded"])
        # deleted ledger rows no longer count
        gone = Accounting.objects.filter(user=self.user, group=self.groups[ALPHA], jv=True)
        gone.update(is_deleted=True)
        response = self.client.get("/GroupWiseDashboard", {"GWD_page_size": "All", "show_zero_due": "true"})
        alpha = {g.GroupName: g for g in response.context["page_obj"]}[ALPHA]
        self.assertEqual(float(alpha.total_jv), 0.0)
        self.assertEqual(float(alpha.due_amount), 11500.0)          # 15000.60 - 3500.25 = 11500.35 -> Round

    def test_new_ledger_rows_move_the_figures(self):
        """A fresh ledger entry shows up in the IPO cell, New Collection and Due Amount on the next load."""
        self.open_dashboard()
        before = self.row(BETA)
        self.add_entry(self.groups[BETA], "credit", "2000.50", self.settled)
        self.add_entry(self.groups[BETA], "debit", "300.25", None, jv=True)       # JV debit: -300.25
        self.open_dashboard()
        after = self.row(BETA)
        self.assertEqual(before["dueText"], "-9000")
        self.assertEqual(after["ipo"][SETTLED_IPO]["collection"], "Collection: 3000")   # 3000.50
        self.assertAlmostEqual(float(after["ipo"][SETTLED_IPO]["exactDue"]), -11000.50, places=2)
        self.assertAlmostEqual(float(after["jvAttr"]), 300.25, places=2)
        self.assertEqual(after["newColl"], "2700")                                    # 3000.50 - 300.25 = 2700.25
        self.assert_row(BETA)

    # ====================================================================================
    # Section 2: filters and toolbar
    # ====================================================================================
    def test_zero_due_switch_hides_and_restores_square_groups(self):
        self.step(1, 4, "default: zero-due groups are hidden")
        self.open_dashboard(show_zero=False)
        expect(self.page.locator("#toggleZeroDue")).not_to_be_checked()
        shown = self.names()
        self.assertIn(ALPHA, shown)
        self.assertIn(BETA, shown)
        self.assertNotIn(DELTA, shown)
        self.assertNotIn(PROBE, shown)
        self.assertEqual(len(shown), self.expected_group_count(False))
        self.assertEqual(self.page.locator("tr.zero-due-row").count(), 0)

        self.step(2, 4, "tick the switch: zero-due groups appear")
        query = self.ajax_reload(lambda: self.switch("toggleZeroDue").click())
        self.assertEqual((query["show_zero_due"], query["page"]), ("true", "1"))
        expect(self.page.locator("#toggleZeroDue")).to_be_checked()
        shown = self.names()
        self.assertIn(DELTA, shown)
        self.assertIn(PROBE, shown)
        self.assertEqual(len(shown), self.expected_group_count(True))
        self.assertIn(f"of {self.expected_group_count(True)} Rows", self.footer_text())
        self.assertEqual(self.row(DELTA)["dueText"], "0")

        self.step(3, 4, "the choice is remembered by the session")
        self.page.goto(f"{self.live_server_url}/GroupWiseDashboard")
        settle(self.page)
        expect(self.page.locator("#toggleZeroDue")).to_be_checked()
        self.assertIn(DELTA, self.names())

        self.step(4, 4, "untick: they disappear again")
        query = self.ajax_reload(lambda: self.switch("toggleZeroDue").click())
        self.assertEqual(query["show_zero_due"], "false")
        self.assertNotIn(DELTA, self.names())
        self.assertIn(f"of {self.expected_group_count(False)} Rows", self.footer_text())

    def test_a_group_that_settles_to_zero_leaves_the_default_view(self):
        self.open_dashboard(show_zero=False)
        self.assertIn(ALPHA, self.names())
        Accounting.objects.create(user=self.user, group=self.groups[ALPHA], ipo=self.settled, amount=Decimal("11500.35"),
                                  amount_type="credit", date_time=timezone.now(), jv=False)
        Accounting.objects.filter(user=self.user, group=self.groups[ALPHA], jv=True).update(is_deleted=True)
        self.assertEqual(self.ledger()[ALPHA]["due_rounded"], 0)
        self.open_dashboard(show_zero=False)
        self.assertNotIn(ALPHA, self.names())
        self.open_dashboard(show_zero=True)
        row = self.row(ALPHA)
        # 15000.60 - (3500.25 + 11500.35) leaves float dust, which the page prints as "-0"
        self.assertEqual(row["dueText"].lstrip("-"), "0")
        self.assertIn("due-zero", row["dueClass"].split())

    def test_search_filters_by_group_name(self):
        self.open_dashboard(show_zero=True)
        total = self.expected_group_count(True)
        self.assertEqual(len(self.names()), total)

        self.step(1, 3, "type ALPHA")
        field = self.page.locator("#searchInput")
        query = self.ajax_reload(lambda: field.fill("ALPHA"))
        self.assertEqual(query["SearchQuery"], "ALPHA")
        names = self.names()
        self.assertIn(ALPHA, names)
        self.assertNotIn(BETA, names)
        self.assertTrue(all("ALPHA" in n for n in names), names)
        self.assertEqual(len(names), self.expected_group_count(True, "ALPHA"))
        self.assertIn(f"Showing 1 to {len(names)} of {len(names)} Rows", self.footer_text())
        self.assertIn("SearchQuery=ALPHA", self.page.url)
        self.assert_row(ALPHA)

        self.step(2, 3, "the search is case-insensitive and partial")
        self.ajax_reload(lambda: field.fill("gwd beta"))
        self.assertEqual(self.names(), [BETA])
        self.assertIn("Showing 1 to 1 of 1 Rows", self.footer_text())
        self.ajax_reload(lambda: field.fill("NO SUCH GROUP"))
        self.assertEqual(self.names(), [])
        self.assertIn("of 0 Rows", self.footer_text())

        self.step(3, 3, "clear it: every group is back")
        query = self.ajax_reload(lambda: field.fill(""))
        self.assertNotIn("SearchQuery", query)
        self.assertEqual(len(self.names()), total)
        self.assertIn(f"of {total} Rows", self.footer_text())

    def test_search_keeps_the_zero_due_filter(self):
        self.open_dashboard(show_zero=False)
        self.ajax_reload(lambda: self.page.locator("#searchInput").fill("GWD"))
        self.assertEqual(sorted(n for n in self.names() if n.startswith("GWD")), sorted([ALPHA, BETA]))
        self.ajax_reload(lambda: self.switch("toggleZeroDue").click())
        self.assertEqual(sorted(n for n in self.names() if n.startswith("GWD")), sorted([ALPHA, BETA, DELTA, PROBE]))

    def test_page_size_selector(self):
        self.open_dashboard(show_zero=True, size="50")
        select = self.page.locator("#GWD_page_size")
        self.assertEqual(select.locator("option").all_inner_texts(), ["50", "100", "200", "All"])
        total = self.expected_group_count(True)

        for size in ("100", "200", "All", "50"):
            with self.subTest(size=size):
                query = self.ajax_reload(lambda: select.select_option(size))
                self.assertEqual((query["GWD_page_size"], query["page"]), (size, "1"))
                shown = len(self.names())
                limit = total if size == "All" else int(size)
                self.assertEqual(shown, min(limit, total))
                self.assertIn(f"Showing 1 to {shown} of {total} Rows", self.footer_text())

        self.step(1, 2, "100 is remembered by the session across a plain visit")
        self.ajax_reload(lambda: select.select_option("100"))
        self.page.goto(f"{self.live_server_url}/GroupWiseDashboard")
        settle(self.page)
        expect(self.page.locator("#GWD_page_size")).to_have_value("100")

        self.step(2, 2, "pagination links keep the page size, sort and filter")
        self.ajax_reload(lambda: self.page.locator("#GWD_page_size").select_option("50"))
        links = self.page.locator("#paginationRow a.page-link")
        if total > 50:
            self.assertGreater(links.count(), 0)
            self.ajax_reload(lambda: self.page.locator("#paginationRow a.page-link", has_text=re.compile(r"^2$")).first.click())
            self.assertIn(f"Showing 51 to {min(100, total)} of {total} Rows", self.footer_text())
        else:
            self.assertEqual(links.count(), 0, "a single page needs no page links")
            self.log("only one page of groups in the shared database; page 2 not exercised", indent=2)

    def test_toggle_all_ipos_hides_every_ipo_column(self):
        self.open_dashboard()
        ipos = [h["title"] for h in self.headers() if h["title"]]
        self.assertIn(SETTLED_IPO, ipos)
        self.assertIn(FLAT_IPO, ipos)
        self.assertEqual(self.visible_headers()[0], "Group Name")
        self.open_columns_menu()
        expect(self.page.locator("#toggleAllIposSwitch")).to_be_checked()

        self.step(1, 3, "switch all IPOs off: only the fixed columns remain")
        self.switch("toggleAllIposSwitch").click()
        expect(self.page.locator("#toggleAllIposSwitch")).not_to_be_checked()
        self.assertEqual(self.visible_headers(), ["Group Name", "JV", "Total", "Collection", "Due Amount"])
        self.assertEqual({r["shown"] for r in self.rows()}, {5})
        self.assertEqual(len([c for c in self.page.evaluate(FOOTER_JS) if c["visible"]]), 5)
        self.assertEqual(self.page.locator("#columnToggleMenu .col-toggle-btn.active-col").count(), 5)
        self.assertEqual(self.row(ALPHA)["dueText"], "10500")           # the figures are untouched, just hidden

        self.step(2, 3, "the choice is saved and survives a reload")
        self.assertEqual(self.page.evaluate("localStorage.getItem('gwd_toggle_all_ipos')"), "false")
        self.open_dashboard()
        self.assertEqual(self.visible_headers(), ["Group Name", "JV", "Total", "Collection", "Due Amount"])

        self.step(3, 3, "switch them back on")
        self.open_columns_menu()
        self.switch("toggleAllIposSwitch").click()
        expect(self.page.locator("#toggleAllIposSwitch")).to_be_checked()
        self.assertEqual(self.visible_headers(), [h["text"] for h in self.headers()])
        self.assertIn(SETTLED_IPO, self.visible_headers())

    def test_single_ipo_column_toggle(self):
        self.open_dashboard()
        all_headers = self.visible_headers()
        self.open_columns_menu()
        item = self.page.locator(f'#columnToggleMenu .col-toggle-btn[data-colname="{FLAT_IPO}"]')
        expect(item).to_have_class(re.compile("active-col"))

        item.click()                                        # hide only the flat IPO
        expect(item).not_to_have_class(re.compile("active-col"))
        self.assertEqual(self.visible_headers(), [h for h in all_headers if h != FLAT_IPO])
        self.assertIn(SETTLED_IPO, self.visible_headers())
        expect(self.page.locator("#toggleAllIposSwitch")).not_to_be_checked()      # no longer "all IPOs"
        self.assertEqual({r["shown"] for r in self.rows()}, {len(all_headers) - 1})

        self.step(1, 2, "a table reload (search) keeps it hidden")
        self.ajax_reload(lambda: self.page.locator("#searchInput").fill("GWD ALPHA"))
        self.assertEqual(self.visible_headers(), [h for h in all_headers if h != FLAT_IPO])
        self.assertEqual({r["shown"] for r in self.rows()}, {len(all_headers) - 1})

        self.step(2, 2, "show it again")
        self.open_columns_menu()
        item.click()
        expect(item).to_have_class(re.compile("active-col"))
        self.assertEqual(self.visible_headers(), all_headers)
        expect(self.page.locator("#toggleAllIposSwitch")).to_be_checked()

    # ====================================================================================
    # Section 3: sorting
    # ====================================================================================
    def assert_ordered(self, keys, ascending, what):
        pairs = list(zip(keys, keys[1:]))
        ok = all(a <= b for a, b in pairs) if ascending else all(a >= b for a, b in pairs)
        self.assertTrue(ok, f"{what} not {'ascending' if ascending else 'descending'}: {keys}")

    def keys_in_view(self, key_of):
        led = self.ledger()
        return [round(key_of(led[n]), 2) for n in self.names()]

    def sort_twice(self, label, key_of):
        """First click -> ascending, second -> descending; the rows must follow `key_of` each time."""
        self.open_dashboard(show_zero=True)
        index, query = self.click_header(label)
        self.assertEqual((query["sort_col"], query["sort_dir"]), (str(index), "asc"))
        expect(self.page.locator("#example thead th").nth(index).locator(".sort-icon")).to_have_class(re.compile("fa-sort-up"))
        asc = self.keys_in_view(key_of)
        self.assert_ordered(asc, True, f"{label} (1st click)")
        names_asc = self.names()

        index, query = self.click_header(label)
        self.assertEqual((query["sort_col"], query["sort_dir"]), (str(index), "desc"))
        expect(self.page.locator("#example thead th").nth(index).locator(".sort-icon")).to_have_class(re.compile("fa-sort-down"))
        self.assert_ordered(self.keys_in_view(key_of), False, f"{label} (2nd click)")
        self.assertEqual(len(self.names()), len(names_asc))

        index, query = self.click_header(label)             # and back to ascending
        self.assertEqual(query["sort_dir"], "asc")
        return names_asc

    def test_sort_by_group_name(self):
        """The unsorted table is already A-Z, so the first click on Group Name gives Z-A, the next A-Z again."""
        self.open_dashboard(show_zero=True)
        default = self.names()
        self.assertEqual(default, sorted(default, key=str.lower))
        index, query = self.click_header("Group Name")
        self.assertEqual((index, query["sort_col"], query["sort_dir"]), (0, "0", "desc"))
        self.assertEqual(self.names(), sorted(default, reverse=True))
        expect(self.page.locator("#example thead th").first.locator(".sort-icon")).to_have_class(re.compile("fa-sort-down"))
        _, query = self.click_header("Group Name")
        self.assertEqual((query["sort_col"], query["sort_dir"]), ("0", "asc"))
        self.assertEqual(self.names(), sorted(default))
        expect(self.page.locator("#example thead th").first.locator(".sort-icon")).to_have_class(re.compile("fa-sort-up"))

    def test_sort_by_an_ipo_column(self):
        names = self.sort_twice(SETTLED_IPO, lambda f: f["cells"][SETTLED_IPO]["net"])
        # ... which puts BETA (-9000 net) before DELTA (0) before ALPHA (11500)
        self.assertLess(names.index(BETA), names.index(DELTA))
        self.assertLess(names.index(DELTA), names.index(ALPHA))

    def test_sort_by_due_amount(self):
        """Negative, zero and positive dues, ordered numerically (not as text)."""
        names = self.sort_twice("Due Amount", lambda f: f["due_rounded"])
        self.assertLess(names.index(BETA), names.index(DELTA))          # -9000 < 0
        self.assertLess(names.index(DELTA), names.index(ALPHA))         # 0 < 10500
        self.open_dashboard(show_zero=True)
        self.click_header("Due Amount")
        self.click_header("Due Amount")                                 # descending now
        names = self.names()
        self.assertLess(names.index(ALPHA), names.index(DELTA))
        self.assertLess(names.index(DELTA), names.index(BETA))

    def test_sort_by_jv_total_and_collection(self):
        self.sort_twice("JV", lambda f: f["jv"])
        self.sort_twice("Total", lambda f: f["total"])
        self.sort_twice("Collection", lambda f: f["new_collection"])

    def test_sorting_survives_the_zero_due_switch_and_search(self):
        self.open_dashboard(show_zero=False)
        self.click_header("Due Amount")
        query = self.ajax_reload(lambda: self.switch("toggleZeroDue").click())
        self.assertEqual((query["sort_col"], query["sort_dir"], query["show_zero_due"]),
                         (str(self.header_index("Due Amount")), "asc", "true"))
        self.assert_ordered([f["due_rounded"] for f in (self.ledger()[n] for n in self.names())], True, "Due Amount")
        query = self.ajax_reload(lambda: self.page.locator("#searchInput").fill("GWD"))
        self.assertEqual(query["sort_col"], str(self.header_index("Due Amount")))
        names = self.names()
        self.assertTrue(all(n.startswith("GWD") for n in names), names)
        self.assertLess(names.index(BETA), names.index(DELTA))

    # ====================================================================================
    # Section 4: exports
    # ====================================================================================
    def export_xlsx(self):
        """Click Download -> Excel and return (filename, openpyxl worksheet rows)."""
        self.page.click("#downloadDropdown")
        with self.page.expect_download(timeout=20000) as info:
            self.page.click("#exportGwdExcelBtn")
        download = info.value
        path = os.path.join(tempfile.mkdtemp(prefix="gwd_export_"), download.suggested_filename)
        download.save_as(path)
        sheet = openpyxl.load_workbook(path).active
        return download.suggested_filename, sheet.title, [list(r) for r in sheet.iter_rows(values_only=True)]

    def test_excel_export_matches_the_table(self):
        self.open_dashboard(show_zero=True, size="50")
        self.ajax_reload(lambda: self.page.locator("#searchInput").fill("GWD"))
        requests = []
        handler = lambda r: requests.append(r.url)
        self.page.on("request", handler)
        self.assertEqual(self.page.evaluate("typeof XLSX"), "object", "static/JS/xlsx.full.min.js did not load")

        filename, sheet_name, grid = self.export_xlsx()
        self.assertRegex(filename, r"^GroupWiseDashboard_\d{4}-\d{2}-\d{2}\.xlsx$")
        self.assertEqual(sheet_name, "GroupWiseDashboard")
        # not 'All' rows: the button fetches the complete, filtered table first
        fetched = [u for u in requests if "/GroupWiseDashboard?" in u]
        self.assertTrue(fetched, "no full-table fetch was made")
        query = parse_qs(urlparse(fetched[-1]).query)
        self.assertEqual((query["GWD_page_size"], query["show_zero_due"], query["page"], query["SearchQuery"]),
                         (["All"], ["true"], ["1"], ["GWD"]))
        self.assertFalse([u for u in requests if "download-Groupwise" in u])
        self.page.remove_listener("request", handler)

        header = [h["text"] for h in self.headers() if h["visible"]]
        self.assertEqual(grid[0], header)
        body = {r[0]: r for r in grid[1:-1]}
        self.assertEqual(list(body), self.names_from_server(show_zero=True, search="GWD"))
        self.assertEqual(grid[-1][0], "Total")                          # footer row

        dom = {r["name"]: r for r in self.rows()}
        col = {(h["title"] or h["text"]): i for i, h in enumerate(x for x in self.headers() if x["visible"])}
        for name in MAIN:
            exp, shown, sheet = self.ledger()[name], dom[name], body[name]
            with self.subTest(group=name):
                for ipo_name, cell in exp["cells"].items():
                    self.assertEqual(sheet[col[ipo_name]], "\n".join(shown["ipo"][ipo_name][k] for k in ("total", "collection", "due")))
                self.assertEqual(float(sheet[col["JV"]]), float(fmt(-exp["jv"])))
                self.assertEqual(float(sheet[col["Total"]]), float(fmt(exp["total"])))
                self.assertEqual(float(sheet[col["Collection"]]), float(fmt(exp["new_collection"])))
                self.assertEqual(float(sheet[col["Due Amount"]]), float(fmt(exp["due"])))
                self.assertEqual(float(sheet[col["Due Amount"]]), number(shown["dueText"]))

        footer = self.page.evaluate(FOOTER_JS)
        for i, cell in enumerate(c for c in footer if c["visible"]):
            lines = [l.strip() for l in cell["text"].strip().splitlines() if l.strip()]
            expected = "\n".join(lines) if len(lines) > 1 else (float(lines[0].replace(",", "")) if re.fullmatch(r"-?[\d,.]+", lines[0]) else lines[0])
            self.assertEqual(grid[-1][i], expected, f"footer column {header[i]}")

    def names_from_server(self, show_zero, search=""):
        response = self.client.get("/GroupWiseDashboard", {"GWD_page_size": "All", "SearchQuery": search,
                                                           "show_zero_due": "true" if show_zero else "false"})
        return [g.GroupName for g in response.context["page_obj"]]

    def test_excel_export_follows_search_zero_due_and_hidden_columns(self):
        self.open_dashboard(show_zero=False, size="50")
        self.ajax_reload(lambda: self.page.locator("#searchInput").fill("GWD"))

        self.step(1, 2, "zero-due switch off + search GWD")
        requests = []
        handler = lambda r: requests.append(r.url)
        self.page.on("request", handler)
        _, _, grid = self.export_xlsx()
        self.page.remove_listener("request", handler)
        query = parse_qs(urlparse([u for u in requests if "/GroupWiseDashboard?" in u][-1]).query)
        self.assertEqual((query["SearchQuery"], query["show_zero_due"]), (["GWD"], ["false"]))
        exported = [r[0] for r in grid[1:-1]]
        self.assertEqual(sorted(exported), sorted([ALPHA, BETA]))        # DELTA and PROBE are square: left out
        self.assertEqual(exported, self.names_from_server(show_zero=False, search="GWD"))

        self.step(2, 2, "all IPO columns hidden: only the fixed columns are exported")
        self.open_columns_menu()
        self.switch("toggleAllIposSwitch").click()
        self.page.keyboard.press("Escape")
        _, _, grid = self.export_xlsx()
        self.assertEqual(grid[0], ["Group Name", "JV", "Total", "Collection", "Due Amount"])
        rows = {r[0]: r for r in grid[1:-1]}
        self.assertEqual(float(rows[ALPHA][4]), 10500.0)
        self.assertEqual(float(rows[BETA][4]), -9000.0)
        self.assertEqual(float(rows[ALPHA][1]), -1000.0)

    def test_legacy_download_groupwise_endpoint(self):
        """/download-Groupwise is no longer behind the button; it still answers with a workbook."""
        response = self.client.get("/download-Groupwise")
        self.assertEqual(response.status_code, 200)
        self.assertIn('filename="GroupWiseDashboard.xlsx"', response["Content-Disposition"])
        sheet = openpyxl.load_workbook(io.BytesIO(response.content)).active
        rows = [list(r) for r in sheet.iter_rows(values_only=True)]
        self.assertEqual(rows[0][-3:], ["Total", "Collection", "Due Amount"])
        self.assertIn(ALPHA, [r[0] for r in rows])
        self.assertEqual(rows[-1][0], "Total")

    def copy_breakdown(self, name):
        self.page.evaluate("navigator.clipboard.writeText('')")
        button = self.page.locator(f'.copy-group-due-btn[data-group-name="{name}"]')
        button.click()
        expect(button.locator("i")).to_have_class(re.compile("fa-check"))      # success tick
        return self.page.evaluate("navigator.clipboard.readText()")

    def expected_breakdown(self, name):
        """The Status.html-style block the button copies, built here from the ledger."""
        exp = self.ledger()[name]
        items = [(ipo_name, cell["due"]) for ipo_name, cell in exp["cells"].items() if round(cell["due"], 2) != 0]
        if round(exp["jv"], 2) != 0:
            items.append(("jv", -exp["jv"]))
        net = exp["due"]
        net = round(net, 2)
        items = [(label, round(v, 2)) for label, v in items]
        value_width = max([len(indian(js_round(v))) for _, v in items] + [len(indian(js_round(net)))])
        label_width = max([13] + [len(label) + 2 for label, _ in items])
        line = lambda label, text: label.ljust(label_width) + text.rjust(value_width)
        body = [line(label, indian(js_round(v))) for label, v in items] + [line("Net Amount", indian(js_round(net)))]
        return name + "\n```\n" + "\n".join(body) + "\n```"

    def test_copy_due_amount_breakdown(self):
        self.open_dashboard()
        for name in (ALPHA, BETA):
            with self.subTest(group=name):
                text = self.copy_breakdown(name)
                self.assertEqual(text, self.expected_breakdown(name))
        alpha = self.copy_breakdown(ALPHA)
        lines = alpha.splitlines()
        self.assertEqual(lines[0], ALPHA)
        self.assertEqual(lines[1], "```")
        self.assertEqual(lines[-1], "```")
        self.assertRegex(lines[2], rf"^{SETTLED_IPO}\s+11,500$")        # the IPO's due, 11500.35
        self.assertRegex(lines[3], r"^jv\s+-1,000$")                    # the JV balance as shown in the JV column
        self.assertRegex(lines[4], r"^Net Amount\s+10,500$")            # the row's Due Amount
        self.assertNotIn(FLAT_IPO, alpha)                               # IPOs with a zero due are left out
        self.assertEqual(self.copy_breakdown(BETA).splitlines()[-2].split()[-1], "-9,000")

    def test_copy_breakdown_does_not_open_the_payment_modal(self):
        self.open_dashboard()
        self.copy_breakdown(ALPHA)
        self.assertEqual(self.page.locator(".modal.show").count(), 0)

    # ====================================================================================
    # Section 5: the three payment modals (pipeline data, handed on to the Accounting page)
    # ====================================================================================
    # These tests ADD ledger entries to the PAYFLOW groups and never remove them: they are the entries the
    # Accounting page has to show. Each test records exactly what it created in test_cases/pipeline_state.py.
    def ensure_payflow(self):
        """PAYFLOW IPOs, groups and orders: created once per database, never reset (see pipeline_state)."""
        base = dict(IPOType="MAINBOARD", IPOPrice=100.0, PreOpenPrice=250.0, LotSizeRetail=100, LotSizeSHNI=1400,
                    LotSizeBHNI=7000, Active=True)
        self.pay_ipos = {key: CurrentIpoName.objects.get_or_create(user=self.user, IPOName=name, defaults=base)[0]
                         for key, name in PAYFLOW_IPOS.items()}
        self.pay_groups = {}
        for name, orders in PAYFLOW_GROUPS.items():
            group = GroupDetail.objects.get_or_create(user=self.user, GroupName=name)[0]
            self.pay_groups[name] = group
            if Order.objects.filter(user=self.user, OrderGroup=group).exists():
                continue
            for ipo_key, amount in orders:
                Order.objects.create(
                    user=self.user, OrderGroup=group, OrderIPOName=self.pay_ipos[ipo_key],
                    OrderType="BUY" if amount >= 0 else "SELL", Rate=1.0, Quantity=1.0, OrderCategory="Premium",
                    Amount=amount, OrderDate=SEED_DATE, OrderTime=SEED_TIME)
            if name in PAYFLOW_SEED_JV:
                self.add_entry(group, *PAYFLOW_SEED_JV[name], ipo=None, jv=True)

    def entries_since(self, mark):
        return list(Accounting.objects.filter(user=self.user, pk__gt=mark).order_by("pk"))

    def last_entry_id(self):
        last = Accounting.objects.filter(user=self.user).order_by("-pk").first()
        return last.pk if last else 0

    def describe(self, entry):
        """An Accounting row as the Accounting page will have to show it."""
        return dict(id=entry.pk, group=entry.group.GroupName, ipo=entry.ipo.IPOName if entry.ipo else None,
                    amount=str(entry.amount), type=entry.amount_type, jv=entry.jv, remark=entry.remark,
                    date_time=timezone.localtime(entry.date_time).strftime("%Y-%m-%d %H:%M:%S"),
                    batch=str(entry.transfer_batch_id) if entry.transfer_batch_id else None)

    def reload_after(self, action):
        """`action` ends in window.location.reload() or a redirect: wait for the new page."""
        with self.page.expect_navigation(timeout=15000):
            action()
        settle(self.page)
        expect(self.page.locator("#example")).to_be_visible()

    def banner(self, selector):
        return self.page.locator(selector).inner_text().strip()

    def test_add_payment_modal(self):
        """Toolbar 'Payment' -> #transactionModal -> add_transaction_group: credit 4,000 against a 10,000 due."""
        self.ensure_payflow()
        group, ipo = "PAYFLOW SINGLE", PAYFLOW_IPOS["one"]
        page, mark_text = self.page, pipeline_state.marker("groupwise")
        self.open_dashboard()
        self.assertEqual(self.row(group)["ipo"][ipo]["due"], "Due: 10000")
        before = self.last_entry_id()

        def fill(amount, amount_type, remark):
            page.select_option("#trans_group", label=group)
            page.select_option("#trans_ipo", label=ipo)
            page.select_option("#trans_amount_type", amount_type)
            page.fill("#trans_amount", amount)
            page.fill("#trans_remark", remark)
            page.evaluate(LOCAL_NOW_JS, "#add_date_time")
            # picking the group / IPO makes the modal fetch dues in the background; a late reply clears the banner
            settle(page)
            page.wait_for_timeout(800)

        page.click("button[data-bs-target='#transactionModal']")
        expect(page.locator("#transactionModal")).to_have_class(re.compile(r"\bshow\b"))

        self.step(1, 3, "a payment of the wrong type is refused (the group OWES, so it must be a Credit)")
        fill("100", "debit", "wrong type")
        page.click("#trans_save_btn")                                  # saved by AJAX: the error lands in the banner
        expect(page.locator("#trans_validation_message")).to_contain_text(f"{ipo} requires a Credit payment")
        self.assertEqual(self.entries_since(before), [])

        self.step(2, 3, "paying more than the outstanding balance is refused")
        fill("12000", "credit", "too much")
        page.click("#trans_save_btn")
        expect(page.locator("#trans_validation_message")).to_contain_text("cannot exceed its outstanding balance of ₹10000.00")
        self.assertEqual(self.entries_since(before), [])

        self.step(3, 3, "a valid payment is saved and shown")
        fill("4000", "credit", mark_text)
        self.reload_after(lambda: page.click("#trans_save_btn"))
        created = self.entries_since(before)
        self.assertEqual(len(created), 1)
        entry = created[0]
        self.assertEqual((entry.group.GroupName, entry.ipo.IPOName, entry.amount, entry.amount_type, entry.jv,
                          entry.is_deleted, entry.remark),
                         (group, ipo, Decimal("4000.00"), "credit", False, False, mark_text))
        self.assertIsNone(entry.transfer_batch_id)
        cell = self.row(group)["ipo"][ipo]
        self.assertEqual((cell["collection"], cell["due"]), ("Collection: 4000", "Due: 6000"))
        self.assertEqual(self.row(group)["dueText"], "6000")
        pipeline_state.record("groupwise", "add_payment", dict(
            modal="Add Payment", entries=[self.describe(entry)], group_due_after={group: 6000}))

    def test_jv_payment_between_groups_modal(self):
        """Toolbar 'JV Payment' -> #combinedTransferModal: PAYFLOW FROM (owes 7,000) pays 3,000 to PAYFLOW TO."""
        self.ensure_payflow()
        source, dest = "PAYFLOW FROM", "PAYFLOW TO"
        one, two = PAYFLOW_IPOS["one"], PAYFLOW_IPOS["two"]
        page, mark_text = self.page, pipeline_state.marker("groupwise")
        self.open_dashboard()
        self.assertEqual(self.row(source)["ipo"][one]["due"], "Due: 7000")
        self.assertEqual(self.row(dest)["ipo"][two]["due"], "Due: -5000")
        before = self.last_entry_id()

        page.click("button[data-bs-target='#combinedTransferModal']")
        expect(page.locator("#combinedTransferModal")).to_have_class(re.compile(r"\bshow\b"))
        save = page.locator("#ct_save")
        expect(save).to_be_disabled()                                  # nothing chosen yet

        self.step(1, 3, "pick the groups: the IPO allocations are worked out from their balances")
        page.select_option("#ct_source_group", label=source)
        page.select_option("#ct_destination_group", label=dest)
        page.select_option("#ct_source_type", "credit")
        expect(page.locator("#ct_destination_type")).to_have_value("Debit")     # the other side is the opposite
        page.fill("#ct_amount", "3000")
        page.fill("#ct_remark", mark_text)
        page.evaluate(LOCAL_NOW_JS, "#ct_date_time")
        source_rows = page.locator("#ct_source_allocations tr[data-ipo-id]")
        dest_rows = page.locator("#ct_destination_allocations tr[data-ipo-id]")
        expect(source_rows).to_have_count(1)
        expect(dest_rows).to_have_count(1)
        expect(source_rows.first.locator(".ct-allocation")).to_have_value("3000.00")
        expect(dest_rows.first.locator(".ct-allocation")).to_have_value("3000.00")
        self.assertEqual(source_rows.first.get_attribute("data-ipo-id"), str(self.pay_ipos["one"].id))
        self.assertEqual(dest_rows.first.get_attribute("data-ipo-id"), str(self.pay_ipos["two"].id))
        expect(page.locator("#ct_status")).to_have_text("Ready to save")
        expect(save).to_be_enabled()
        settle(page)
        page.wait_for_timeout(800)                                     # let both IPO lists finish re-drawing the rows
        expect(source_rows).to_have_count(1)
        expect(dest_rows).to_have_count(1)
        expect(source_rows.first.locator(".ct-allocation")).to_have_value("3000.00")
        expect(dest_rows.first.locator(".ct-allocation")).to_have_value("3000.00")

        self.step(2, 3, "save: one entry on each side, tied together by a transfer batch")
        self.reload_after(lambda: save.click())
        created = self.entries_since(before)
        shown = [self.describe(e) for e in created]
        self.assertEqual(len(created), 2, shown)
        src, dst = created
        self.assertTrue(src.ipo and dst.ipo, f"each side must be an IPO entry, got {shown}")
        self.assertEqual((src.group.GroupName, src.ipo.IPOName, src.amount, src.amount_type, src.jv),
                         (source, one, Decimal("3000.00"), "credit", False))
        self.assertEqual((dst.group.GroupName, dst.ipo.IPOName, dst.amount, dst.amount_type, dst.jv),
                         (dest, two, Decimal("3000.00"), "debit", False))
        self.assertIsNotNone(src.transfer_batch_id)
        self.assertEqual(src.transfer_batch_id, dst.transfer_batch_id)
        self.assertEqual(src.remark, f"Transfer to {dest} - {mark_text}")
        self.assertEqual(dst.remark, f"Transfer from {source} - {mark_text}")
        self.assertEqual(src.date_time, dst.date_time)

        self.step(3, 3, "both rows on the dashboard moved by the transfer")
        self.assertEqual(self.row(source)["ipo"][one]["due"], "Due: 4000")      # 7000 - 3000
        self.assertEqual(self.row(dest)["ipo"][two]["due"], "Due: -2000")       # -5000 + 3000
        self.assertEqual((self.row(source)["dueText"], self.row(dest)["dueText"]), ("4000", "-2000"))
        pipeline_state.record("groupwise", "jv_payment", dict(
            modal="JV Payment Between Groups", entries=[self.describe(src), self.describe(dst)],
            group_due_after={source: 4000, dest: -2000}))

    def test_add_multiple_ipo_payments_modal(self):
        """Double-click a Due Amount cell -> 'Add Multiple IPO Payments': PAYFLOW MULTI owes 6,000 + 4,000."""
        self.ensure_payflow()
        group = "PAYFLOW MULTI"
        one, two = PAYFLOW_IPOS["one"], PAYFLOW_IPOS["two"]
        page, mark_text = self.page, pipeline_state.marker("groupwise")
        self.open_dashboard()
        self.assertEqual(self.row(group)["dueText"], "10000")
        before = self.last_entry_id()

        page.locator("#example tbody tr").filter(has=page.locator(f"th:text-is('{group}')")) \
            .locator(".due-amount-cell span").dblclick()
        modal = page.locator("#multipleIPOModal")
        expect(modal).to_have_class(re.compile(r"\bshow\b"))
        rows = modal.locator(".ipo-payment-row")
        expect(rows).to_have_count(2)
        row_for = lambda ipo: modal.locator(f".ipo-payment-row[data-ipo-id='{self.pay_ipos[ipo].id}']")

        self.step(1, 3, "each IPO is pre-filled with its full due, as a Credit")
        expect(row_for("one").locator(".ipo-amount")).to_have_value("6000")
        expect(row_for("two").locator(".ipo-amount")).to_have_value("4000")
        expect(row_for("one").locator(".ipo-amount-type")).to_have_value("credit")
        expect(modal.locator(".total-credit")).to_have_text("Credit: 10000.00")
        expect(modal.locator(".Net-amount")).to_have_text("Net Amount: 10000.00")

        self.step(2, 3, "part-pay the first IPO, pay the second in full; the totals follow")
        row_for("one").locator(".ipo-amount").fill("2500")
        page.fill("#multi_remark", mark_text)
        expect(modal.locator(".total-credit")).to_have_text("Credit: 6500.00")
        expect(row_for("one").locator(".ipo-remark-preview")).to_have_text(mark_text)
        self.reload_after(lambda: page.click("#saveMultipleIPOPayments"))

        self.step(3, 3, "two entries, one per IPO, same time and remark")
        created = self.entries_since(before)
        self.assertEqual(len(created), 2)
        by_ipo = {e.ipo.IPOName: e for e in created}
        self.assertEqual(set(by_ipo), {one, two})
        for ipo_name, amount in ((one, "2500.00"), (two, "4000.00")):
            entry = by_ipo[ipo_name]
            self.assertEqual((entry.group.GroupName, entry.amount, entry.amount_type, entry.jv, entry.remark),
                             (group, Decimal(amount), "credit", False, mark_text))
        self.assertEqual(created[0].date_time, created[1].date_time)
        row = self.row(group)
        self.assertEqual((row["ipo"][one]["due"], row["ipo"][two]["due"]), ("Due: 3500", "Due: 0"))
        self.assertEqual(row["dueText"], "3500")
        pipeline_state.record("groupwise", "multi_ipo_payment", dict(
            modal="Add Multiple IPO Payments", entries=[self.describe(e) for e in created],
            group_due_after={group: 3500}))

    def test_transfer_needs_both_groups_and_an_amount_before_it_can_be_saved(self):
        self.ensure_payflow()
        self.open_dashboard()
        before = self.last_entry_id()
        self.page.click("button[data-bs-target='#combinedTransferModal']")
        save, status = self.page.locator("#ct_save"), self.page.locator("#ct_status")
        expect(save).to_be_disabled()
        self.page.select_option("#ct_source_group", label="PAYFLOW FROM")
        expect(status).to_have_text("Select the To Group.")
        self.page.select_option("#ct_destination_group", label="PAYFLOW TO")
        expect(status).to_have_text("Enter the amount to transfer.")
        expect(save).to_be_disabled()
        self.page.fill("#ct_amount", "100")
        expect(save).to_be_enabled()
        self.page.fill("#ct_amount", "")
        expect(save).to_be_disabled()
        self.assertEqual(self.entries_since(before), [])

    def test_a_remark_is_capped_at_250_characters_in_the_payment_modal(self):
        self.ensure_payflow()
        self.open_dashboard()
        self.page.click("button[data-bs-target='#transactionModal']")
        self.assertEqual(self.page.get_attribute("#trans_remark", "maxlength"), "250")
        self.page.fill("#trans_remark", "r" * 250)
        self.assertEqual(len(self.page.input_value("#trans_remark")), 250)

    def test_toolbar_buttons_open_their_modals(self):
        self.open_dashboard()
        self.page.click("button[data-bs-target='#transactionModal']")
        expect(self.page.locator("#transactionModal")).to_have_class(re.compile(r"\bshow\b"))
        expect(self.page.locator("#trans_group")).to_be_visible()
        self.open_dashboard()
        self.page.click("button[data-bs-target='#combinedTransferModal']")
        expect(self.page.locator("#combinedTransferModal")).to_have_class(re.compile(r"\bshow\b"))

    def test_double_clicking_a_due_cell_prefills_the_payment_modal(self):
        self.open_dashboard()
        row = self.page.locator("#example tbody tr").filter(has=self.page.locator(f"th:text-is('{ALPHA}')"))
        row.locator(f".amount-cell[data-ipo-name='{SETTLED_IPO}'] .cell-due").dblclick()
        expect(self.page.locator("#transactionModal")).to_have_class(re.compile(r"\bshow\b"))
        expect(self.page.locator("#trans_amount")).to_have_value("11500.35")
        expect(self.page.locator("#trans_amount_type")).to_have_value("credit")
        self.assertEqual(self.page.eval_on_selector("#trans_group", "e => e.selectedOptions[0].text.trim()"), ALPHA)
        self.assertEqual(self.page.eval_on_selector("#trans_ipo", "e => e.selectedOptions[0].text.trim()"), SETTLED_IPO)

    # ------------------------------------------------------------------ Add Multiple IPO Payments, in depth
    def open_multi_modal(self, group):
        """Double-click the group's Due Amount cell and wait for the modal to be filled from the server."""
        self.page.locator("#example tbody tr").filter(has=self.page.locator(f"th:text-is('{group}')")) \
            .locator(".due-amount-cell span").dblclick()
        modal = self.page.locator("#multipleIPOModal")
        expect(modal).to_have_class(re.compile(r"\bshow\b"))
        expect(modal.locator(".totals-row")).to_be_visible()
        return modal

    def save_multi_and_read_alert(self):
        """Click 'Save All Payments' for a save that is refused: return the alert text."""
        seen = []
        self.page.once("dialog", lambda dialog: (seen.append(dialog.message), dialog.accept()))
        self.page.click("#saveMultipleIPOPayments")
        for _ in range(50):
            if seen:
                break
            self.page.wait_for_timeout(100)
        self.assertTrue(seen, "no alert was shown")
        return seen[0]

    def test_multiple_ipo_modal_shows_the_jv_entry_the_ipo_rows_and_live_totals(self):
        """ALPHA holds a JV credit of 1,000 and owes 11,500.35 on one IPO: the modal offers both."""
        self.open_dashboard()
        modal = self.open_multi_modal(ALPHA)
        expect(modal.locator("#multi_group_name")).to_have_value(ALPHA)

        self.step(1, 4, "the JV entry settles the JV balance: Debit 1,000 (the shown JV is -1000)")
        jv = modal.locator(".jv-entry-row")
        expect(jv.locator(".jv-amount")).to_have_value("1000")
        expect(jv.locator(".jv-amount-type")).to_have_value("debit")
        expect(jv.locator(".jv-include")).to_be_checked()
        expect(jv).to_contain_text("Current JV: -1000")

        self.step(2, 4, "the IPO row is pre-filled with its full due as a Credit")
        rows = modal.locator(".ipo-payment-row")
        expect(rows).to_have_count(1)                                    # only the IPO with a balance is offered
        ipo_row = modal.locator(f".ipo-payment-row[data-ipo-id='{self.settled.id}']")
        expect(ipo_row.locator(".ipo-amount")).to_have_value("11500.35")
        expect(ipo_row.locator(".ipo-amount-type")).to_have_value("credit")
        expect(ipo_row).to_contain_text("Original: ₹11500.35")
        expect(ipo_row.locator("input[readonly]").first).to_have_value(SETTLED_IPO)
        expect(modal.locator(".total-credit")).to_have_text("Credit: 11500.35")
        expect(modal.locator(".total-debit")).to_have_text("Debit: 1000.00")
        expect(modal.locator(".Net-amount")).to_have_text("Net Amount: 10500.35")
        expect(modal.locator(".total-amount")).to_have_value("10500.35")
        expect(modal.locator(".balance-display")).to_have_value("Credit Balance: 10500.35")

        self.step(3, 4, "Include boxes and amounts drive the totals and the balance message")
        jv.locator(".jv-include").uncheck()
        expect(modal.locator(".total-debit")).to_have_text("Debit: 0.00")
        expect(modal.locator(".balance-display")).to_have_value("Credit Balance: 11500.35")
        jv.locator(".jv-include").check()
        ipo_row.locator(".ipo-amount").fill("1000")
        expect(modal.locator(".Net-amount")).to_have_text("Net Amount: 0.00")
        expect(modal.locator(".balance-display")).to_have_value("Balanced")
        ipo_row.locator(".ipo-amount").fill("400")
        expect(modal.locator(".balance-display")).to_have_value("Debit Balance: 600.00")
        ipo_row.locator(".ipo-amount-type").select_option("debit")
        expect(modal.locator(".total-debit")).to_have_text("Debit: 1400.00")

        self.step(4, 4, "the common remark is copied into every row's remark")
        self.page.fill("#multi_remark", "common note")
        expect(jv.locator(".jv-remark-preview")).to_have_text("common note")
        expect(ipo_row.locator(".ipo-remark-preview")).to_have_text("common note")

    def test_multiple_ipo_modal_refuses_bad_payments_and_empty_selections(self):
        self.open_dashboard()
        before = self.last_entry_id()
        modal = self.open_multi_modal(ALPHA)
        ipo_row = modal.locator(f".ipo-payment-row[data-ipo-id='{self.settled.id}']")
        jv = modal.locator(".jv-entry-row")

        self.step(1, 3, "more than the IPO's outstanding balance")
        ipo_row.locator(".ipo-amount").fill("12000")
        jv.locator(".jv-include").uncheck()
        self.assertIn("cannot exceed its outstanding balance of ₹11500.35", self.save_multi_and_read_alert())

        self.step(2, 3, "the wrong type for that balance (the group owes: Credit)")
        ipo_row.locator(".ipo-amount").fill("100")
        ipo_row.locator(".ipo-amount-type").select_option("debit")
        self.assertIn(f"{SETTLED_IPO} requires a Credit payment", self.save_multi_and_read_alert())

        self.step(3, 3, "nothing included")
        ipo_row.locator(".ipo-include").uncheck()
        self.assertIn("select at least one entry", self.save_multi_and_read_alert())
        self.assertEqual(self.entries_since(before), [])                    # none of the refusals wrote anything
        expect(self.page.locator("#saveMultipleIPOPayments")).to_be_enabled()

    def test_multiple_ipo_modal_saves_the_jv_entry_with_the_ipo_payment(self):
        """PAYFLOW WITH JV: a JV debit of 800 clears its JV credit while 3,000 is paid on the IPO."""
        self.ensure_payflow()
        group, ipo = "PAYFLOW WITH JV", PAYFLOW_IPOS["one"]
        mark_text = pipeline_state.marker("groupwise")
        self.open_dashboard()
        row = self.row(group)
        self.assertEqual((row["jvText"], row["dueText"]), ("-800", "4200"))          # 5000 - JV credit 800
        before = self.last_entry_id()

        modal = self.open_multi_modal(group)
        expect(modal.locator(".jv-amount")).to_have_value("800")
        expect(modal.locator(".jv-amount-type")).to_have_value("debit")
        ipo_row = modal.locator(f".ipo-payment-row[data-ipo-id='{self.pay_ipos['one'].id}']")
        expect(ipo_row.locator(".ipo-amount")).to_have_value("5000")
        ipo_row.locator(".ipo-amount").fill("3000")
        self.page.fill("#multi_remark", mark_text)
        expect(modal.locator(".Net-amount")).to_have_text("Net Amount: 2200.00")      # 3000 credit - 800 debit
        self.reload_after(lambda: self.page.click("#saveMultipleIPOPayments"))

        created = self.entries_since(before)
        self.assertEqual(len(created), 2)
        by_kind = {("jv" if e.jv else "ipo"): e for e in created}
        jv_entry, ipo_entry = by_kind["jv"], by_kind["ipo"]
        self.assertEqual((jv_entry.group.GroupName, jv_entry.ipo, jv_entry.amount, jv_entry.amount_type, jv_entry.remark),
                         (group, None, Decimal("800.00"), "debit", mark_text))
        self.assertEqual((ipo_entry.group.GroupName, ipo_entry.ipo.IPOName, ipo_entry.amount, ipo_entry.amount_type,
                          ipo_entry.jv, ipo_entry.remark),
                         (group, ipo, Decimal("3000.00"), "credit", False, mark_text))
        self.assertEqual(jv_entry.date_time, ipo_entry.date_time)

        row = self.row(group)
        self.assertEqual(row["jvText"], "0")                                          # +800 credit, -800 debit
        self.assertEqual(row["ipo"][ipo]["collection"], "Collection: 3000")
        self.assertEqual((row["newColl"], row["dueText"]), ("3000", "2000"))
        pipeline_state.record("groupwise", "multi_ipo_payment_with_jv", dict(
            modal="Add Multiple IPO Payments (with JV entry)", entries=[self.describe(e) for e in created],
            group_due_after={group: 2000}))
