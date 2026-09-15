"""Standalone Playwright/Django calculation regression.

Run: .venv/bin/python tests/verify_calculations.py --autoinject
Add --group-reference /path/to/group.xlsx --client-reference /path/to/client.xlsx
to compare workbook expectations directly with Playwright-rendered table cells.
Add --reference-report /path/to/report.json for the browser comparisons.
--headed shows Chromium. --fixture /path/to/fixture.json replaces --autoinject;
--export-fixture /path/to/fixture.json saves the fully seeded dataset for reuse.

The supplied client workbook covers page 1 (50 rows), not the whole client ledger.
Group reference amounts are whole rupees; UI category amounts show tenths, so
reference billing checks allow half a rupee. The independent raw-data oracle
also verifies category billing exactly at the UI's one-decimal precision.
Missing original Method/strike metadata is not verified or inferred.

Run with the repository's Python, directly (no pytest or existing e2e helpers).
Only the launcher below configures Django; importing the parsing/oracle utilities
does not open a database. Source cells are data, never executable instructions.
"""
import argparse
import csv
import json
import logging
import re
import sys
import tempfile
import unittest
from collections import Counter, defaultdict
from datetime import date, datetime, time, timedelta
from decimal import Decimal, InvalidOperation, ROUND_HALF_EVEN
from pathlib import Path

LOG = logging.getLogger("billing-calculations")
ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = Path.home() / "Downloads/DEEPA JEWELLERS 15-09-2026- 16-18 .xlsx"
USERNAME = "calculation-test-broker"
PASSWORD = "calculation-test-only-password"
IPO_NAME = "DEEPA JEWELLERS"
INVESTORS = ("RETAIL", "SHNI", "BHNI")
ORDER_HEADERS = ("Group", "OrderType", "Order Category", "Investor Type", "Qty",
                 "Rate", "Amount", "Order Date", "Order Time")
DETAIL_HEADERS = ("Group", "Order Category", "Investor Type", "Order Type", "Rate",
                  "AllotedQty", "Pre-Open Price", "Amount")


def number(value):
    """Strict finite decimal parser: blanks/dashes/errors are not silently zero."""
    text = re.sub(r"[\s,₹]", "", str(value)).replace("−", "-")
    if text.startswith("(") and text.endswith(")"):
        text = "-" + text[1:-1]
    if not re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)", text):
        raise ValueError(f"Invalid numeric value: {value!r}")
    try:
        result = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(f"Invalid numeric value: {value!r}") from exc
    if not result.is_finite():
        raise ValueError(f"Non-finite numeric value: {value!r}")
    return result


def records(path, sheet, required):
    """Read every nonempty data row, preserving duplicates and source row numbers."""
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.reader(stream))
    elif path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            rows = list(workbook[sheet].values)
        finally:
            workbook.close()
    else:
        raise ValueError("Source must be .csv or .xlsx")
    if not rows:
        raise ValueError(f"Empty source: {path}")
    headers = [str(x).strip() if x is not None else "" for x in rows[0]]
    if len(set(headers)) != len(headers) or set(required) - set(headers):
        raise ValueError(f"Invalid headers in {sheet}: required {required}, got {headers}")
    result = []
    for line, values in enumerate(rows[1:], 2):
        if all(v is None or str(v).strip() == "" for v in values):
            continue
        if len(values) != len(headers):
            raise ValueError(f"{sheet} row {line}: incorrect number of columns")
        item = dict(zip(headers, values))
        item["_line"] = line
        result.append(item)
    return result


def parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(str(value).strip(), fmt).date()
        except ValueError:
            pass
    raise ValueError(f"Unsupported date: {value!r}")


def dimensions(row, side_column):
    group = str(row["Group"] or "").strip().upper()
    side = str(row[side_column]).strip().upper()
    categories = {c.lower(): c for c in ("Kostak", "Subject To", "Premium", "CALL", "PUT")}
    category = categories.get(str(row["Order Category"]).strip().lower())
    investor = str(row["Investor Type"]).strip().upper()
    if not group or side not in ("BUY", "SELL") or category is None:
        raise ValueError(f"Row {row['_line']}: invalid group/type/category")
    if investor not in (*INVESTORS, "PREMIUM", "OPTIONS"):
        raise ValueError(f"Row {row['_line']}: unsupported investor {investor!r}")
    if category in ("Kostak", "Subject To") and investor not in INVESTORS:
        raise ValueError(f"Row {row['_line']}: invalid investor for {category}")
    return group, category, investor, side


def parse_dataset(args):
    orders = []
    for raw in records(args.source, args.orders_sheet, ORDER_HEADERS):
        group, category, investor, side = dimensions(raw, "OrderType")
        qty = number(raw["Qty"])
        if qty < 0 or qty != qty.to_integral_value():
            raise ValueError(f"Row {raw['_line']}: Qty must be a nonnegative integer")
        timestamp = raw["Order Time"]
        timestamp = timestamp if isinstance(timestamp, time) else time.fromisoformat(str(timestamp).strip())
        strike = raw.get("Strike", raw.get("Method"))
        orders.append(dict(group=group, category=category, investor=investor, side=side,
                           qty=qty, rate=number(raw["Rate"]), amount=number(raw["Amount"]),
                           date=parse_date(raw["Order Date"]), time=timestamp,
                           method=str(strike).strip() if strike not in (None, "") else None))
    if len(orders) != args.expected_rows:
        raise ValueError(f"Expected {args.expected_rows} orders, found {len(orders)} excluding header. "
                         "The supplied Sheet1 contains 840 orders + 1 header. "
                         "Use --expected-rows 840 to explicitly accept that file; no rows are invented.")
    details_path = args.details_source or args.source
    if details_path.suffix.lower() == ".csv" and args.details_source is None:
        raise ValueError("CSV orders require --details-source with the allotment export; "
                         "missing allotments cannot be inferred as zero.")
    details = []
    for raw in records(details_path, args.details_sheet, DETAIL_HEADERS):
        group, category, investor, side = dimensions(raw, "Order Type")
        qty = None if raw["AllotedQty"] in (None, "") else number(raw["AllotedQty"])
        if qty is not None and (qty < 0 or qty != qty.to_integral_value()):
            raise ValueError(f"Row {raw['_line']}: AllotedQty must be a nonnegative integer")
        details.append(dict(group=group, category=category, investor=investor, side=side,
                            rate=number(raw["Rate"]), allotted=qty,
                            preopen=number(raw["Pre-Open Price"]), amount=number(raw["Amount"]),
                            pan=str(raw.get("PAN No") or "").strip().upper()))
    if len(details) != args.expected_details:
        raise ValueError(f"Expected {args.expected_details} allotment records, found {len(details)}")
    keys = {match_key(o) for o in orders}
    if any(match_key(d) not in keys for d in details):
        raise ValueError("Allotment record has no matching group/category/investor/side/rate order")
    LOG.info("Parsed %s orders, %s allotment records, %s groups", len(orders), len(details),
             len({o['group'] for o in orders}))
    return orders, details


def match_key(row):
    return tuple(row[k] for k in ("group", "category", "investor", "side", "rate"))


def oracle(orders, details):
    """Independent Python math only. Never reads models, view helpers or UI values."""
    totals = defaultdict(lambda: defaultdict(Decimal))
    options = defaultdict(lambda: defaultdict(Decimal))
    for order in orders:
        group, cat, inv = (order[k] for k in ("group", "category", "investor"))
        sign = 1 if order["side"] == "BUY" else -1
        # Premium/options are not split by investor in this application's UI.
        inv = inv if cat in ("Kostak", "Subject To") else "ALL"
        bucket = totals[group, cat, inv]
        bucket["Count"] += sign * order["qty"]
        bucket["Billing"] += order["amount"]  # Amount already carries its own sign.
        if cat in ("CALL", "PUT"):
            strike = order["method"] or "-"
            options[group, strike][cat] += order["amount"]
            options[group, strike]["Shares"] += sign * order["qty"]
    for detail in details:
        if detail["category"] not in ("Kostak", "Subject To"):
            continue
        bucket = totals[detail["group"], detail["category"], detail["investor"]]
        sign = 1 if detail["side"] == "BUY" else -1
        if detail["allotted"] not in (None, 0):
            bucket["Alloted"] += sign
            bucket["Allotted Shares"] += sign * detail["allotted"]
    return totals, options


def seed_database(orders, details):
    """Bulk insert every order and detail, atomically, only in the isolated test DB.

    Export has no parent order IDs. Details attach to the first matching order
    by all five exported dimensions. This preserves these aggregate calculations,
    but does not reconstruct production per-order/PAN relationships.
    """
    from django.conf import settings
    from django.contrib.auth.models import Group
    from django.db import connection, transaction
    from home.models import CustomUser, CurrentIpoName, GroupDetail, ClientDetail, Order, OrderDetail
    if str(connection.settings_dict["NAME"]) != settings.CALCULATION_DB:
        raise RuntimeError("Refusing to seed outside the calculation test database")
    with transaction.atomic():
        user = CustomUser.objects.create_user(USERNAME, password=PASSWORD,
                                             Expiry_Date=date.today() + timedelta(days=30))
        user.groups.add(Group.objects.create(name="Broker"))
        ipo = CurrentIpoName.objects.create(user=user, IPOType="MAINBOARD", IPOName=IPO_NAME,
                                            IPOPrice=100, PreOpenPrice=221,
                                            LotSizeRetail=1, LotSizeSHNI=1, LotSizeBHNI=1)
        groups = {name: GroupDetail.objects.create(user=user, GroupName=name)
                  for name in sorted({o["group"] for o in orders})}
        objects = [Order(user=user, OrderGroup=groups[o["group"]], OrderIPOName=ipo,
                         OrderType=o["side"], OrderCategory=o["category"], InvestorType=o["investor"],
                         Quantity=float(o["qty"]), Rate=float(o["rate"]), Amount=float(o["amount"]),
                         OrderDate=o["date"], OrderTime=o["time"], Method=o["method"])
                   for o in orders]
        Order.objects.bulk_create(objects, batch_size=200)
        parents = {}
        for source, obj in zip(orders, objects):
            parents.setdefault(match_key(source), obj)
        clients = {}
        for detail in details:
            key = detail["group"], detail["pan"]
            if detail["pan"] and key not in clients:
                clients[key] = ClientDetail.objects.create(user=user, Group=groups[detail["group"]],
                                                          PANNo=detail["pan"])
        OrderDetail.objects.bulk_create([
            OrderDetail(user=user, Order=parents[match_key(d)],
                        OrderDetailPANNo=clients.get((d["group"], d["pan"])),
                        AllotedQty=None if d["allotted"] is None else float(d["allotted"]),
                        PreOpenPrice=float(d["preopen"]), Amount=float(d["amount"]))
            for d in details], batch_size=400)
    LOG.info("Bulk inserted ALL %s orders and %s details", len(orders), len(details))


def load_references(args):
    """Read expected values only. Actual values must come from Playwright DOM reads."""
    from openpyxl import load_workbook

    def read(path):
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            return list(workbook.active.values)
        finally:
            workbook.close()

    groups, clients, client_total = {}, [], None
    if args.group_reference:
        rows = read(args.group_reference)
        columns = [(cat, inv, metric) for cat in ("Kostak", "Subject To")
                   for inv in INVESTORS for metric in ("Count", "Alloted", "Billing")]
        columns += [("Premium", "ALL", "Shares"), ("Premium", "ALL", "Billing"),
                    ("CALL", "ALL", "Call Amount"), ("PUT", "ALL", "Put Amount"),
                    ("Total", "ALL", "Shares"), ("Total", "ALL", "Amount")]
        if [str(x).strip() for x in rows[1]] != ["Tally", "Group Name"] + [c[2] for c in columns]:
            raise ValueError("Unexpected flattened group-reference headers")
        for line, row in enumerate(rows[2:], 3):
            if len(row) != 26:
                raise ValueError(f"Group reference row {line}: expected 26 columns")
            group = str(row[1]).strip().upper()
            if group in groups:
                raise ValueError(f"Duplicate group reference: {group}")
            groups[group] = {key: number(value) for key, value in zip(columns, row[2:])}
    if args.client_reference:
        rows = read(args.client_reference)
        header = [None, "Group", "Order Category", "Premium Strike Price", "Investor Type",
                  "Order Type", "Rate", "PAN No", "Pre-Open Price", "Alloted Qty", "Amount"]
        if list(rows[1]) != header or rows[-1][0] != "Total":
            raise ValueError("Unexpected client-reference headers/footer")
        for line, row in enumerate(rows[2:-1], 3):
            if len(row) != len(header):
                raise ValueError(f"Client reference row {line}: invalid width")
            item = dict(zip(header[1:], row[1:]))
            for metric in ("Rate", "Pre-Open Price", "Alloted Qty", "Amount"):
                item[metric] = number(item[metric])
            item["_line"] = line
            clients.append(item)
        client_total = number(rows[-1][10])
    return groups, clients, client_total


# Build logical columns from header text and spans. CSS styling and column order
# may change without changing assertions. Read cells through Playwright locators.
HEADER_COLUMNS = """table => {
    const rows = [...table.tHead.rows];
    if (rows[0].cells.length === 1) rows.shift(); // modal category title
    const grid = rows.map(() => []);
    rows.forEach((row, r) => {
        let c = 0;
        [...row.cells].forEach(cell => {
            while (grid[r][c] !== undefined) c++;
            const label = cell.textContent.replace(/\\s+/g, ' ').trim().toLowerCase();
            for (let y=r; y<r+cell.rowSpan; y++) {
                if (!grid[y]) grid[y] = [];
                for (let x=c; x<c+cell.colSpan; x++) grid[y][x] = label;
            }
            c += cell.colSpan;
        });
    });
    return grid[0].map((_, c) => [...new Set(grid.map(row => row[c]).filter(Boolean))]);
}"""


def column_index(columns, *labels):
    wanted = [s.lower() for s in labels]
    matches = [i for i, path in enumerate(columns) if all(s in path for s in wanted)]
    if len(matches) != 1:
        raise AssertionError(f"Expected one column for {labels}; found {matches}: {columns}")
    return matches[0]


def make_case(args):
    from django.contrib.staticfiles.testing import StaticLiveServerTestCase
    from django.core.management import call_command
    from home.models import CustomUser, CurrentIpoName, GroupDetail, Order, OrderDetail
    from playwright.sync_api import sync_playwright, expect

    class CalculationVerification(StaticLiveServerTestCase):
        def setUp(self):
            # Python equivalent of beforeEach: parse, calculate, seed, verify completeness.
            self.orders, self.details = parse_dataset(args)
            self.expected, self.option_expected = oracle(self.orders, self.details)
            if args.autoinject:
                seed_database(self.orders, self.details)
            else:
                call_command("loaddata", str(args.fixture), verbosity=0)
            self.verify_seed()
            self.group_reference, self.client_reference, self.client_total = load_references(args)
            self.browser_reference_results = []
            self.addCleanup(self.save_reference_report)
            self.groups = list(GroupDetail.objects.filter(user=self.user).order_by("GroupName"))
            self.assertEqual({g.GroupName for g in self.groups}, {o["group"] for o in self.orders})
            if args.export_fixture:
                with args.export_fixture.open("w", encoding="utf-8") as stream:
                    call_command("dumpdata", "auth.group", "home.customuser", "home.currentiponame",
                                 "home.groupdetail", "home.clientdetail", "home.order", "home.orderdetail", stdout=stream)
            self.pw = sync_playwright().start()
            self.addCleanup(self.pw.stop)
            browser = self.pw.chromium.launch(headless=not args.headed)
            self.addCleanup(browser.close)
            self.page = browser.new_page(viewport={"width": 1600, "height": 1000})
            self.page.set_default_timeout(15000)
            LOG.info("Logging in to the temporary application")
            self.page.goto(self.live_server_url + "/login")
            self.page.get_by_placeholder("User Name", exact=True).fill(USERNAME)
            password = self.page.get_by_placeholder("Password", exact=True)
            password.fill(PASSWORD)
            password.press("Enter")
            self.page.wait_for_url(re.compile(r"^(?!.*\/login).+$"))
            LOG.info("Authenticated; starting calculation assertions")
            self.checked = 0

        def verify_seed(self):
            self.user = CustomUser.objects.get(username=USERNAME)
            self.ipo = CurrentIpoName.objects.get(user=self.user, IPOName=IPO_NAME)
            self.assertEqual(self.ipo.IPOType, "MAINBOARD")
            self.assertEqual(Order.objects.count(), len(self.orders), "Full order count differs")
            self.assertEqual(OrderDetail.objects.count(), len(self.details), "Full detail count differs")
            # Compare multisets, not just sums/counts: duplicated or omitted rows must fail.
            actual = Counter()
            for o in Order.objects.select_related("OrderGroup"):
                self.assertEqual((o.user_id, o.OrderIPOName_id), (self.user.pk, self.ipo.pk))
                self.assertNotIn(str(o.Telly).lower(), ("true", "1"))
                actual[(o.OrderGroup.GroupName, o.OrderCategory, o.InvestorType, o.OrderType,
                        number(o.Quantity), number(o.Rate), number(o.Amount),
                        o.OrderDate, o.OrderTime, o.Method)] += 1
            fields = ("group", "category", "investor", "side", "qty", "rate", "amount", "date", "time", "method")
            self.assertEqual(actual, Counter(tuple(o[k] for k in fields) for o in self.orders),
                             "Seeded order multiset differs from source")
            actual_details = Counter()
            for d in OrderDetail.objects.select_related("Order__OrderGroup", "OrderDetailPANNo"):
                self.assertEqual(d.user_id, self.user.pk)
                o = d.Order
                actual_details[(o.OrderGroup.GroupName, o.OrderCategory, o.InvestorType, o.OrderType,
                                number(o.Rate), None if d.AllotedQty is None else number(d.AllotedQty),
                                number(d.PreOpenPrice), number(d.Amount),
                                d.OrderDetailPANNo.PANNo if d.OrderDetailPANNo else "")] += 1
            fields = ("group", "category", "investor", "side", "rate", "allotted", "preopen", "amount", "pan")
            self.assertEqual(actual_details, Counter(tuple(d[k] for k in fields) for d in self.details),
                             "Seeded allotment multiset differs from source")

        def save_reference_report(self):
            if not args.reference_report:
                return
            expected_count = len(self.group_reference) * 24 + (len(self.client_reference) * 4 + 1
                                                               if args.client_reference else 0)
            complete = len(self.browser_reference_results) == expected_count
            report = dict(actual_value_source="Playwright rendered DOM", complete=complete,
                          passed=complete and all(r["passed"] for r in self.browser_reference_results),
                          expected_numeric_checks=expected_count,
                          checks=self.browser_reference_results)
            args.reference_report.parent.mkdir(parents=True, exist_ok=True)
            args.reference_report.write_text(json.dumps(report, indent=2), encoding="utf-8")

        def check_reference_cell(self, cell, expected, label, tolerance=Decimal(0)):
            """Expected comes from the workbook; actual comes only from Chromium."""
            with self.subTest(browser_reference=label):
                text = cell.inner_text()
                actual = number(text)
                result = dict(label=label, url=self.page.url, expected=str(expected),
                              ui_text=text, ui_value=str(actual), tolerance=str(tolerance),
                              passed=abs(actual - expected) <= tolerance)
                self.browser_reference_results.append(result)
                if not result["passed"]:
                    LOG.error("%s workbook=%s browser=%s difference=%s", label, expected,
                              actual, actual - expected)
                self.assertTrue(result["passed"], str(result))

        def verify_group_reference_in_browser(self, table, group):
            columns = table.evaluate(HEADER_COLUMNS)
            row = table.locator("tbody tr").filter(has=self.page.get_by_text(IPO_NAME, exact=True))
            expect(row).to_have_count(1)
            cells = row.locator(":scope > th, :scope > td")
            for (category, investor, metric), expected in self.group_reference[group].items():
                labels = ((category, investor, metric) if investor != "ALL" else
                          ("OPTIONS" if category in ("CALL", "PUT") else category, metric))
                # Reference export stores whole rupees; live category Billing
                # cells store tenths. Half-rupee precision is all this reference
                # can establish. The separate raw-data oracle still checks tenths exactly.
                tolerance = Decimal("0.5") if metric in ("Billing", "Call Amount", "Put Amount") else Decimal(0)
                self.check_reference_cell(cells.nth(column_index(columns, *labels)), expected,
                                          f"Group workbook: {group}/{category}/{investor}/{metric}", tolerance)

        def verify_client_reference_in_browser(self):
            # The supplied export is the first 50-row client-billing page.
            # Keep the entire source dataset seeded; do not replace it with this subset.
            response = self.page.goto(f"{self.live_server_url}/{self.ipo.pk}/Billing?page=1")
            self.assertEqual(response.status, 200)
            table = self.page.locator("#example")
            expect(table).to_be_visible()
            columns = table.evaluate(HEADER_COLUMNS)
            rows = table.locator("tbody > tr")
            expect(rows).to_have_count(len(self.client_reference))
            used = set()
            for item in self.client_reference:
                label = f"Client workbook row {item['_line']}: {item['Group']}/{item['Investor Type']}"
                matches = rows
                for key in ("Group", "PAN No", "Order Category", "Investor Type", "Order Type"):
                    matches = matches.filter(has=self.page.get_by_text(str(item[key]).strip(), exact=True))
                # Match identity + rate, never Amount or Alloted Qty; those are
                # the measurements under test. A DOM row cannot satisfy two references.
                matching_rows = []
                for i in range(matches.count()):
                    candidate = matches.nth(i)
                    row_id = candidate.locator(".billing-row-checkbox").get_attribute("value")
                    rate = number(candidate.locator(":scope > td").nth(column_index(columns, "Rate")).inner_text())
                    if rate == item["Rate"] and row_id not in used:
                        matching_rows.append((candidate, row_id))
                self.assertTrue(matching_rows, f"{label}: no unused matching browser row")
                row, row_id = matching_rows[0]
                used.add(row_id)
                cells = row.locator(":scope > td")
                for metric in ("Rate", "Pre-Open Price", "Alloted Qty", "Amount"):
                    self.check_reference_cell(cells.nth(column_index(columns, metric)), item[metric],
                                              f"{label}/{metric}")
            self.assertEqual(len(used), len(self.client_reference))
            # Read the actual displayed page footer, not a sum of workbook rows.
            footer = table.locator("tfoot > tr").first.locator(":scope > th, :scope > td")
            self.check_reference_cell(footer.nth(column_index(columns, "Amount")), self.client_total,
                                      "Client workbook: displayed page Amount total")
            LOG.info("Playwright verified %s client rows and their displayed footer against workbook",
                     len(self.client_reference))

        def check(self, cell, expected, group, category, investor, metric, decimals=0):
            label = f"Group={group} Category={category} Investor={investor} Metric={metric}"
            with self.subTest(calculation=label):
                actual = number(cell.inner_text())
                rounded = expected.quantize(Decimal(1).scaleb(-decimals), rounding=ROUND_HALF_EVEN)
                if actual != rounded:
                    LOG.error("%s UI=%s expected=%s raw=%s difference=%s", label, actual,
                              rounded, expected, actual - rounded)
                self.assertEqual(actual, rounded, f"{label}; raw={expected}; UI={actual}; expected={rounded}")
                self.checked += 1

        def verify_table(self, table, group, modal_category=None):
            expect(table).to_have_count(1)
            columns = table.evaluate(HEADER_COLUMNS)
            row = table.locator("tbody tr").filter(has=self.page.get_by_text(IPO_NAME, exact=True))
            expect(row).to_have_count(1)
            cells = row.locator(":scope > th, :scope > td")
            for category in ("Kostak", "Subject To"):
                if modal_category and modal_category != category:
                    continue
                for investor in INVESTORS:
                    for metric in ("Count", "Alloted", "Billing"):
                        labels = (investor, metric) if modal_category else (category, investor, metric)
                        cell = cells.nth(column_index(columns, *labels))
                        self.check(cell, self.expected[group, category, investor][metric], group,
                                   category, investor, metric, 1 if metric == "Billing" else 0)
            if modal_category in (None, "Premium"):
                for label, metric in (("Shares", "Count"), ("Billing", "Billing")):
                    self.check(cells.nth(column_index(columns, "Premium", label)),
                               self.expected[group, "Premium", "ALL"][metric], group,
                               "Premium", "ALL", label, 1 if label == "Billing" else 0)
            if modal_category is None:
                for cat, metric in (("CALL", "Call Amount"), ("PUT", "Put Amount")):
                    self.check(cells.nth(column_index(columns, "OPTIONS", metric)),
                               self.expected[group, cat, "ALL"]["Billing"], group, cat, "ALL", metric, 1)
                shares = self.expected[group, "Premium", "ALL"]["Count"] + sum(
                    self.expected[group, cat, inv]["Allotted Shares"]
                    for cat in ("Kostak", "Subject To") for inv in INVESTORS)
                amount = sum(o["amount"] for o in self.orders if o["group"] == group)
                for metric, value in (("Shares", shares), ("Amount", amount)):
                    self.check(cells.nth(column_index(columns, "Total", metric)), value,
                               group, "Total", "ALL", metric)

        def verify_options(self, table, group):
            # Modal can have several strike rows with a rowspan IPO-name header.
            expect(table).to_have_count(1)
            expected = {strike: values for (name, strike), values in self.option_expected.items()
                        if name == group and any(values.values())}
            if not expected:
                expected = {"-": {"CALL": Decimal(0), "PUT": Decimal(0), "Shares": Decimal(0)}}
            rows = table.locator("tbody tr")
            expect(rows).to_have_count(len(expected))
            seen = set()
            for i in range(rows.count()):
                # IPO is a th; all five numerical/strike columns are td regardless of rowspan.
                cells = rows.nth(i).locator(":scope > td")
                expect(cells).to_have_count(5)
                strike = cells.nth(0).inner_text().strip()
                self.assertIn(strike, expected, f"Unexpected strike for {group}: {strike}")
                self.assertNotIn(strike, seen)
                seen.add(strike)
                value = expected[strike]
                for index, metric, amount, places in (
                    (1, "Call Amount", value.get("CALL", Decimal(0)), 1),
                    (2, "Put Amount", value.get("PUT", Decimal(0)), 1),
                    (3, "Shares", value["Shares"], 0),
                    (4, "Amount", value.get("CALL", Decimal(0)) + value.get("PUT", Decimal(0)), 0),
                ):
                    self.check(cells.nth(index), amount, group, "Options", f"Strike={strike}", metric, places)

        def test_all_group_and_share_calculations(self):
            if args.group_reference:
                self.assertEqual(set(self.group_reference), {g.GroupName for g in self.groups})
            if args.client_reference:
                self.verify_client_reference_in_browser()
            # Direct group routes avoid All Groups' five-group lazy-load limit.
            for group in self.groups:
                with self.subTest(group=group.GroupName):
                    LOG.info("Checking %s", group.GroupName)
                    response = self.page.goto(f"{self.live_server_url}/group-billing-details/{group.pk}/")
                    self.assertEqual(response.status, 200)
                    main = self.page.locator(f"#mainboardBillingTable_{group.pk}")
                    self.verify_table(main, group.GroupName)
                    if args.group_reference:
                        self.verify_group_reference_in_browser(main, group.GroupName)
                    self.page.get_by_role("button", name="Share", exact=True).click()
                    expect(self.page.locator("#shareGroupModal")).to_be_visible()
                    empty_toggle = self.page.locator("#hideEmptyRowsToggle")
                    if empty_toggle.is_checked():
                        self.page.locator('label[for="hideEmptyRowsToggle"]').click()
                    expect(empty_toggle).not_to_be_checked()
                    for category in ("Kostak", "Subject To", "Premium", "Options"):
                        table = self.page.locator("#shareGroupModalTablesContainer table").filter(
                            has=self.page.locator("thead > tr:first-child > th").filter(
                                has_text=re.compile(r"^\s*" + re.escape(category) + r"\s*$")))
                        if category == "Options":
                            self.verify_options(table, group.GroupName)
                        else:
                            self.verify_table(table, group.GroupName, category)
                    LOG.info("Verified page + Share modal: %s", group.GroupName)
            LOG.info("Completed %s numeric assertions across %s groups", self.checked, len(self.groups))
            LOG.info("Completed %s direct workbook-to-browser numeric assertions",
                     len(self.browser_reference_results))

    return CalculationVerification


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--orders-sheet", default="Sheet1")
    parser.add_argument("--details-sheet", default="Sheet2")
    parser.add_argument("--details-source", type=Path)
    parser.add_argument("--group-reference", type=Path)
    parser.add_argument("--client-reference", type=Path)
    parser.add_argument("--reference-report", type=Path, help="Save worksheet-expected versus browser-rendered comparisons as JSON")
    parser.add_argument("--expected-rows", type=int, default=840)
    parser.add_argument("--expected-details", type=int, default=9388)
    parser.add_argument("--headed", action="store_true")
    parser.add_argument("--failfast", action="store_true", help="Stop at the first failed calculation or locator")
    parser.add_argument("--export-fixture", type=Path, help="Save the fully seeded Django fixture for manual editing/reuse")
    mode = parser.add_mutually_exclusive_group(required=True)
    mode.add_argument("--autoinject", action="store_true")
    mode.add_argument("--fixture", type=Path, help="Load a manually prepared Django fixture instead of auto-seeding")
    args = parser.parse_args()
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
    # Validate before Django creates even the isolated database.
    parse_dataset(args)
    sys.path.insert(0, str(ROOT))
    from django.conf import settings
    from userproject import settings as project_settings
    if settings.configured:
        raise RuntimeError("Run this file directly in a fresh Python process")
    with tempfile.TemporaryDirectory(prefix="ipo-calculations-") as run_dir:
        db = str(Path(run_dir) / "calculations.sqlite3")
        (Path(run_dir) / "static").mkdir()
        (Path(run_dir) / "media").mkdir()
        config = {k: getattr(project_settings, k) for k in dir(project_settings) if k.isupper()}
        config.update(DATABASES={"default": {"ENGINE": "django.db.backends.sqlite3", "NAME": db,
                                             "TEST": {"NAME": db}}},
                      CALCULATION_DB=db, STATIC_ROOT=str(Path(run_dir) / "static"),
                      MEDIA_ROOT=str(Path(run_dir) / "media"),
                      STATICFILES_STORAGE="django.contrib.staticfiles.storage.StaticFilesStorage",
                      EMAIL_BACKEND="django.core.mail.backends.locmem.EmailBackend",
                      PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"],
                      SECRET_KEY="calculation-test-only")
        settings.configure(**config)
        import django
        django.setup()
        from django.test.runner import DiscoverRunner
        runner = DiscoverRunner(verbosity=2, interactive=False, parallel=1, failfast=args.failfast)
        runner.setup_test_environment()
        old_config = None
        try:
            old_config = runner.setup_databases()
            suite = unittest.defaultTestLoader.loadTestsFromTestCase(make_case(args))
            result = runner.run_suite(suite)
            return 0 if result.wasSuccessful() else 1
        finally:
            if old_config is not None:
                runner.teardown_databases(old_config)
            runner.teardown_test_environment()


if __name__ == "__main__":
    sys.exit(main())
